# Teste da aba Plano de Trabalho (backend).
#
# Como rodar:  python3 docs/testes/plano-back.py
#   - Sempre: funções PURAS (parser da planilha, de-para, resolução nome→login,
#     regras de gravação, atraso).
#   - Com PLANO_TEST_DB=<dsn de um Postgres DESCARTÁVEL>: também as ROTAS, pelo
#     test client do Flask — importação, idempotência, "a tela vence",
#     permissões (cp_totvs, cp_cliente, consultor próprio/alheio, usuário-chave,
#     leitor), simulação, regras 1–3 e histórico. O teste CRIA o schema mínimo
#     do cockpit nesse banco e APAGA tudo no começo: nunca aponte para o Supabase.
#
# A planilha é SINTÉTICA (montada aqui com openpyxl), com as mesmas armadilhas
# do arquivo real da Açosul: duas grafias de "Produto", linha de total sem
# atividade, linhas vazias formatadas, empresas sobre a coluna de responsável,
# rotina repetida (FISA170) e datas que variam por módulo. O repo é público —
# dado de cliente não entra aqui.
import io, os, pathlib, sys
from datetime import date, datetime

RAIZ = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "api"))
os.environ.setdefault("SESSION_SECRET", "teste-plano")
if os.environ.get("PLANO_TEST_DB"):
    os.environ["DATABASE_URL"] = os.environ["PLANO_TEST_DB"]
import index as A          # noqa: E402
import openpyxl            # noqa: E402

OK = 0


def ok(cond, msg):
    global OK
    if not cond:
        raise AssertionError(msg)
    OK += 1


def planilha(status_fat2="Finalizado", obs_fin=None, sem_marco_unico=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financeiro"
    ws.append(["Sequência", "Etapa", "Tipo de Agenda", "Atividade", "Data De", "Data Até",
               "Nro. Horas", "Responsável", "Usuário/Consultor", "Produto da Atividade",
               "Status", "Observações"])
    ws.append([1, "Parametrização", "Individual", "Re-parametrização da base",
               datetime(2026, 10, 17), datetime(2026, 10, 23), 0, "Totvs", "Ana", "Não há",
               "Finalizado", None])
    ws.append([2, "Treinamento", "Individual", "Treinamento de processos",
               datetime(2026, 10, 20), datetime(2026, 10, 21), 18, "Totvs", "Ana",
               "MIT037/MIT010", "Finalizado com Restrições", obs_fin or "MITs aguardando aprovação"])
    ws.append([3, "Parametrização", "Individual", "Entrega dos cadastros",
               datetime(2026, 10, 22), datetime(2026, 11, 13), 0, "Cliente X", "Lucas",
               "MIT006 preenchida", "Agendado", None])
    ws.append([None, None, None, None, None, None, 18, None, None, None, None, None])  # total
    for _ in range(5):
        ws.append([None] * 12)
    ws2 = wb.create_sheet("Estoque_Custos")
    ws2.append(["Sequência", "Etapa", "Tipo de Agenda", "Atividade", "Data De", "Data Até",
                "Nro. Horas", "Responsável", "Usuário/Consultor", "Produto",
                "Status", "Observações"])
    ws2.append([1, "Treinamento", "Coletiva", "Treinamento estoque",
                datetime(2026, 11, 25), datetime(2026, 11, 26), 9, "Totvs", "Bruno Souza",
                None, status_fat2, None])
    ws2.append([2, "Protótipo Integrado", "Coletiva", "Primeiro protótipo",
                datetime(2026, 12, 2), datetime(2026, 12, 4), 13.5, "Totvs", "Bruno",
                None, None, None])
    ws3 = wb.create_sheet("Lista de Tarefas_ClienteX")
    ws3.append([None, None, None, "Cliente X ", "Outra Empresa ", None])
    ws3.append(["Módulo", "Cadastro/Entidade Básica", "Consultor", "Pessoa Responsável (Cliente X)",
                None, None, "Ambiente", "Base Liberada - A partir ", "Data Limite - Treinamentos",
                "Data Limite - Protótipo", "Data Limite - Go Live"])
    go = datetime(2026, 12, 17)
    ws3.append(["Financeiro", "MATA070 - Bancos", "Ana Lima", "Lucas", "Pedro", None, "Produção",
                datetime(2026, 10, 22), datetime(2026, 11, 20), datetime(2026, 11, 27), go])
    ws3.append(["Fiscal", "FISA170 - Perfis Fiscais", "Bruno Souza", None, None, None, "Produção",
                datetime(2026, 10, 22), datetime(2026, 10, 29), datetime(2026, 11, 27),
                datetime(2026, 12, 18) if sem_marco_unico else go])
    ws3.append(["Estoque", "MATA010 - Produtos", "Bruno Souza", None, None, None, "Produção",
                datetime(2026, 10, 21), datetime(2026, 10, 29), datetime(2026, 11, 27), go])
    ws3.append(["Fiscal", "FISA170 - cBenef", "Bruno Souza", None, None, None, "Produção",
                datetime(2026, 10, 22), datetime(2026, 10, 29), datetime(2026, 11, 27), go])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HOJE = date(2026, 10, 27)

# ── funções puras ───────────────────────────────────────────────────────────
abas = A._plano_le_xlsx(planilha())
ok(all(len(v) < 20 for v in abas.values()), "linhas vazias do fim cortadas")
p = A._plano_parse(abas, "CLIENTE X COMERCIO LTDA", hoje=HOJE)
ok(len(p["atividades"]) == 5, f"5 atividades (veio {len(p['atividades'])})")
ok(len(p["cadastros"]) == 4, "4 cadastros")
prod = next(c for c in p["cadastros"] if c["rotina"] == "MATA010")
ok(prod["modulo"] == "Estoque/Custos" and prod["origem_chave"] == "C|ESTOQUECUSTOS|MATA010|PRODUTOS",
   f"'Estoque' do cadastro casa com a aba 'Estoque_Custos' ({prod['modulo']})")
ok(p["ignoradas"] == 1, "linha de total ignorada e contada")
fin2 = next(a for a in p["atividades"] if a["origem_chave"] == "A|FINANCEIRO|2")
ok(fin2["produto"] == "MIT037/MIT010", "'Produto da Atividade' lido")
ok(fin2["status"] == "finalizado_restricoes", "de-para Finalizado com Restrições")
ok(fin2["horas"] == 18.0 and fin2["data_de"] == "2026-10-20", "horas e data")
ok(fin2["responsavel_lado"] == "totvs", "Totvs → lado totvs")
ok("sem motivo" in (fin2["inconsistencia"] or ""), "restrição sem motivo apontada")
fin3 = next(a for a in p["atividades"] if a["origem_chave"] == "A|FINANCEIRO|3")
ok(fin3["responsavel_lado"] == "cliente", "empresa do cliente → lado cliente")
est1 = next(a for a in p["atividades"] if a["modulo"] == "Estoque/Custos" and a["sequencia"] == 1)
ok("começa em 25/11/2026" in (est1["inconsistencia"] or ""), "finalizado futuro apontado, não barrado")
ok(est1["status"] == "finalizado", "importa fiel o status da planilha")
est2 = next(a for a in p["atividades"] if a["modulo"] == "Estoque/Custos" and a["sequencia"] == 2)
ok(est2["status"] == "agendado" and est2["produto"] is None, "status vazio → agendado; 'Produto' lido")
bancos = next(c for c in p["cadastros"] if c["rotina"] == "MATA070")
ok(bancos["responsavel_nome"] == "Lucas", "coluna da empresa do cliente (não a outra)")
ok(bancos["entidade"] == "Bancos" and bancos["consultor_nome"] == "Ana Lima", "rotina/entidade/consultor")
ok(bancos["data_de"] == "2026-10-22" and bancos["data_ate"] == "2026-11-20", "período do cadastro por linha")
ok(len({c["origem_chave"] for c in p["cadastros"]}) == 4, "FISA170 repetida gera chaves distintas")
ok(p["marcos"] == {"prototipo": "2026-11-27", "go_live": "2026-12-17"},
   f"só coluna com data única vira marco ({p['marcos']})")
p2 = A._plano_parse(A._plano_le_xlsx(planilha(sem_marco_unico=True)), "Cliente X", hoje=HOJE)
ok("go_live" not in p2["marcos"], "go-live divergente não vira marco")
try:
    A._plano_parse({"X": [["a", "b"], [1, 2]]}, "", hoje=HOJE)
    ok(False, "planilha sem cabeçalho deveria falhar")
except ValueError:
    ok(True, "")

ok(A._plano_status_de("FINALIZADO", "atividade") == ("finalizado", None), "de-para caixa alta")
ok(A._plano_status_de("Xpto", "atividade")[1], "status desconhecido avisa")
ok(A._plano_status_de("Validado", "atividade")[1], "status de cadastro numa atividade avisa")
ok(A._plano_rotina("MATA030 - Clientes") == ("MATA030", "Clientes"), "rotina")
ok(A._plano_rotina("Cadastro de feriados") == (None, "Cadastro de feriados"), "sem rotina")

C = [{"email": "fabricio.sandrade@totvs.com.br", "nome": "Fabricio Sandrade"},
     {"email": "mauricio.vitto@totvs.com.br", "nome": "Mauricio Vitto"},
     {"email": "mauricio.souza@totvs.com.br", "nome": "Mauricio Souza"},
     {"email": "elayne.climaco@totvs.com.br", "nome": "Elayne Climaco"}]
ok(A._plano_resolve("Fabrício Andrade", C) == ("fabricio.sandrade@totvs.com.br", False),
   "sobrenome contido no login (Andrade ⊂ Sandrade)")
ok(A._plano_resolve("Maurício Vitto", C) == ("mauricio.vitto@totvs.com.br", False), "nome completo")
ok(A._plano_resolve("Maurício", C) == (None, True), "ambíguo não adivinha")
ok(A._plano_resolve("Lucas", C) == (None, False), "sem login")
ok(A._plano_resolve("ELAYNE", C)[0] == "elayne.climaco@totvs.com.br", "caixa/acento")

at = {"tipo": "atividade", "status": "agendado", "data_de": date(2026, 11, 25),
      "data_ate": date(2026, 11, 26), "motivo_restricao": None}
ok("começa em 25/11/2026" in A._plano_valida(at, {"status": "finalizado"}, HOJE)[0], "regra 1")
passada = {**at, "data_de": date(2026, 10, 20), "data_ate": date(2026, 10, 21)}
ok(A._plano_valida(passada, {"status": "finalizado"}, HOJE) == (None, False), "finaliza passada")
ok("exige o motivo" in A._plano_valida(passada, {"status": "finalizado_restricoes"}, HOJE)[0], "regra 2")
ok(A._plano_valida(passada, {"status": "finalizado_restricoes", "motivo_restricao": "x"}, HOJE)[0] is None,
   "restrição com motivo")
ok(A._plano_valida(at, {"data_ate": "2026-11-28"}, HOJE) == (None, True), "regra 3: reagendou")
ok(A._plano_valida({**at, "status": "finalizado"}, {"data_ate": "2026-11-28"}, HOJE)[1] is False,
   "mexer em data de finalizada não conta reagendamento")
ok("anterior" in A._plano_valida(at, {"data_ate": "2026-11-01"}, HOJE)[0], "datas invertidas")
inc = {**at, "status": "finalizado"}
ok(A._plano_valida(inc, {"observacoes": "x"}, HOJE)[0] is None,
   "item inconsistente aceita editar observação (regra só no que mudou)")
ok(A._plano_valida({"tipo": "cadastro", "status": "nao_iniciado"}, {"status": "finalizado"}, HOJE)[0],
   "status de atividade num cadastro")

ok(A._plano_atrasado({"tipo": "atividade", "status": "agendado", "data_ate": "2026-10-26"}, {}, HOJE),
   "atividade vencida")
ok(not A._plano_atrasado({"tipo": "atividade", "status": "finalizado", "data_ate": "2026-10-26"}, {}, HOJE),
   "finalizada não atrasa")
ok(A._plano_atrasado({"tipo": "cadastro", "status": "em_cadastro", "data_ate": None},
                     {"prototipo": "2026-10-01"}, HOJE), "cadastro sem limite cai no marco do protótipo")
ok(not A._plano_atrasado({"tipo": "cadastro", "status": "entregue", "data_ate": "2026-10-01"}, {}, HOJE),
   "cadastro entregue não atrasa")
ok(A._cp_hoje().isoformat() == datetime.now(A.TZ_SP).date().isoformat(), "hoje em UTC-3")
print(f"puras: {OK} asserções OK")

# ── rotas (Postgres descartável) ────────────────────────────────────────────
DSN = os.environ.get("PLANO_TEST_DB")
if not DSN:
    print("PLANO_TEST_DB não definido — rotas não testadas.")
    sys.exit(0)

import psycopg2  # noqa: E402

SCHEMA = (RAIZ / "docs" / "sql" / "2026-10-27-plano-trabalho.sql").read_text(encoding="utf-8")
with psycopg2.connect(DSN) as c, c.cursor() as cur:
    c.autocommit = True
    cur.execute("drop schema if exists cockpit cascade; create schema cockpit;")
    cur.execute("""
      create table cockpit.clientes (customer text primary key, nome text, tipo text);
      create table cockpit.usuarios_login (email text primary key, senha_hash text, nome text,
        ativo boolean default true, is_admin boolean default false, perfil text);
      create table cockpit.usuario_clientes (email text, customer text, abas text[], papel text,
        primary key (email, customer));""")
    cur.execute(SCHEMA)
    cur.execute("""
      insert into cockpit.clientes values ('CLIX0000', 'CLIENTE X COMERCIO', 'novo');
      insert into cockpit.usuarios_login (email, nome, perfil) values
        ('vivi@totvs.com.br', 'Viviane CP', 'comum'),
        ('ana.lima@totvs.com.br', 'Ana Lima', 'consultoria'),
        ('bruno.souza@totvs.com.br', 'Bruno Souza', 'consultoria'),
        ('lucas@clientex.com.br', 'Lucas Cliente', 'cliente'),
        ('key@clientex.com.br', 'Usuario Chave', 'cliente'),
        ('leitor@totvs.com.br', 'Leitor', 'leitor'),
        ('semaba@clientex.com.br', 'Sem Aba', 'cliente');
      insert into cockpit.usuarios_login (email, nome, perfil, is_admin) values
        ('adm@totvs.com.br', 'Admin', 'admin', true);
      insert into cockpit.usuario_clientes values
        ('vivi@totvs.com.br', 'CLIX0000', null, 'cp_totvs'),
        ('ana.lima@totvs.com.br', 'CLIX0000', null, 'consultor'),
        ('bruno.souza@totvs.com.br', 'CLIX0000', null, 'consultor'),
        ('lucas@clientex.com.br', 'CLIX0000', '{plano}', 'cp_cliente'),
        ('key@clientex.com.br', 'CLIX0000', '{plano}', 'usuario_chave'),
        ('leitor@totvs.com.br', 'CLIX0000', null, 'cp_totvs'),
        ('semaba@clientex.com.br', 'CLIX0000', '{cadastros}', 'cp_cliente');""")

app = A.app.test_client()
PJ = "/api/plano/CLIX0000/CLIX000001"


def como(email, view_as=None):
    app.set_cookie(A.COOKIE_NAME, A.make_session(email, email, view_as=view_as), domain="localhost")


def importa(dados, nome="plano.xlsx"):
    return app.post(PJ + "/importar", data={"arquivo": (io.BytesIO(dados), nome)},
                    content_type="multipart/form-data")


como("key@clientex.com.br")
ok(importa(planilha()).status_code == 403, "usuário-chave não importa")
como("semaba@clientex.com.br")
ok(app.get(PJ).status_code == 403, "sem a aba liberada não lê")
como("vivi@totvs.com.br")
r = app.post(PJ + "/marcos", json={"marcos": {"go_live": "2027-01-04"}})
ok(r.status_code == 200 and r.json["marcos"]["go_live"] == "2027-01-04", "CP grava go-live oficial")
r = importa(planilha())
ok(r.status_code == 200, f"importa ({r.status_code} {r.json})")
j = r.json
ok(j["criados"] == 9 and j["alterados"] == 0, f"9 criados ({j['criados']}/{j['alterados']})")
ok(j["ignorados"] == 1, "total ignorado")
ok(j["sem_login"] == [] and j["ambiguos"] == [], f"todos resolvidos ({j['sem_login']}/{j['ambiguos']})")
ok(len(j["inconsistencias"]) == 2, f"2 inconsistências ({j['inconsistencias']})")
g = app.get(PJ).json
ok(g["marcos"]["go_live"] == "2027-01-04", "go-live da planilha NÃO sobrescreve o oficial")
ok(g["marcos"]["prototipo"] == "2026-11-27", "protótipo veio da planilha")
ok(g["permissao"] == "tudo", "cp_totvs = tudo")
itens = {i["origem_chave"]: i for i in g["itens"]}
ok(itens["A|FINANCEIRO|1"]["responsavel_email"] == "ana.lima@totvs.com.br", "Ana → login")
ok(itens["C|FINANCEIRO|MATA070|BANCOS"]["consultor_email"] == "ana.lima@totvs.com.br", "consultor do cadastro")

r = importa(planilha())
ok(r.json["criados"] == 0 and r.json["alterados"] == 0, "reimportar igual não muda nada")

fin1 = itens["A|FINANCEIRO|1"]["id"]
est1 = itens["A|ESTOQUECUSTOS|1"]["id"]
fin3 = itens["A|FINANCEIRO|3"]["id"]

# consultora Ana: só os itens dela e só status/observação
como("ana.lima@totvs.com.br")
ok(app.get(PJ).json["permissao"] == "proprio", "consultor = proprio")
r = app.post(f"/api/plano/item/{fin1}", json={"observacoes": "ok, base revisada"})
ok(r.status_code == 200, f"consultora edita a própria ({r.status_code} {r.json})")
ok(app.post(f"/api/plano/item/{fin1}", json={"atividade": "outra"}).status_code == 403,
   "consultor não muda descrição")
ok(app.post(f"/api/plano/item/{est1}", json={"observacoes": "x"}).status_code == 403,
   "consultor não mexe no item de outro")

# leitor com papel de CP continua só lendo
como("leitor@totvs.com.br")
ok(app.get(PJ).json["permissao"] is None, "leitor não grava nem com papel")
ok(app.post(f"/api/plano/item/{fin1}", json={"observacoes": "x"}).status_code == 403, "leitor 403")

# CP do cliente edita tudo
como("lucas@clientex.com.br")
ok(app.get(PJ).json["permissao"] == "tudo", "cp_cliente = tudo")
r = app.post(f"/api/plano/item/{est1}", json={"status": "agendado"})
ok(r.status_code == 200 and r.json["item"]["inconsistencia"] is None,
   "corrigir na tela limpa a inconsistência")
r = app.post(f"/api/plano/item/{est1}", json={"status": "finalizado"})
ok(r.status_code == 422, "regra 1 na rota")
r = app.post(f"/api/plano/item/{fin3}", json={"data_ate": "2026-11-14"})
ok(r.status_code == 200 and r.json["reagendou"] and r.json["item"]["qtd_reagendamentos"] == 1,
   "regra 3 na rota")
ok(app.post(f"/api/plano/item/{fin3}", json={"status": "finalizado_restricoes"}).status_code == 422,
   "regra 2 na rota")
r = app.post(PJ + "/item", json={"tipo": "atividade", "modulo": "Financeiro",
                                 "atividade": "Nova reunião", "data_de": "2026-11-03",
                                 "data_ate": "2026-11-03", "horas": 2})
ok(r.status_code == 200 and r.json["item"]["ordem"] == 4, "cria atividade no fim do módulo")
novo = r.json["item"]["id"]
ok(app.post(f"/api/plano/item/{novo}/inativar").status_code == 200, "inativa")
ok(all(i["id"] != novo for i in app.get(PJ).json["itens"]), "inativo some da lista")

h = app.get(f"/api/plano/item/{fin3}/hist").json["hist"]
campos = [x["campo"] for x in h]
ok("data_ate" in campos and "criado" in campos, f"histórico ({campos})")
ok(h[0]["origem"] == "tela" and h[-1]["origem"] == "importacao", "origem tela/importação")

# reimportar NÃO desfaz a tela
como("vivi@totvs.com.br")
r = importa(planilha())
ok(r.json["preservados"] >= 2, f"tela vence na reimportação ({r.json['preservados']})")
g = {i["id"]: i for i in app.get(PJ).json["itens"]}
ok(g[est1]["status"] == "agendado" and not g[est1]["inconsistencia"], "status da tela mantido")
ok(g[fin3]["data_ate"] == "2026-11-14", "data da tela mantida")
ok(g[fin1]["observacoes"] == "ok, base revisada", "observação da tela mantida")

# planilha mudou de verdade num campo NÃO tocado na tela → atualiza
r = importa(planilha(obs_fin="Itens pendentes: movimento bancário"))
fin2 = next(i for i in app.get(PJ).json["itens"] if i["origem_chave"] == "A|FINANCEIRO|2")
ok(r.json["alterados"] == 1 and fin2["observacoes"].startswith("Itens pendentes"),
   "campo não tocado na tela acompanha a planilha")

# simulação não grava
como("adm@totvs.com.br", view_as="vivi@totvs.com.br")
ok(app.post(f"/api/plano/item/{fin1}", json={"observacoes": "y"}).status_code == 409, "ver como → 409")
como("adm@totvs.com.br")
ok(app.get(PJ).json["permissao"] == "tudo", "admin = tudo")
ok(app.post("/api/plano/item/nao-e-uuid", json={}).status_code == 404, "id inválido → 404")
print(f"total: {OK} asserções OK (puras + rotas)")
