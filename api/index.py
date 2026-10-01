#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Projetos · TOTVS SC — backend serverless (Vercel)
=================================================
Porte do Projetos/server.py + pci_client.py para função serverless.

- Lista de projetos: lida de cockpit.projetos (Supabase), sincronizada de HORA
  em HORA pelo pg_cron do Supabase (jobid 7 "projetos-sync" -> GET
  /api/cron/sync com Bearer CRON_SECRET) e sob demanda pelo botão "Sincronizar
  API" (POST /api/sync?page=N, paginado pelo front para não estourar os 60s).
- RECORTE: só entram projetos da regional (região do coordenador titular ou
  auxiliar em SYNC_REGIOES), de clientes de cockpit.clientes, ou de clientes da
  regional com projeto ainda vivo. Ver no_recorte() — e a MESMA regra em SQL em
  purga_fora_do_recorte().
- Detalhe (mapa/cronograma): AO VIVO na API PCI, com cache curto em memória.
- Login e recorte por cliente: mesmos do ecossistema (cockpit.usuarios_login +
  cockpit.usuario_clientes). Admin vê todos os clientes.
"""
from __future__ import annotations

import base64, csv, email.message, hashlib, hmac, imaplib, io, json, os, re, secrets, \
    time, unicodedata, urllib.parse
from pathlib import Path

import psycopg2, psycopg2.extras, requests
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from psycopg2.extras import execute_values
from flask import Flask, Response, request, redirect, make_response
from werkzeug.exceptions import HTTPException

BASE_DIR = Path(__file__).resolve().parent.parent
WEB_DIR = BASE_DIR / "web"          # NÃO usar "public/": a Vercel serviria estático
                                    # antes da função e furaria a porta de login.

DATABASE_URL = os.environ.get("DATABASE_URL", "")
TASKS_BASE = os.environ.get("TASKS_SC_BASE_URL", "https://api.tscst.com.br/restAPI").rstrip("/")
TASKS_USER = os.environ.get("TASKS_USERNAME", "")
TASKS_PASS = os.environ.get("TASKS_PASSWORD", "")
SESSION_SECRET = os.environ.get("SESSION_SECRET", "dev-insecure-secret")
CRON_SECRET = os.environ.get("CRON_SECRET", "")

# ── Recorte do sync (regional SC Sul + clientes do Cockpit) ─────────────────
# A API PCI devolve os ~1.7 mil projetos de TODA a TOTVS SC. Só interessam os da
# regional e os dos clientes atendidos: sem isso a lista vira palheiro e o cron
# grava 700 projetos que ninguém abre. Um projeto entra se a REGIÃO do
# coordenador (titular OU auxiliar) estiver em SYNC_REGIOES, ou se o cliente
# estiver em cockpit.clientes. SYNC_REGIOES vazio = sem filtro (volta ao antigo).
SYNC_REGIOES = {r.strip() for r in os.environ.get("SYNC_REGIOES", "201,202,211").split(",")
                if r.strip()}
# O cron apaga o que está fora do recorte (projeto que trocou de coordenador/
# região some da base). SYNC_PURGE=0 desliga.
SYNC_PURGE = os.environ.get("SYNC_PURGE", "1").strip().lower() not in ("0", "false", "off")
# Terceiro critério: cliente DA regional cujo projeto ainda está vivo, mesmo com
# coordenador de outra região (ex.: PRODUZA/203, GROWTH/601, SOLFACIL/302 — todos
# clientes 201). Sem o corte por status isso arrastaria 644 projetos finalizados.
SYNC_ENCERRADOS = {"finalizado", "cancelado"}
SESSION_TTL = 12 * 3600
COOKIE_NAME = "proj_sess"

LISTA_URL = f"{TASKS_BASE}/custom/tscst/pci/api/v1/projetos"
MAPA_URL = f"{TASKS_BASE}/PCITConectaProjetos/mapa"
CRONO_URL = f"{TASKS_BASE}/PCITConectaProjetos/cronograma"

app = Flask(__name__)

# ── Roteamento na Vercel — o contrato "?__path=" ────────────────────────────
# Com `destination: "/api/index"` (seco), a Vercel entrega à função o caminho de
# DESTINO: o Flask recebe PATH_INFO=/api/index em TODA request, nenhuma rota casa
# e tudo cai no catch-all `/<path:asset>` → 404 "Rota de API desconhecida." (o
# site inteiro morre, inclusive `/` e `/api/health`). O `vercel.json` passa o
# caminho original em `?__path=/$1` e este middleware o devolve ao PATH_INFO.
# Os dois andam em par: mexeu no destination, mexa aqui.
FUNC_PATHS = ("/api/index", "/api/index.py")


class _VercelRewritePath:
    """Devolve ao PATH_INFO o caminho original vindo em ?__path=.
    Inerte quando o PATH_INFO já chega certo (dev local ou se a Vercel voltar a
    preservar o caminho)."""

    def __init__(self, wsgi):
        self.wsgi = wsgi

    def __call__(self, environ, start_response):
        if environ.get("PATH_INFO", "") in FUNC_PATHS:
            pares = urllib.parse.parse_qsl(environ.get("QUERY_STRING", ""),
                                           keep_blank_values=True)
            resto = [(k, v) for k, v in pares if k != "__path"]
            path = next((v for k, v in pares if k == "__path"), "") \
                or environ.get("HTTP_X_VERCEL_ORIGINAL_PATH", "")
            if path:
                if "?" in path:                      # query que veio no próprio $1
                    path, extra = path.split("?", 1)
                    resto += urllib.parse.parse_qsl(extra, keep_blank_values=True)
                environ["PATH_INFO"] = path if path.startswith("/") else "/" + path
                environ["QUERY_STRING"] = urllib.parse.urlencode(resto)
        return self.wsgi(environ, start_response)


app.wsgi_app = _VercelRewritePath(app.wsgi_app)


class PCIUnavailable(Exception):
    """API PCI indisponível (timeout/5xx persistente) — erro esperado."""


# ── util ────────────────────────────────────────────────────────────────────
def _json(o, code=200):
    return Response(json.dumps(o, ensure_ascii=False, default=str), status=code,
                    mimetype="application/json")


def _err(code, msg):
    return _json({"ok": False, "error": msg}, code)


def _slug(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


@app.errorhandler(Exception)
def _on_error(e):
    if isinstance(e, HTTPException):
        return e
    if isinstance(e, PCIUnavailable):
        return _json({"ok": False, "error": str(e), "api_indisponivel": True}, 503)
    return _json({"ok": False, "error": f"{type(e).__name__}: {e}"}, 500)


# ── Postgres (Supabase) ─────────────────────────────────────────────────────
def db():
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL não configurada.")
    c = psycopg2.connect(DATABASE_URL, connect_timeout=10)
    c.autocommit = True
    return c


def q(sql, params=None, one=False):
    with db() as c, c.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute(sql, params or ())
        if cur.description is None:
            return None
        rows = cur.fetchall()
        return (rows[0] if rows else None) if one else rows


def execute(sql, params=None):
    with db() as c, c.cursor() as cur:
        cur.execute(sql, params or ())


# ── Sessão / login (mesmo padrão do ecossistema) ────────────────────────────
def _sign(p):
    return hmac.new(SESSION_SECRET.encode(), p.encode(), hashlib.sha256).hexdigest()


def make_session(email, nome, view_as=None):
    d = {"e": email, "n": nome, "x": int(time.time()) + SESSION_TTL}
    if view_as:
        d["v"] = view_as          # admin simulando a visão de outro usuário
    raw = json.dumps(d)
    b = base64.urlsafe_b64encode(raw.encode()).decode()
    return f"{b}.{_sign(b)}"


def read_session():
    tok = request.cookies.get(COOKIE_NAME, "")
    if not tok or "." not in tok:
        return None
    b, sig = tok.rsplit(".", 1)
    if not hmac.compare_digest(sig, _sign(b)):
        return None
    try:
        d = json.loads(base64.urlsafe_b64decode(b.encode()).decode())
    except Exception:
        return None
    return d if int(d.get("x", 0)) > time.time() else None


def current_user():
    """Usuário REAL do login (nunca o simulado). Use para auditoria/escrita."""
    s = read_session()
    return s.get("e") if s else None


def is_admin(email):
    if not email:
        return False
    r = q("select coalesce(is_admin,false) adm from cockpit.usuarios_login "
          "where lower(email)=%s", (email.lower(),), one=True)
    return bool(r and r["adm"])


def effective_user():
    """Usuário cuja VISÃO vale. É o simulado (view_as) somente se quem está
    logado for admin de verdade — a simulação jamais amplia acesso, só restringe."""
    s = read_session()
    if not s:
        return None
    alvo = s.get("v")
    if alvo and is_admin(s.get("e")):
        return alvo
    return s.get("e")


def require_auth():
    return None if current_user() else _err(401, "Não autenticado.")


def require_admin():
    if not current_user():
        return _err(401, "Não autenticado.")
    if not is_admin(current_user()):
        return _err(403, "Apenas administradores.")
    return None


def verify_scrypt(stored, senha):
    """scrypt$salt$hash. O Cockpit gera com Node scryptSync(pw, saltString, 64):
    o salt vai como STRING (o hex em UTF-8). Tentamos essa variante primeiro e,
    como fallback, o salt decodificado de hex (formato legado)."""
    try:
        scheme, salt_hex, hash_hex = stored.split("$", 2)
    except ValueError:
        return False
    if scheme != "scrypt":
        return False
    dklen = len(hash_hex) // 2
    variants = [salt_hex.encode()]
    try:
        variants.append(bytes.fromhex(salt_hex))
    except ValueError:
        pass
    for salt in variants:
        for n in (16384, 32768, 8192, 65536):
            try:
                dk = hashlib.scrypt(senha.encode(), salt=salt, n=n, r=8, p=1,
                                    dklen=dklen, maxmem=132 * 1024 * 1024)
            except Exception:
                continue
            if hmac.compare_digest(dk.hex(), hash_hex):
                return True
    return False


# ── Acesso por cliente (admin vê tudo; comum só os liberados) ───────────────
def allowed_customers():
    email = effective_user()      # respeita o "ver como"
    if not email:
        return set()
    row = q("select coalesce(is_admin,false) adm from cockpit.usuarios_login "
            "where lower(email)=%s", (email.lower(),), one=True)
    if row and row["adm"]:
        return None
    rows = q("select customer from cockpit.usuario_clientes where lower(email)=%s",
             (email.lower(),))
    return {r["customer"] for r in rows}


def deny_customer(cust):
    a = allowed_customers()
    return None if (a is None or cust in a) else _err(403, "Sem acesso a este cliente.")


# ── Perfil, abas liberadas e a ÁREA DE CLIENTES ─────────────────────────────
# TRÊS EIXOS, e confundi-los abre acesso:
#   allowed_customers() → QUAIS CLIENTES ele vê
#   abas_liberadas()    → QUE ABAS ele vê NAQUELE cliente
#   perfil              → ele é da TOTVS (interno) ou é o cliente
#
# Esconder o botão na tela NÃO é proteger: quem não é interno é barrado aqui,
# no servidor (require_interno), mesmo chamando a rota na mão.
PERFIS = ("admin", "comum", "consultoria", "cliente", "leitor")
# 'consultoria' = consultor que trabalha nos projetos (TOTVS ou parceiro): tem o
# MESMO poder de 'comum', só o rótulo é próprio, para dar para separar quem é
# consultoria de quem é do time fixo sem inventar um terceiro eixo de permissão.
# ARMADILHA: a tabela usuarios_login é compartilhada com o cockpit-unico-tsc, e
# lá perfil desconhecido degrada para 'leitor' — o perfil novo tem que entrar no
# CHECK do banco E no src/lib/auth/perfis.ts antes de ser usado de verdade.
PERFIS_INTERNOS = ("admin", "comum", "consultoria", "leitor")   # gente da TOTVS

# Catálogo das abas do detalhe do projeto. 'cliente' diz se a aba PODE ser
# liberada para um usuário do cliente — 'consumo' fica de fora de propósito
# (é margem/valoração interna, não se mostra para o cliente).
ABAS = [
    {"id": "resumo",    "label": "Resumo",               "cliente": True},
    {"id": "crono",     "label": "Cronograma",           "cliente": True},
    {"id": "modulos",   "label": "Por Módulo",           "cliente": True},
    {"id": "etapas",    "label": "Por Etapa",            "cliente": True},
    {"id": "consumo",   "label": "Consumo",              "cliente": False},
    {"id": "gaps",      "label": "GAPs",                 "cliente": True},
    {"id": "tarefas",   "label": "Tarefas",              "cliente": True},
    {"id": "cadastros", "label": "Cadastros e Movimentos", "cliente": True},
    {"id": "cobertura", "label": "Cobertura",            "cliente": True},
    {"id": "empresas",  "label": "Empresas e Compartilhamento", "cliente": True},
    {"id": "estrutura", "label": "Estrutura de Empresas", "cliente": False},
    {"id": "compara",   "label": "Compara (dicionário)",  "cliente": False},
    {"id": "prototipo", "label": "Protótipo",            "cliente": True},
    {"id": "transicao", "label": "Transição",            "cliente": True},
]
ABAS_IDS = [a["id"] for a in ABAS]
ABAS_CLIENTE_OK = {a["id"] for a in ABAS if a["cliente"]}
# O que a área de clientes libera quando a liberação existe mas ninguém marcou
# aba nenhuma (coluna abas NULL). Protótipo de 03/09/2026: só Cadastros e
# Movimentos. Lista vazia ('{}') é diferente: aí é "nenhuma aba", de propósito.
ABAS_PADRAO_CLIENTE = ["cadastros"]


def perfil_do(email):
    """admin | comum | cliente | leitor. MESMA regra do cockpit-unico-tsc
    (src/lib/auth/perfis.ts): perfil VAZIO cai no is_admin legado; perfil
    PREENCHIDO mas desconhecido degrada para 'leitor', nunca para 'comum' —
    a tabela é compartilhada e um perfil novo pode chegar antes do deploy."""
    if not email:
        return "leitor"
    r = q("select perfil, coalesce(is_admin,false) adm from cockpit.usuarios_login "
          "where lower(email)=%s and ativo", (email.lower(),), one=True)
    if not r:
        return "leitor"
    p = (r["perfil"] or "").strip()
    if p in PERFIS:
        return p
    if not p:
        return "admin" if r["adm"] else "comum"
    return "leitor"


def eh_interno(email=None):
    """Interno = TOTVS. Usa o usuário EFETIVO: 'ver como' um cliente tem que
    mostrar a tela do cliente, senão a simulação não serve para conferir."""
    return perfil_do(email or effective_user()) in PERFIS_INTERNOS


def abas_liberadas(customer=None):
    """Abas do usuário efetivo. None = todas (interno).
    Sem customer devolve o mapa {customer: set(abas)} — é o que o front usa
    para montar a barra de abas do projeto aberto."""
    email = effective_user()
    if not email:
        return set() if customer else {}
    if perfil_do(email) in PERFIS_INTERNOS:
        return None
    rows = q("select customer, abas from cockpit.usuario_clientes where lower(email)=%s",
             (email.lower(),))
    mapa = {}
    for r in rows:
        lista = ABAS_PADRAO_CLIENTE if r["abas"] is None else list(r["abas"])
        mapa[r["customer"]] = {a for a in lista if a in ABAS_CLIENTE_OK}
    return mapa if customer is None else mapa.get(customer, set())


def deny_aba(customer, *abas):
    """Barra o acesso a uma aba NESTE cliente. Passe mais de uma quando a rota
    serve duas telas — /api/monitmov alimenta Cadastros E Cobertura."""
    if (d := deny_customer(customer)):
        return d
    lib = abas_liberadas(customer)
    if lib is None or (set(abas) & lib):
        return None
    return _err(403, "Esta aba não está liberada para você neste cliente.")


def require_interno():
    """Portão das ações de OPERAÇÃO: importar medição, script SQL, ambiente,
    coletar, cenários, publicar painel, sincronizar, decidir GAP. O usuário do
    cliente consulta e exporta — não opera."""
    if (r := require_auth()):
        return r
    if not eh_interno():
        return _err(403, "Ação restrita à equipe TOTVS.")
    return None


def hash_scrypt(senha):
    """Gera no ÚNICO formato que o ecossistema entende: scrypt$<saltHex>$<hashHex>,
    com o salt usado como STRING (o hex em UTF-8), N=16384 r=8 p=1 dklen=64 —
    igual ao scryptSync do Node no cockpit-unico-tsc. Gravar bcrypt aqui
    derruba o login de TODOS os dashboards: eles leem a mesma tabela."""
    salt_hex = secrets.token_hex(16)   # 32 hex -> hash de 168 chars, igual ao do cockpit
    dk = hashlib.scrypt(senha.encode(), salt=salt_hex.encode(), n=16384, r=8, p=1,
                        dklen=64, maxmem=132 * 1024 * 1024)
    return f"scrypt${salt_hex}${dk.hex()}"


def senha_provisoria():
    """Legível ao telefone: sem 0/O/1/l/I."""
    alfabeto = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnpqrstuvwxyz23456789"
    return "".join(secrets.choice(alfabeto) for _ in range(12))


_RE_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+[.][A-Za-z]{2,}")


def emails_do_texto(v):
    """Aceita lista OU um blocão colado (vírgula, ponto-e-vírgula, quebra de
    linha, 'Nome <email>'). É isso que faz 'adicionar e-mails' ser prático:
    o admin cola o que veio do cliente e a tela entende."""
    txt = v if isinstance(v, str) else " ".join(str(x) for x in (v or []))
    vistos, out = set(), []
    for m in _RE_EMAIL.findall(txt):
        e = m.strip().lower()
        if e not in vistos:
            vistos.add(e)
            out.append(e)
    return out


# ── OAuth Tasks SC (a API PCI usa o mesmo) ──────────────────────────────────
_tok = {"t": None, "exp": 0}


def _token(force=False):
    now = time.time()
    if not force and _tok["t"] and _tok["exp"] - 120 > now:
        return _tok["t"]
    r = requests.post(f"{TASKS_BASE}/api/oauth2/v1/token",
                      data={"grant_type": "password", "username": TASKS_USER,
                            "password": TASKS_PASS},
                      headers={"Content-Type": "application/x-www-form-urlencoded"},
                      timeout=30)
    if r.status_code not in (200, 201):
        raise RuntimeError(f"OAuth falhou HTTP {r.status_code}: {r.text[:200]}")
    d = r.json()
    _tok["t"] = d["access_token"]
    _tok["exp"] = now + int(d.get("expires_in", 3600))
    return _tok["t"]


def pci_get(url, params):
    """GET na API PCI com retry (401/5xx) e o encoding mentiroso (cp1252)."""
    ultimo = None
    for i in range(3):
        try:
            r = requests.get(url, params=params, timeout=(8, 40),
                             headers={"Authorization": f"Bearer {_token(force=i >= 1)}",
                                      "Accept": "application/json"})
        except (requests.exceptions.ConnectTimeout, requests.exceptions.ConnectionError,
                requests.exceptions.ReadTimeout) as ex:
            ultimo = ex
            time.sleep(1.5 * (i + 1))
            continue
        if r.status_code == 401 and i < 2:
            continue
        # 500 incluído: a API da TOTVS devolve 500 em soluços transitórios.
        if r.status_code in (500, 502, 503, 504) and i < 2:
            ultimo = f"HTTP {r.status_code}"
            time.sleep(1.5 * (i + 1))
            continue
        r.raise_for_status()
        txt = r.content.decode("utf-8", errors="replace")
        if "�" not in txt:
            return json.loads(txt)
        try:
            return json.loads(r.content.decode("cp1252"))
        except Exception:
            return json.loads(txt)
    raise PCIUnavailable(f"API Totvs SC indisponível após 3 tentativas — {ultimo}. "
                         f"Instabilidade temporária; tente novamente em instantes.")


# ── Páginas ─────────────────────────────────────────────────────────────────
def serve(name, ctype="text/html; charset=utf-8"):
    f = WEB_DIR / name
    if not f.exists():
        return _err(404, f"{name} não encontrado.")
    return Response(f.read_bytes(), mimetype=ctype)


@app.get("/login")
def page_login():
    return serve("login.html")


@app.get("/")
def page_root():
    return serve("index.html") if current_user() else redirect("/login", 302)


@app.get("/index.html")
def page_index():
    return serve("index.html") if current_user() else redirect("/login", 302)


# ── Auth API ────────────────────────────────────────────────────────────────
@app.post("/api/login")
def api_login():
    b = request.get_json(silent=True) or {}
    email = (b.get("email") or "").strip().lower()
    senha = b.get("senha") or ""
    if not email or not senha:
        return _err(400, "Informe e-mail e senha.")
    row = q("select email, nome, senha_hash, ativo from cockpit.usuarios_login "
            "where lower(email)=%s", (email,), one=True)
    if not row or not row["ativo"]:
        return _err(401, "Usuário não autorizado.")
    if not verify_scrypt(row["senha_hash"], senha):
        return _err(401, "Credenciais inválidas.")
    resp = make_response(_json({"ok": True, "email": row["email"], "nome": row["nome"]}))
    resp.set_cookie(COOKIE_NAME, make_session(row["email"], row["nome"] or ""),
                    max_age=SESSION_TTL, httponly=True, secure=True, samesite="Lax", path="/")
    return resp


@app.post("/api/logout")
def api_logout():
    r = make_response(_json({"ok": True}))
    r.set_cookie(COOKIE_NAME, "", max_age=0, path="/")
    return r


@app.get("/api/me")
def api_me():
    s = read_session()
    if not s:
        return _err(401, "Não autenticado.")
    real = s["e"]
    adm = is_admin(real)
    alvo = s.get("v") if (s.get("v") and adm) else None
    efetivo = alvo or real
    perfil = perfil_do(efetivo)
    lib = abas_liberadas()                 # None = interno (todas as abas)
    return _json({"ok": True, "email": real, "nome": s.get("n"), "is_admin": adm,
                  "view_as": alvo, "efetivo": efetivo,
                  "perfil": perfil,
                  "interno": perfil in PERFIS_INTERNOS,
                  # modo 'cliente' = ÁREA DE CLIENTES: mesma tela, sem as ações
                  # de operação e só com as abas liberadas.
                  "modo": "interno" if perfil in PERFIS_INTERNOS else "cliente",
                  "abas_catalogo": ABAS,
                  "abas": None if lib is None else {c: sorted(a) for c, a in lib.items()}})


@app.get("/api/usuarios")
def api_usuarios():
    """Lista de usuários para o seletor 'ver como' — só admin."""
    if (r := require_admin()):
        return r
    rows = q("""select u.email, u.nome, coalesce(u.is_admin,false) as is_admin, u.ativo,
                       (select count(*) from cockpit.usuario_clientes uc
                         where lower(uc.email)=lower(u.email)) as n_clientes
                from cockpit.usuarios_login u order by u.nome""")
    return _json({"ok": True, "usuarios": rows})


@app.post("/api/view-as")
def api_view_as():
    """Liga/desliga a simulação. Só admin. Body: {email} ou {email: null} p/ sair."""
    if (r := require_admin()):
        return r
    body = request.get_json(silent=True) or {}
    alvo = (body.get("email") or "").strip().lower() or None
    if alvo:
        u = q("select email from cockpit.usuarios_login where lower(email)=%s",
              (alvo,), one=True)
        if not u:
            return _err(404, "Usuário não encontrado.")
        alvo = u["email"]
    s = read_session()
    resp = make_response(_json({"ok": True, "view_as": alvo}))
    resp.set_cookie(COOKIE_NAME, make_session(s["e"], s.get("n"), view_as=alvo),
                    max_age=SESSION_TTL, httponly=True, secure=True,
                    samesite="Lax", path="/")
    return resp


# ── ACESSOS — e-mails x cliente x abas (tela "Acessos", só admin) ───────────
# A liberação vive em cockpit.usuario_clientes: (email, customer) já existia e
# agora carrega a coluna abas text[]. NULL = herda (interno vê tudo; cliente
# recebe ABAS_PADRAO_CLIENTE). '{}' = liberado no cliente e sem nenhuma aba.
@app.get("/api/acessos")
def api_acessos():
    if (r := require_admin()):
        return r
    usuarios = q("""select u.email, u.nome, u.perfil, coalesce(u.is_admin,false) as is_admin,
                           u.ativo, u.last_login
                      from cockpit.usuarios_login u
                     order by coalesce(u.nome, u.email)""")
    liberacoes = q("""select uc.email, uc.customer, uc.abas, uc.papel,
                             c.nome as cliente_nome
                        from cockpit.usuario_clientes uc
                        left join cockpit.clientes c on c.customer = uc.customer
                       order by uc.customer, uc.email""")
    clientes = q("select customer, nome from cockpit.clientes order by nome")
    # Papel e módulos vêm JUNTO das abas porque é a mesma pergunta do admin:
    # "o que essa pessoa faz neste cliente?". Ter que abrir o modal do protótipo
    # só para dizer que fulano valida Faturamento era o atrito que sobrava.
    return _json({"ok": True, "usuarios": usuarios, "liberacoes": liberacoes,
                  "clientes": clientes, "abas": ABAS, "perfis": list(PERFIS),
                  "abas_padrao_cliente": ABAS_PADRAO_CLIENTE,
                  "papeis": [{"id": p, "label": PROTO_PAPEL_LABEL[p],
                              "totvs": p in PROTO_PAPEIS_TOTVS} for p in PROTO_PAPEIS],
                  "modulos_cliente": modulos_por_cliente(),
                  "modulos_liberados": q("""select customer, email, modulo, papel
                                              from cockpit.proto_usuario_modulos
                                             order by customer, email, papel, modulo""")})


@app.post("/api/acessos")
def api_acessos_salvar():
    """Libera N e-mails em N clientes com as MESMAS abas, numa tacada.

    Cria o login de quem ainda não existe com uma senha provisória, devolvida
    UMA ÚNICA VEZ nesta resposta — não fica gravada em lugar nenhum, nem em log.

    Body: {emails: "colado ou lista", customers: [...], abas: [...],
           perfil: "cliente", criar_login: true,
           papel: null|"consultor"|…, papel_modulo: "executa"|"valida",
           modulos: ["Faturamento", …] ou ["*"]}
    """
    if (r := require_admin()):
        return r
    b = request.get_json(silent=True) or {}
    emails = emails_do_texto(b.get("emails"))
    customers = [str(c).strip() for c in (b.get("customers") or []) if str(c).strip()]
    abas = [a for a in (b.get("abas") or []) if a in ABAS_IDS]
    perfil = (b.get("perfil") or "cliente").strip().lower()
    criar = b.get("criar_login") is not False

    if not emails:
        return _err(400, "Nenhum e-mail válido no que foi colado.")
    if not customers:
        return _err(400, "Escolha pelo menos um cliente.")
    if perfil not in PERFIS:
        return _err(400, f"Perfil desconhecido: {perfil}")
    if perfil == "admin":
        return _err(400, "Perfil admin não se concede por aqui — admin vê todos os clientes.")
    # Aba fora do catálogo do cliente vira acesso indevido silencioso.
    if perfil == "cliente" and (fora := [a for a in abas if a not in ABAS_CLIENTE_OK]):
        return _err(400, "Abas não liberáveis para o perfil cliente: " + ", ".join(fora))
    # Cliente inexistente em cockpit.clientes = liberação que não aparece na tela.
    conhecidos = {r["customer"] for r in q("select customer from cockpit.clientes")}
    if (nc := [c for c in customers if c not in conhecidos]):
        return _err(409, "Cliente não cadastrado em cockpit.clientes: " + ", ".join(nc))

    criados, existentes = [], []
    for e in emails:
        row = q("select email, perfil, ativo from cockpit.usuarios_login "
                "where lower(email)=%s", (e,), one=True)
        if row:
            existentes.append({"email": row["email"], "perfil": row["perfil"],
                               "ativo": row["ativo"]})
            continue
        if not criar:
            return _err(409, f"{e} não tem login e 'criar login' está desligado.")
        pw = senha_provisoria()
        execute("""insert into cockpit.usuarios_login (email, senha_hash, nome, ativo, perfil, created_by)
                   values (%s,%s,%s,true,%s,%s)""",
                (e, hash_scrypt(pw), e.split("@")[0].replace(".", " ").title(),
                 perfil, current_user()))
        criados.append({"email": e, "senha_provisoria": pw})

    # Perfil dos que já existiam: só REBAIXA para 'cliente' se quem chamou pediu
    # explicitamente. Promover alguém por engano numa tela de liberação seria
    # exatamente o tipo de acidente que esta tela existe para evitar.
    if b.get("aplicar_perfil"):
        for e in [x["email"] for x in existentes]:
            execute("update cockpit.usuarios_login set perfil=%s, updated_at=now() "
                    "where lower(email)=%s and coalesce(is_admin,false)=false",
                    (perfil, e.lower()))

    # A FK de usuario_clientes aponta para usuarios_login(email) — casing EXATO.
    # Inserir o e-mail em minúsculo quando o login foi criado "Fulano@x.com"
    # estoura a FK. Sempre resolver o e-mail canônico do banco.
    canon = {r["email"].lower(): r["email"] for r in
             q("select email from cockpit.usuarios_login where lower(email) = any(%s)",
               (emails,))}
    if (sem := [e for e in emails if e not in canon]):
        return _err(500, "Login não encontrado após a criação: " + ", ".join(sem))
    pares = [(canon[e], c, abas) for e in emails for c in customers]
    with db() as conn, conn.cursor() as cur:
        execute_values(cur,
            """insert into cockpit.usuario_clientes (email, customer, abas, created_by)
               values %s
               on conflict (email, customer) do update set abas = excluded.abas""",
            [(e, c, a, current_user()) for e, c, a in pares])
    # ── Papel e módulos NESTE cliente ──────────────────────────────────────
    # É a mesma tabela que o modal "Equipe do protótipo" grava. Está aqui
    # porque liberar o consultor e dizer o que ele valida é UMA decisão só:
    # separar em duas telas era o que fazia sobrar consultor liberado no
    # cliente e sem módulo nenhum — e ele descobria isso ao ser barrado.
    alvos = [canon[e] for e in emails]
    papel = (b.get("papel") or "").strip() or None
    papel_mod = (b.get("papel_modulo") or "").strip() or None
    modulos = [str(m).strip() for m in (b.get("modulos") or []) if str(m).strip()]
    ignorados = []

    if papel:
        if papel not in PROTO_PAPEIS:
            return _err(400, f"Papel desconhecido: {papel}")
        # Papel de coordenação/validação é da TOTVS. Confere DEPOIS de criar o
        # login e de aplicar o perfil — senão recusaria o consultor que acabou
        # de ser criado com o perfil certo nesta mesma chamada.
        if papel in PROTO_PAPEIS_TOTVS:
            externos = [e for e in alvos if perfil_do(e) not in PERFIS_INTERNOS]
            if externos:
                return _err(409, "Papel da TOTVS não pode ser dado a login de cliente: "
                                 + ", ".join(externos))
        execute("""update cockpit.usuario_clientes set papel=%s
                    where customer=any(%s) and lower(email)=any(%s)""",
                (papel, customers, emails))

    if modulos:
        if papel_mod not in ("executa", "valida"):
            return _err(400, "Para liberar módulos, diga se a pessoa EXECUTA ou VALIDA.")
        # Módulo que não existe naquele cliente vira liberação fantasma: fica
        # gravada, não aparece em lugar nenhum e engana na próxima conferência.
        # '*' é o coringa "todos" e passa sempre.
        catalogo = {c: set(modulos_do_cliente(c)) for c in customers}
        linhas = []
        for c in customers:
            for m in modulos:
                if m == "*" or m in catalogo[c]:
                    linhas += [(c, e, m, papel_mod, None, current_user()) for e in alvos]
                else:
                    ignorados.append(f"{c} · {m}")
        # SUBSTITUI o conjunto daquele papel nos pares marcados: o que a tela
        # mostra desmarcado tem que sair, senão desmarcar não faria nada.
        execute("""delete from cockpit.proto_usuario_modulos
                    where customer=any(%s) and lower(email)=any(%s) and papel=%s""",
                (customers, emails, papel_mod))
        if linhas:
            with db() as conn, conn.cursor() as cur:
                execute_values(cur,
                    """insert into cockpit.proto_usuario_modulos
                         (customer, email, modulo, papel, processo, created_by) values %s
                       on conflict do nothing""", linhas)

    return _json({"ok": True, "emails": emails, "customers": customers, "abas": abas,
                  "criados": criados, "existentes": existentes,
                  "liberacoes": len(pares), "papel": papel,
                  "modulos": modulos if modulos else [], "papel_modulo": papel_mod,
                  "modulos_ignorados": sorted(set(ignorados))})


@app.delete("/api/acessos")
def api_acessos_remover():
    """Tira a liberação de um par (email, customer). Não apaga o login."""
    if (r := require_admin()):
        return r
    email = (request.args.get("email") or "").strip().lower()
    customer = (request.args.get("customer") or "").strip()
    if not email or not customer:
        return _err(400, "Informe email e customer.")
    execute("delete from cockpit.usuario_clientes where lower(email)=%s and customer=%s",
            (email, customer))
    return _json({"ok": True})


@app.post("/api/acessos/senha")
def api_acessos_senha():
    """Nova senha provisória. Devolvida uma única vez — não fica gravada."""
    if (r := require_admin()):
        return r
    b = request.get_json(silent=True) or {}
    email = (b.get("email") or "").strip().lower()
    if not q("select 1 from cockpit.usuarios_login where lower(email)=%s", (email,), one=True):
        return _err(404, "Usuário não encontrado.")
    pw = senha_provisoria()
    execute("update cockpit.usuarios_login set senha_hash=%s, updated_at=now() "
            "where lower(email)=%s", (hash_scrypt(pw), email))
    return _json({"ok": True, "email": email, "senha_provisoria": pw})


@app.post("/api/acessos/ativo")
def api_acessos_ativo():
    if (r := require_admin()):
        return r
    b = request.get_json(silent=True) or {}
    email = (b.get("email") or "").strip().lower()
    ativo = bool(b.get("ativo"))
    if email == (current_user() or "").lower() and not ativo:
        return _err(409, "Você não pode se desativar.")
    execute("update cockpit.usuarios_login set ativo=%s, updated_at=now() "
            "where lower(email)=%s", (ativo, email))
    return _json({"ok": True, "email": email, "ativo": ativo})


@app.get("/api/health")
def api_health():
    info = {"ok": True, "service": "projetos-vercel",
            "recorte": {"regioes": sorted(SYNC_REGIOES), "purga": SYNC_PURGE},
            "env": {"DATABASE_URL": bool(DATABASE_URL), "TASKS_USERNAME": bool(TASKS_USER),
                    "TASKS_PASSWORD": bool(TASKS_PASS), "SESSION_SECRET": SESSION_SECRET != "dev-insecure-secret"},
            "db": False}
    try:
        r = q("select 1 ok", one=True)
        info["db"] = bool(r and r.get("ok") == 1)
        c = q("select count(*) n, max(synced_at) s from cockpit.projetos", one=True)
        info["projetos"] = c["n"]
        info["ultimo_sync"] = c["s"]
    except Exception as e:
        info["ok"] = False
        info["db_error"] = f"{type(e).__name__}: {e}"
    return _json(info)


# ── Raio X · maturidade da equipe (schema pdi) ──────────────────────────────
# Matriz consultor × módulo com escala 0..5 (0 não conhece → 5 especialista).
# Toda escrita passa pelas funções pdi.fn_*, que validam a escala e deixam a
# trigger registrar o histórico.
#
# QUEM VÊ O QUÊ — a regra toda mora aqui, não no front:
#   perfil 'consultoria'  → só a PRÓPRIA linha, e só edita a própria
#                           autoavaliação. Nunca vê a nota que o gestor deu a
#                           outra pessoa, nem o recorte da célula.
#   admin / comum / leitor → a célula inteira (é o CP e o time fixo).
#   cliente                → nada: require_interno() barra antes.
# O vínculo login→consultor é o e-mail em pdi.consultores.email, preenchido na
# sub-aba Cadastro. SEM vínculo, um 'consultoria' não vê linha nenhuma — falhar
# fechado é melhor que mostrar a pessoa errada.
RX_UUID = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-"
                     r"[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")


def _rx_uuid(v):
    v = (str(v or "")).strip()
    return v if RX_UUID.match(v) else None


def _rx_nivel(v):
    """None limpa a célula; fora de 0..5 é erro, não silêncio."""
    if v in (None, ""):
        return None, None
    try:
        n = int(v)
    except (TypeError, ValueError):
        return None, "nível precisa ser um número de 0 a 5."
    if not 0 <= n <= 5:
        return None, "nível fora da escala 0..5."
    return n, None


def _rx_par(b):
    cons, mod = _rx_uuid(b.get("consultor")), _rx_uuid(b.get("modulo"))
    if not cons or not mod:
        return None, None, _err(400, "consultor e módulo precisam ser UUID.")
    return cons, mod, None


def _rx_ctx():
    """(email, perfil, so_eu, meu_id). so_eu=True ⇒ enxerga uma linha só."""
    email = effective_user() or ""
    perfil = perfil_do(email)
    so_eu = perfil == "consultoria"
    meu = q("select id from pdi.consultores where ativo and lower(email)=%s",
            (email.lower(),), one=True) if email else None
    return email, perfil, so_eu, (meu["id"] if meu else None)


def _rx_gestor():
    """Portão das ações de gestão: nota do gestor, meta, cadastro, frente,
    vínculo de e-mail e arquivamento. Consultor não mexe em nada disso."""
    if (r := require_interno()):
        return r
    if perfil_do(effective_user() or "") == "consultoria":
        return _err(403, "Ação restrita à coordenação.")
    return None


@app.get("/api/raiox")
def api_raiox():
    if (r := require_interno()):
        return r
    email, perfil, so_eu, meu_id = _rx_ctx()
    if so_eu and not meu_id:
        return _json({"sem_vinculo": True, "eu": {"email": email, "perfil": perfil,
                      "so_eu": True, "consultor_id": None, "pode_gerir": False},
                      "niveis": [], "grupos": [], "consultores": [], "modulos": [],
                      "avaliacoes": [], "historico": []})
    if so_eu:
        row = q("select pdi.fn_snapshot_consultor(%s::uuid) as s", (meu_id,), one=True)
    else:
        row = q("select pdi.fn_snapshot() as s", one=True)
    s = (row or {}).get("s") or {}
    s["eu"] = {"email": email, "perfil": perfil, "so_eu": so_eu,
               "consultor_id": meu_id, "pode_gerir": not so_eu}
    return _json(s)


@app.post("/api/raiox/nivel")
def api_raiox_nivel():
    if (r := require_interno()):
        return r
    b = request.get_json(silent=True) or {}
    campo = (b.get("campo") or "").strip()
    if campo not in ("auto", "gestor", "meta"):
        return _err(400, "campo deve ser auto, gestor ou meta.")
    nivel, erro = _rx_nivel(b.get("valor"))
    if erro:
        return _err(400, erro)
    cons, mod, falha = _rx_par(b)
    if falha:
        return falha

    _, _, so_eu, meu_id = _rx_ctx()
    if so_eu:
        # O consultor preenche a própria autoavaliação — e só isso.
        if campo != "auto":
            return _err(403, "Você só pode preencher a sua autoavaliação.")
        if not meu_id or cons != meu_id:
            return _err(403, "Você só pode avaliar a sua própria linha.")

    row = q("select pdi.fn_set_nivel(%s::uuid, %s::uuid, %s, %s::smallint) as r",
            (cons, mod, campo, nivel), one=True)
    return _json(row["r"] if row else {"ok": True})


@app.post("/api/raiox/detalhe")
def api_raiox_detalhe():
    if (r := _rx_gestor()):
        return r
    b = request.get_json(silent=True) or {}
    cons, mod, falha = _rx_par(b)
    if falha:
        return falha
    prazo = (b.get("prazo") or "").strip() or None
    row = q("select pdi.fn_set_detalhe(%s::uuid, %s::uuid, %s, %s::date) as r",
            (cons, mod, (b.get("obs") or "").strip(), prazo), one=True)
    return _json(row["r"] if row else {"ok": True})


@app.post("/api/raiox/frente")
def api_raiox_frente():
    if (r := _rx_gestor()):
        return r
    b = request.get_json(silent=True) or {}
    cons = _rx_uuid(b.get("consultor"))
    frente = (b.get("frente") or "").strip()
    if not cons:
        return _err(400, "consultor precisa ser UUID.")
    if not frente:
        return _err(400, "a frente não pode ficar vazia.")
    row = q("select pdi.fn_set_frente(%s::uuid, %s) as r", (cons, frente), one=True)
    return _json(row["r"] if row else {"ok": True})


@app.post("/api/raiox/email")
def api_raiox_email():
    """Vincula o login do consultor à linha dele. É isto que faz o recorte
    'só a minha linha' funcionar — sem e-mail, o consultor não vê nada."""
    if (r := _rx_gestor()):
        return r
    b = request.get_json(silent=True) or {}
    cons = _rx_uuid(b.get("consultor"))
    if not cons:
        return _err(400, "consultor precisa ser UUID.")
    row = q("select pdi.fn_set_email(%s::uuid, %s) as r",
            (cons, (b.get("email") or "").strip()), one=True)
    return _json(row["r"] if row else {"ok": True})


@app.post("/api/raiox/modulo")
def api_raiox_modulo():
    if (r := _rx_gestor()):
        return r
    b = request.get_json(silent=True) or {}
    codigo = re.sub(r"\s+", "_", (b.get("codigo") or "").strip()).upper()
    nome = (b.get("nome") or "").strip()
    grupo = (b.get("grupo") or "").strip()
    if not codigo or not nome or not grupo:
        return _err(400, "código, nome e grupo são obrigatórios.")
    row = q("select pdi.fn_upsert_modulo(%s, %s, %s, %s, %s) as r",
            (codigo, nome, (b.get("sigla") or "").strip(), grupo,
             bool(b.get("critico"))), one=True)
    return _json(row["r"] if row else {"ok": True})


@app.post("/api/raiox/consultor")
def api_raiox_consultor():
    if (r := _rx_gestor()):
        return r
    b = request.get_json(silent=True) or {}
    codigo = (b.get("codigo") or "").strip().upper()
    nome = (b.get("nome") or "").strip()
    if not codigo or not nome:
        return _err(400, "código e nome são obrigatórios.")
    row = q("select pdi.fn_upsert_consultor(%s, %s, %s, %s, %s) as r",
            (codigo, nome, (b.get("funcao") or "").strip(),
             (b.get("celula") or "201").strip(),
             (b.get("frente") or "Funcional").strip() or "Funcional"), one=True)
    return _json(row["r"] if row else {"ok": True})


@app.post("/api/raiox/arquivar")
def api_raiox_arquivar():
    if (r := _rx_gestor()):
        return r
    b = request.get_json(silent=True) or {}
    tipo = (b.get("tipo") or "").strip()
    alvo = _rx_uuid(b.get("id"))
    if tipo not in ("modulo", "consultor"):
        return _err(400, "tipo deve ser modulo ou consultor.")
    if not alvo:
        return _err(400, "id precisa ser UUID.")
    row = q("select pdi.fn_arquivar(%s, %s::uuid) as r", (tipo, alvo), one=True)
    return _json(row["r"] if row else {"ok": True})


# ── Dados ───────────────────────────────────────────────────────────────────
@app.get("/api/clientes")
def api_clientes():
    if (r := require_auth()):
        return r
    allowed = allowed_customers()
    rows = q("select customer, nome from cockpit.clientes order by nome")
    out = [{"codigo": r["customer"], "nome": r["nome"], "chave": _slug(r["nome"])}
           for r in rows if allowed is None or r["customer"] in allowed]
    return _json({"ok": True, "clientes": out})


@app.get("/api/projetos")
def api_projetos():
    """Lista da tabela sincronizada, já recortada pelos clientes liberados."""
    if (r := require_auth()):
        return r
    allowed = allowed_customers()
    rows = q("select * from cockpit.projetos order by nome_cliente_projeto, codigo_projeto")
    nomes = {r["customer"]: r["nome"] for r in q("select customer, nome from cockpit.clientes")}
    projetos, sync = [], None
    for r in rows:
        cli = r["codigo_cliente_projeto"]
        if allowed is not None and cli not in allowed:
            continue
        sync = sync or r["synced_at"]
        p = dict(r.get("raw") or {})
        p.update({
            "codigo_projeto": r["codigo_projeto"],
            "codigo_cliente_projeto": cli,
            "nome_cliente_projeto": r["nome_cliente_projeto"],
            "descricao_projeto": r["descricao_projeto"],
            "nome_coordenador_projeto": r["nome_coordenador_projeto"],
            "status_projeto": r["status_projeto"],
            "tipo_projeto": r["tipo_projeto"],
            "versao_projeto": r["versao_projeto"],
            "_nome_cliente_local": nomes.get(cli) or r["nome_cliente_projeto"],
        })
        projetos.append(p)
    return _json({"ok": True, "total": len(projetos), "projetos": projetos,
                  "sincronizado_em": sync})


@app.post("/api/sync")
def api_sync():
    """Sincroniza UMA página da API PCI para cockpit.projetos.
    O front chama page=1,2,3… enquanto hasNext for true (evita timeout)."""
    if (r := require_interno()):
        return r
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de sincronizar.")
    page = int(request.args.get("page", 1))
    size = int(request.args.get("pageSize", 200))
    data = pci_get(LISTA_URL, {"page": page, "pageSize": size})
    if not isinstance(data, dict):
        return _err(502, "Resposta inesperada da API PCI.")
    items, ignorados = filtra_recorte(data.get("items") or [])
    for p in items:
        cod = p.get("codigo_projeto")
        if not cod:
            continue
        execute("""
            insert into cockpit.projetos (codigo_projeto, codigo_cliente_projeto,
              nome_cliente_projeto, descricao_projeto, nome_coordenador_projeto,
              status_projeto, tipo_projeto, versao_projeto, raw, synced_at)
            values (%s,%s,%s,%s,%s,%s,%s,%s,%s,now())
            on conflict (codigo_projeto) do update set
              codigo_cliente_projeto=excluded.codigo_cliente_projeto,
              nome_cliente_projeto=excluded.nome_cliente_projeto,
              descricao_projeto=excluded.descricao_projeto,
              nome_coordenador_projeto=excluded.nome_coordenador_projeto,
              status_projeto=excluded.status_projeto,
              tipo_projeto=excluded.tipo_projeto,
              versao_projeto=excluded.versao_projeto,
              raw=excluded.raw, synced_at=now()
        """, (cod, p.get("codigo_cliente_projeto"), p.get("nome_cliente_projeto"),
              p.get("descricao_projeto"), p.get("nome_coordenador_projeto"),
              p.get("status_projeto"), p.get("tipo_projeto"), p.get("versao_projeto"),
              json.dumps(p)))
    return _json({"ok": True, "page": page, "gravados": len(items),
                  "ignorados": ignorados, "hasNext": bool(data.get("hasNext"))})


# ── Recorte: quem entra na base ─────────────────────────────────────────────
_CLIENTES_CACHE = {"t": 0.0, "v": frozenset()}


def clientes_cockpit(force=False):
    """customers liberados em cockpit.clientes — cache de 5 min.
    Se o banco falhar, devolve o último valor conhecido (nunca vazio por erro:
    conjunto vazio + regiões vazias apagaria a base inteira na purga)."""
    if not force and _CLIENTES_CACHE["v"] and time.time() - _CLIENTES_CACHE["t"] < 300:
        return _CLIENTES_CACHE["v"]
    try:
        v = frozenset(str(r["customer"] or "").strip()
                      for r in q("select customer from cockpit.clientes"))
    except Exception:
        return _CLIENTES_CACHE["v"]
    if v:
        _CLIENTES_CACHE.update(t=time.time(), v=v)
    return v


def no_recorte(p, clientes=None):
    """True se o projeto é da regional (região do coordenador titular ou auxiliar)
    ou de um cliente do Cockpit."""
    if not SYNC_REGIOES:
        return True
    for k in ("regiao_coordenador_projeto", "regiao_coordenador_auxiliar"):
        if str(p.get(k) or "").strip() in SYNC_REGIOES:
            return True
    if str(p.get("regiao_cliente_projeto") or "").strip() in SYNC_REGIOES \
            and str(p.get("status_projeto") or "").strip().lower() not in SYNC_ENCERRADOS:
        return True
    cli = str(p.get("codigo_cliente_projeto") or "").strip()
    return bool(cli) and cli in (clientes if clientes is not None else clientes_cockpit())


def filtra_recorte(items, clientes=None):
    if not SYNC_REGIOES:
        return list(items), 0
    clientes = clientes_cockpit() if clientes is None else clientes
    dentro = [p for p in items if no_recorte(p, clientes)]
    return dentro, len(items) - len(dentro)


def purga_fora_do_recorte():
    """Apaga de cockpit.projetos o que não casa mais com o recorte. Roda no fim do
    cron; a regra SQL é a MESMA do no_recorte() — mexeu numa, mexa na outra."""
    if not (SYNC_PURGE and SYNC_REGIOES):
        return 0
    regs = list(SYNC_REGIOES)
    with db() as c, c.cursor() as cur:
        cur.execute("""
            delete from cockpit.projetos
             where coalesce(btrim(raw->>'regiao_coordenador_projeto'), '') <> all(%s)
               and coalesce(btrim(raw->>'regiao_coordenador_auxiliar'), '') <> all(%s)
               and not (coalesce(btrim(raw->>'regiao_cliente_projeto'), '') = any(%s)
                        and lower(coalesce(btrim(status_projeto), '')) <> all(%s))
               and coalesce(btrim(codigo_cliente_projeto), '') not in (
                     select coalesce(btrim(customer), '') from cockpit.clientes)
        """, (regs, regs, regs, list(SYNC_ENCERRADOS)))
        return cur.rowcount or 0


# ── Sync completo para o CRON (pg_cron + pg_net) ────────────────────────────
def _upsert_projetos(items):
    """Upsert em LOTE (execute_values) — muito mais rápido que 1 insert por linha,
    o que é o que permite sincronizar as ~9 páginas dentro dos 60s da função."""
    rows = [(p.get("codigo_projeto"), p.get("codigo_cliente_projeto"),
             p.get("nome_cliente_projeto"), p.get("descricao_projeto"),
             p.get("nome_coordenador_projeto"), p.get("status_projeto"),
             p.get("tipo_projeto"), p.get("versao_projeto"), json.dumps(p))
            for p in items if p.get("codigo_projeto")]
    if not rows:
        return 0
    with db() as c, c.cursor() as cur:
        execute_values(cur, """
            insert into cockpit.projetos (codigo_projeto, codigo_cliente_projeto,
              nome_cliente_projeto, descricao_projeto, nome_coordenador_projeto,
              status_projeto, tipo_projeto, versao_projeto, raw, synced_at)
            values %s
            on conflict (codigo_projeto) do update set
              codigo_cliente_projeto=excluded.codigo_cliente_projeto,
              nome_cliente_projeto=excluded.nome_cliente_projeto,
              descricao_projeto=excluded.descricao_projeto,
              nome_coordenador_projeto=excluded.nome_coordenador_projeto,
              status_projeto=excluded.status_projeto,
              tipo_projeto=excluded.tipo_projeto,
              versao_projeto=excluded.versao_projeto,
              raw=excluded.raw, synced_at=now()
        """, rows, template="(%s,%s,%s,%s,%s,%s,%s,%s,%s,now())", page_size=200)
    return len(rows)


def _cron_autorizado():
    """O cron não tem sessão: autentica por Authorization: Bearer <CRON_SECRET>.
    Aceita também ?secret= para facilitar teste manual."""
    if not CRON_SECRET:
        return False
    auth = request.headers.get("Authorization", "")
    tok = auth[7:].strip() if auth.lower().startswith("bearer ") else ""
    return hmac.compare_digest(tok or request.args.get("secret", ""), CRON_SECRET)


@app.route("/api/cron/sync", methods=["GET", "POST"])
def api_cron_sync():
    """Sincroniza TODAS as páginas da API PCI numa tacada. Disparado de hora em
    hora pelo pg_cron do Supabase (via pg_net). Não exige login — exige o Bearer."""
    if not _cron_autorizado():
        return _err(401, "CRON_SECRET inválido ou ausente.")
    ini = time.time()
    page, total, paginas, ignorados = 1, 0, 0, 0
    clientes = clientes_cockpit(force=True)
    while page <= 50:
        data = pci_get(LISTA_URL, {"page": page, "pageSize": 200})
        if not isinstance(data, dict):
            break
        items, fora = filtra_recorte(data.get("items") or [], clientes)
        ignorados += fora
        total += _upsert_projetos(items)
        paginas += 1
        if not data.get("hasNext"):
            break
        page += 1
    apagados = purga_fora_do_recorte()
    dur = int((time.time() - ini) * 1000)
    try:
        execute("""insert into cockpit.sync_log
                   (source, status, started_at, finished_at, duration_ms,
                    tickets_processed, tickets_upserted)
                   values ('projetos-vercel','success', to_timestamp(%s), now(), %s, %s, %s)""",
                (ini, dur, total, total))
    except Exception:
        pass
    return _json({"ok": True, "paginas": paginas, "projetos": total,
                  "ignorados": ignorados, "apagados": apagados,
                  "regioes": sorted(SYNC_REGIOES), "duration_ms": dur})


# ── MONITCAD — status dos cadastros (aba "Cadastros") ───────────────────────
def _data_iso(v):
    """Aceita '20260427' (AAAAMMDD) ou '2026-04-27'; devolve ISO ou None."""
    s = str(v or "").strip()
    if re.fullmatch(r"\d{8}", s):
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return s if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s) else None


def _dt_hora(v):
    """('11/08/2026 11:43') -> ('2026-08-11', '11:43'). Aceita também ISO e AAAAMMDD."""
    s = str(v or "").strip()
    if not s:
        return None, None
    hora = None
    if " " in s:
        s, _, h = s.partition(" ")
        hora = h.strip()[:8] or None
    if m := re.fullmatch(r"(\d{2})[/-](\d{2})[/-](\d{4})", s):
        d, mes, a = m.groups()
        return f"{a}-{mes}-{d}", hora
    return _data_iso(s), hora


# Nomes de coluna aceitos no CSV (sem acento, maiúsculo, sem espaço/underscore)
# QTD_REAL / QTD_ESTIMADA estão aqui por um motivo concreto: o script SQL que a
# própria aba gera usa esses nomes, e em 18/08/2026 uma medição da Olim entrou
# com as 112 tabelas e realizado = 0 porque a coluna não era reconhecida. Ao
# acrescentar um alias novo, lembre que coluna ignorada vira ZERO silencioso.
CSV_COLS = {
    "MODULO": "modulo", "TABELA": "tabela", "DESCRICAO": "descricao",
    "QTDE": "realizado", "QTD": "realizado", "QUANTIDADE": "realizado",
    "REALIZADO": "realizado", "REGISTROS": "realizado", "CONTAGEM": "realizado",
    "QTDREAL": "realizado", "QTDATUAL": "realizado", "QTDREGISTROS": "realizado",
    "ESTIMATIVA": "estimativa", "META": "estimativa", "ESTIMADA": "estimativa",
    "PREVISTO": "estimativa", "QTDESTIMADA": "estimativa",
    "QTDESTIMATIVA": "estimativa", "QTDPREVISTA": "estimativa",
    "FILTRO": "filtro", "ETAPA": "etapa", "RESPONSAVEL": "responsavel",
    "STATUS": "status", "DATAPREV": "data_prev", "PREVISAO": "data_prev",
    "DTLEITURA": "_dt", "DATA": "_dt", "DATAMEDICAO": "_dt", "SEMANA": "_semana",
}

# Tabelas de CONFIGURAÇÃO do Protheus, não de cadastro do cliente: a contagem é
# do dicionário/parametrização que já vem com o produto (dezenas de milhares de
# linhas que ninguém "cadastra" no projeto), então elas poluíam a lista e nunca
# teriam estimativa. Retiradas da análise em 23/08/2026. Filtramos na importação
# para que um CSV gerado por um script antigo não as traga de volta.
# SX5 por GRUPO (SX5_S4, SX5_T3...) continua valendo — aquilo é cadastro de verdade.
TABELAS_IGNORADAS = {"SX5", "SX6", "SX7"}

# Documento e saldo NÃO são carga de cadastro: nota fiscal, pedido, título,
# ordem de produção, saldo de estoque e saldo contábil nascem da operação, não do
# trabalho de cadastramento do projeto. Somados junto, distorciam a contagem da
# aba Cadastros — e a pergunta que eles respondem ("a operação já rodou?") é a da
# aba Movimentos/Cobertura, que trabalha por cenário, não por contagem de tabela.
#
# Estas tabelas continuam sendo IMPORTADAS (o histórico não se perde); só entram
# com painel='movimento' e a aba Cadastros as filtra fora. Reversível: basta um
# update em cockpit.monitcad_tabelas.painel.
#
# Classificar aqui, e não só no script do Protheus, é o que impede a tabela de
# voltar quando alguém roda um script antigo — mesmo motivo do TABELAS_IGNORADAS.
TABELAS_MOVIMENTO = {
    "SC1", "SC7", "SC8",                       # compras: solicitação, pedido, cotação
    "SC5", "SC6",                              # pedido de venda
    "SF1", "SD1", "SF2", "SD2", "SF3",         # notas de entrada/saída e livros
    "SE1", "SE2", "SEB",                       # títulos e retorno bancário
    "SC2",                                     # ordem de produção
    "SB2", "SB8", "SB9", "SBF", "SBJ",         # saldos de estoque
    "SN3", "SN4",                              # saldos e movimentos do ativo
    "CV3",                                     # saldos contábeis
    "SL1", "SL2",                              # venda frente de loja
    "AB3", "AB4", "AB5",                       # orçamentos de serviço
    "AB6", "AB7", "AB8", "AB9", "ABC",         # ordens de serviço e apontamentos
    "CN9", "CNB",                              # contratos: cabeçalho e itens da planilha
}


def _painel_da_tabela(tabela):
    return "movimento" if (tabela or "").strip().upper() in TABELAS_MOVIMENTO else "cadastro"

# Só para a mensagem de erro — os nomes na forma em que o usuário escreve
QTD_ACEITOS = "QTDE, QTD, QTD_REAL, QUANTIDADE, REALIZADO, REGISTROS"


def _norm_col(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def _csv_para_body(texto):
    """Converte o CSV de status de cadastros no mesmo formato do
    historico-semanal.json. Uma medição por data encontrada em DT_LEITURA."""
    amostra = texto[:4096]
    try:
        dial = csv.Sniffer().sniff(amostra, delimiters=",;\t|")
    except csv.Error:
        dial = csv.excel
    linhas = list(csv.reader(io.StringIO(texto), dial))
    if not linhas:
        raise ValueError("CSV vazio.")
    cabec = [CSV_COLS.get(_norm_col(c)) for c in linhas[0]]
    if "tabela" not in cabec:
        raise ValueError("CSV sem a coluna TABELA. Esperado: MODULO, TABELA, "
                         "DESCRICAO, DT_LEITURA, SEMANA, QTDE.")
    # Sem coluna de quantidade a medição entraria inteira zerada e pareceria uma
    # base vazia. Falhar aqui é MUITO mais barato do que descobrir depois.
    if "realizado" not in cabec:
        raise ValueError(f"CSV sem coluna de quantidade — aceito: {QTD_ACEITOS}. "
                         "Cabeçalho recebido: "
                         + ", ".join(c.strip() for c in linhas[0] if c.strip()))

    por_data = {}
    for linha in linhas[1:]:
        if not any((c or "").strip() for c in linha):
            continue
        reg = {}
        for col, val in zip(cabec, linha):
            if col:
                reg[col] = (val or "").strip()
        if not reg.get("tabela"):
            continue
        dt, hora = _dt_hora(reg.pop("_dt", ""))
        if not dt:
            raise ValueError("CSV sem data válida em DT_LEITURA (ex.: 11/08/2026 11:43).")
        semana = reg.pop("_semana", "")
        med = por_data.setdefault(dt, {
            "data_iso": dt, "hora_medicao": hora,
            "semana": int(re.sub(r"^.*W", "", semana) or 0) or None,
            "tabelas": [],
        })
        for k in ("realizado", "estimativa"):
            if k in reg:
                reg[k] = float(str(reg[k]).replace(".", "").replace(",", ".") or 0)
        if "data_prev" in reg:
            reg["data_prev"] = _dt_hora(reg["data_prev"])[0]
        med["tabelas"].append(reg)

    if not por_data:
        raise ValueError("CSV sem linhas de tabela.")
    return {"medicoes": [por_data[d] for d in sorted(por_data)]}


AMBIENTES = ("producao", "teste")


def _ambiente():
    """Produção e teste são bases DIFERENTES: nunca somar as duas na mesma
    leitura. O ambiente vem sempre explícito do front."""
    a = (request.args.get("ambiente") or "producao").strip().lower()
    return a if a in AMBIENTES else "producao"


@app.get("/api/monitcad/<customer>")
def api_monitcad(customer):
    """Última medição + histórico de carga dos cadastros do cliente, no ambiente
    pedido. Sem medição naquele ambiente, devolve vazio:true."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "cadastros")):
        return d
    amb = _ambiente()
    proj = q("select * from cockpit.monitcad_projetos where customer=%s",
             (customer,), one=True)
    meds = q("""select id, data_medicao, semana, hora_medicao, origem
                  from cockpit.monitcad_medicoes
                 where customer=%s and ambiente=%s order by data_medicao""",
             (customer, amb))
    if not meds:
        return _json({"ok": True, "customer": customer, "projeto": proj, "ambiente": amb,
                      "vazio": True, "tabelas": [], "serie": [], "total_medicoes": 0})
    ult = meds[-1]
    # painel='cadastro': documento e saldo vivem na aba Movimentos (ver
    # TABELAS_MOVIMENTO). O filtro precisa estar em TODAS as consultas da aba —
    # deixar de fora a dos módulos ou a da série faria os totais brigarem com a
    # lista logo abaixo deles.
    tabelas = q("""select tabela, descricao, modulo, filtro, realizado, estimativa,
                          percentual, data_prev, etapa, responsavel, status
                     from cockpit.monitcad_tabelas
                    where medicao_id=%s and painel='cadastro'
                    order by modulo nulls last, realizado desc, tabela""", (ult["id"],))
    # Sem monitcad.estimativas cadastradas não existe % de carga — só contagem.
    tem_est = any(float(t["estimativa"] or 0) > 0 for t in tabelas)
    modulos = q("""select coalesce(modulo,'(sem módulo)') as modulo,
                          count(*) as tabelas,
                          count(*) filter (where realizado > 0) as com_carga,
                          sum(realizado) as registros
                     from cockpit.monitcad_tabelas
                    where medicao_id=%s and painel='cadastro'
                    group by 1 order by 4 desc, 1""", (ult["id"],))
    serie = q("""select m.data_medicao, m.semana,
                        sum(t.realizado)  as realizado,
                        sum(t.estimativa) as estimativa,
                        count(*) filter (where t.realizado > 0) as tabelas_com_carga,
                        case when sum(t.estimativa) > 0
                             then round(100.0 * sum(t.realizado) / sum(t.estimativa), 1)
                             else null end as pct
                   from cockpit.monitcad_medicoes m
                   join cockpit.monitcad_tabelas t on t.medicao_id = m.id
                  where m.customer=%s and m.ambiente=%s and t.painel='cadastro'
                  group by m.data_medicao, m.semana
                  order by m.data_medicao""", (customer, amb))
    return _json({"ok": True, "customer": customer, "projeto": proj, "ambiente": amb,
                  "vazio": False, "ultima_medicao": ult["data_medicao"],
                  "total_medicoes": len(meds), "tem_estimativa": tem_est,
                  "tabelas": tabelas, "modulos": modulos, "serie": serie})


SQL_EVO_MEDICOES = """
with med as (
  select id, data_medicao, hora_medicao,
         to_char(data_medicao,'IYYY')||'-W'||lpad(to_char(data_medicao,'IW'),2,'0') as semana_iso,
         row_number() over (partition by to_char(data_medicao,'IYYY-IW')
                            order by data_medicao desc, id desc) = 1 as fim_semana
    from cockpit.monitcad_medicoes
   where customer = %s and ambiente = %s
)
select m.data_medicao, m.semana_iso, m.fim_semana, m.hora_medicao,
       count(t.*)                              as tabelas,
       count(*) filter (where t.realizado > 0) as com_carga,
       sum(t.realizado)                        as realizado,
       sum(t.estimativa)                       as estimativa
  from med m
  join cockpit.monitcad_tabelas t on t.medicao_id = m.id and t.painel = 'cadastro'
 group by 1, 2, 3, 4
 order by 1
"""

SQL_EVO_MATRIZ = """
with med as (
  select id, data_medicao
    from cockpit.monitcad_medicoes
   where customer = %s and ambiente = %s
), ult as (
  select id from med order by data_medicao desc, id desc limit 1
)
select t.tabela,
       max(t.descricao) as descricao,
       max(t.modulo)    as modulo,
       max(t.estimativa) filter (where t.medicao_id = (select id from ult)) as estimativa,
       jsonb_object_agg(to_char(m.data_medicao,'YYYY-MM-DD'), t.realizado)  as serie
  from med m
  join cockpit.monitcad_tabelas t on t.medicao_id = m.id and t.painel = 'cadastro'
 group by t.tabela
 order by max(t.modulo) nulls last, t.tabela
"""


@app.get("/api/monitcad/<customer>/evolucao")
def api_monitcad_evolucao(customer):
    """Evolução dos cadastros entre medições: eixo de medições + matriz
    tabela x medição. Ranking e consolidação semanal são derivados no front a
    partir daqui — o banco devolve o retrato bruto, uma vez só.

    A estimativa vem SÓ da última medição: é o retrato vigente. Pegar o max()
    de todas faria uma estimativa antiga, já corrigida para baixo, continuar
    valendo. E `serie` só tem chave nas datas em que a tabela foi medida — a
    ausência da chave significa fora do escopo naquela data, que é diferente de
    medida com zero."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "cadastros")):
        return d
    amb = _ambiente()
    medicoes = q(SQL_EVO_MEDICOES, (customer, amb))
    if not medicoes:
        return _json({"ok": True, "customer": customer, "ambiente": amb,
                      "vazio": True, "medicoes": [], "tabelas": []})
    return _json({"ok": True, "customer": customer, "ambiente": amb,
                  "vazio": False, "medicoes": medicoes,
                  "tabelas": q(SQL_EVO_MATRIZ, (customer, amb))})


@app.delete("/api/monitcad/<customer>/medicao")
def api_monitcad_medicao_remover(customer):
    """Apaga uma medição inteira — a daquela data, naquele ambiente.

    Só admin: não há desfazer e a medição some do histórico de todo mundo. O
    caso real é o import que entrou zerado: enquanto ele está lá, a série mostra
    uma queda que não existiu."""
    if (r := require_admin()):
        return r
    # A simulação ('ver como') só restringe: apagar medição enquanto se olha a
    # tela de outro é o jeito mais fácil de destruir histórico sem perceber.
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de apagar medições.")
    if (d := deny_customer(customer)):
        return d
    amb = _ambiente()
    data = _data_iso(request.args.get("data"))
    if not data:
        return _err(400, "Informe a data da medição (AAAA-MM-DD).")

    ids = q("""select id from cockpit.monitcad_medicoes
                where customer=%s and ambiente=%s and data_medicao=%s""",
            (customer, amb, data))
    if not ids:
        return _err(404, f"Nenhuma medição de {data} na base {amb} deste cliente.")
    linhas = 0
    for a in ids:
        r0 = q("select count(*) n from cockpit.monitcad_tabelas where medicao_id=%s",
               (a["id"],), one=True)
        linhas += int(r0["n"] if r0 else 0)
        execute("delete from cockpit.monitcad_tabelas where medicao_id=%s", (a["id"],))
        execute("delete from cockpit.monitcad_medicoes where id=%s", (a["id"],))

    # ultima_medicao tem que recuar junto, senão o cabeçalho do projeto passa a
    # apontar para uma medição que não existe mais
    execute("""update cockpit.monitcad_projetos p
                  set ultima_medicao = (select max(m.data_medicao)
                                          from cockpit.monitcad_medicoes m
                                         where m.customer = p.customer
                                           and m.ambiente = 'producao'),
                      updated_at = now()
                where p.customer=%s""", (customer,))
    return _json({"ok": True, "customer": customer, "ambiente": amb, "data": data,
                  "medicoes": len(ids), "linhas": linhas})


# sm0 = P0 (leitura da SM0) e empresas = P1 da aba Empresas e Compartilhamento
SCRIPT_TIPOS = ("cadastros", "movimentos", "sm0", "empresas")


def _tipo_script():
    t = (request.args.get("tipo") or "cadastros").strip().lower()
    return t if t in SCRIPT_TIPOS else "cadastros"


@app.get("/api/monitcad/<customer>/script")
def api_monitcad_script(customer):
    """Script SQL de contagem customizado do cliente.

    Um por tipo (cadastros/movimentos) e válido para as DUAS bases: o escopo de
    tabelas é o mesmo em produção e em teste, o que muda é onde rodar — e disso
    cuida o cabeçalho, que o front reescreve na hora de exibir."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    tipo = _tipo_script()
    row = q("""select customer, tipo, sql, dialeto, sufixo, updated_by, updated_at
                 from cockpit.monitcad_scripts
                where customer=%s and tipo=%s""", (customer, tipo), one=True)
    return _json({"ok": True, "customer": customer, "tipo": tipo, "script": row})


@app.post("/api/monitcad/<customer>/script")
def api_monitcad_script_salvar(customer):
    """Salva o script colado. Quem enxerga o cliente pode salvar — mesma régua
    da importação de medição —, e fica registrado quem foi."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de salvar o script.")
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")

    body = request.get_json(silent=True) or {}
    sql = (body.get("sql") or "").strip()
    if len(sql) < 20:
        return _err(400, "Cole o script antes de salvar.")
    # O script nunca é executado aqui — vai para o console do Protheus. A única
    # checagem é que exista um SELECT, o que também deixa passar CTE (WITH ...).
    if not re.search(r"\bselect\b", sql, re.I):
        return _err(400, "O script precisa conter um SELECT.")

    tipo = _tipo_script()
    execute("""insert into cockpit.monitcad_scripts
                 (customer, tipo, sql, dialeto, sufixo, updated_by, updated_at)
               values (%s,%s,%s,%s,%s,%s, now())
               on conflict (customer, tipo) do update set
                 sql=excluded.sql, dialeto=excluded.dialeto, sufixo=excluded.sufixo,
                 updated_by=excluded.updated_by, updated_at=now()""",
            (customer, tipo, sql, (body.get("dialeto") or "")[:20] or None,
             (body.get("sufixo") or "")[:10] or None, current_user()))
    return _json({"ok": True, "customer": customer, "tipo": tipo, "bytes": len(sql)})


@app.delete("/api/monitcad/<customer>/script")
def api_monitcad_script_remover(customer):
    """Apaga o script salvo — o modal volta a abrir com o gerado."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de remover o script.")
    execute("delete from cockpit.monitcad_scripts where customer=%s and tipo=%s",
            (customer, _tipo_script()))
    return _json({"ok": True, "customer": customer, "tipo": _tipo_script(), "removido": True})


def _gravar_cadastros(customer, amb, body, origem="upload"):
    """Núcleo da carga de cadastros — o MESMO caminho para o arquivo subido à
    mão e para a coleta automática pela REST do cliente. De propósito: as regras
    que decidem o que entra (TABELAS_IGNORADAS) e em que painel entra
    (_painel_da_tabela) não podem ganhar uma segunda cópia, que é exatamente
    como uma tabela de movimento volta a aparecer em Cadastros.
    `medicoes` já vem validada pelo chamador."""
    medicoes = body.get("medicoes") or []
    execute("""insert into cockpit.monitcad_projetos
                 (customer, slug, projeto, cliente_nome, gp_totvs_nome, gp_cliente_nome, updated_at)
               values (%s,%s,%s,%s,%s,%s, now())
               on conflict (customer) do update set
                 projeto=coalesce(excluded.projeto, cockpit.monitcad_projetos.projeto),
                 cliente_nome=coalesce(excluded.cliente_nome, cockpit.monitcad_projetos.cliente_nome),
                 gp_totvs_nome=coalesce(excluded.gp_totvs_nome, cockpit.monitcad_projetos.gp_totvs_nome),
                 gp_cliente_nome=coalesce(excluded.gp_cliente_nome, cockpit.monitcad_projetos.gp_cliente_nome),
                 updated_at=now()""",
            (customer, _slug(body.get("projeto") or "") or None, body.get("projeto"),
             body.get("cliente"), body.get("gp_totvs"), body.get("gp_cliente")))

    n_med = n_tab = 0
    ultima = None
    for m in medicoes:
        dt = _data_iso(m.get("data_iso") or m.get("data_medicao"))
        if not dt:
            continue
        # regrava a medição inteira DAQUELE ambiente — produção e teste convivem
        # na mesma data sem uma apagar a outra
        antigos = q("""select id from cockpit.monitcad_medicoes
                        where customer=%s and data_medicao=%s and ambiente=%s""",
                    (customer, dt, amb))
        for a in antigos:
            execute("delete from cockpit.monitcad_tabelas where medicao_id=%s", (a["id"],))
            execute("delete from cockpit.monitcad_medicoes where id=%s", (a["id"],))
        row = q("""insert into cockpit.monitcad_medicoes
                     (customer, ambiente, data_medicao, semana, hora_medicao, origem, payload)
                   values (%s,%s,%s,%s,%s,%s,%s) returning id""",
                (customer, amb, dt, m.get("semana"), m.get("hora_medicao"), origem,
                 json.dumps({k: v for k, v in m.items() if k != "tabelas"})), one=True)
        mid = row["id"]
        n_med += 1
        ultima = max(ultima or dt, dt)
        linhas = []
        for t in (m.get("tabelas") or []):
            if (t.get("tabela") or "").strip().upper() in TABELAS_IGNORADAS:
                continue
            est = float(t.get("estimativa") or 0)
            real = float(t.get("realizado") or 0)
            pct = t.get("percentual")
            if pct is None:
                pct = round(100.0 * real / est, 1) if est > 0 else 0
            linhas.append((mid, customer, amb, dt, t.get("tabela"), t.get("descricao"),
                           t.get("modulo"), t.get("filtro"), real, est, pct,
                           _data_iso(t.get("data_prev")), t.get("etapa"),
                           t.get("responsavel"), t.get("status"),
                           _painel_da_tabela(t.get("tabela"))))
        if linhas:
            with db() as c, c.cursor() as cur:
                execute_values(cur, """
                    insert into cockpit.monitcad_tabelas
                      (medicao_id, customer, ambiente, data_medicao, tabela, descricao,
                       modulo, filtro, realizado, estimativa, percentual, data_prev,
                       etapa, responsavel, status, painel)
                    values %s""", linhas, page_size=200)
            n_tab += len(linhas)

    if ultima and amb == "producao":     # o marco do projeto é a base de produção
        execute("update cockpit.monitcad_projetos set ultima_medicao=%s, updated_at=now() "
                "where customer=%s", (ultima, customer))
    return {"medicoes": n_med, "tabelas": n_tab, "ultima_medicao": ultima}


@app.post("/api/monitcad/<customer>/upload")
def api_monitcad_upload(customer):
    """Carga manual enquanto o job do Protheus não roda. Aceita DOIS formatos:
    o historico-semanal.json e o CSV de status de cadastros (MODULO, TABELA,
    DESCRICAO, DT_LEITURA, SEMANA, QTDE). Idempotente: regrava a medição inteira
    quando a data já existe."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de subir medições.")

    if (f := request.files.get("arquivo")):
        bruto, nome = f.read(), (f.filename or "")
    else:
        bruto, nome = request.get_data(), ""
    if not bruto:
        return _err(400, "Arquivo vazio.")
    # O Protheus exporta em cp1252 — tenta os encodings na ordem mais provável
    texto = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    texto = texto if texto is not None else bruto.decode("utf-8", errors="replace")

    ehcsv = nome.lower().endswith(".csv") or "csv" in (request.content_type or "") \
        or not texto.lstrip().startswith(("{", "["))
    if ehcsv:
        try:
            body = _csv_para_body(texto)
        except Exception as e:
            return _err(400, f"CSV inválido: {e}")
    else:
        try:
            body = json.loads(texto)
        except Exception as e:
            return _err(400, f"JSON inválido: {e}")

    if not isinstance(body, dict):
        return _err(400, "Envie o historico-semanal.json ou o CSV de cadastros.")
    medicoes = body.get("medicoes") or []
    if not isinstance(medicoes, list) or not medicoes:
        return _err(400, "Arquivo sem medições.")

    # FK: o customer precisa existir em cockpit.clientes
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")
    amb = _ambiente()
    r = _gravar_cadastros(customer, amb, body)
    return _json({"ok": True, "customer": customer, "ambiente": amb,
                  "formato": "csv" if ehcsv else "json", **r})



# ── EMPRESAS E COMPARTILHAMENTO (MONITEMP) ──────────────────────────────────
# Cadastros conta "quanto foi carregado"; esta aba responde "ONDE foi
# carregado": em qual empresa (SRA010, SRA030…), com que conteúdo no campo
# filial e se isso bate com o compartilhamento da SX2 daquela empresa. Carga de
# legado costuma errar exatamente aqui — tudo na empresa 01, filial em branco
# numa tabela exclusiva, código de filial de outra empresa.
#
# Dois passos, de propósito:
#   P0  leitura MANUAL da SM0 (SYS_COMPANY) → cockpit.monitemp_sm0
#   P1  medição empresa × tabela × filial   → cockpit.monitemp_medicoes/_itens
# O script do P1 é MONTADO a partir do P0 (empresas e filiais viram literais):
# sem conferir a SM0 com o cliente antes, o script mede a estrutura errada.
# Gerador dos dois: web/assets/js/monitemp-sql.js (fonte única, browser + node).
# Spec: docs/specs/2026-09-18-aba-empresas-compartilhamento.md

EMP_SM0_COLS = {
    "EMPRESA": "empresa", "FILIAL": "filial", "NOMEEMPRESA": "nome_empresa",
    "NOMEFILIAL": "nome_filial", "CNPJ": "cnpj", "LEIAUTE": "leiaute",
    "SIZEFIL": "sizefil", "SX2": "sx2", "QTDTABELAS": "qtd_tabelas",
    "TABELASEXISTENTES": "tabelas", "DTLEITURA": "_dt", "SEMANA": "_semana",
}
EMP_COLS = {
    "EMPRESA": "empresa", "NOMEEMPRESA": "nome_empresa", "LEIAUTE": "leiaute",
    "TIPO": "tipo", "TABELA": "tabela", "DESCRICAO": "descricao",
    "TABELAFISICA": "tabela_fisica", "SITUACAO": "situacao", "SX2": "sx2",
    "FILIAL": "filial", "FILIALTIPO": "filial_tipo", "NOMEFILIAL": "nome_filial",
    "QTDE": "qtde", "QTD": "qtde", "DTLEITURA": "_dt", "SEMANA": "_semana",
}


def _csv_linhas(texto, mapa, obrig):
    """CSV → lista de dicts pelas colunas de `mapa`. Mesmo sniff do importador
    de cadastros (';' do console e da REST, ',' de planilha)."""
    try:
        dial = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dial = csv.excel
    linhas = list(csv.reader(io.StringIO(texto), dial))
    if not linhas:
        raise ValueError("CSV vazio.")
    cabec = [mapa.get(_norm_col(c)) for c in linhas[0]]
    falta = [c for c in obrig if mapa[c] not in cabec]
    if falta:
        raise ValueError("CSV sem a(s) coluna(s) " + ", ".join(falta) + ". Cabeçalho recebido: "
                         + ", ".join(c.strip() for c in linhas[0] if c.strip()))
    out = []
    for linha in linhas[1:]:
        if not any((c or "").strip() for c in linha):
            continue
        out.append({col: (val or "").strip() for col, val in zip(cabec, linha) if col})
    return out


def _csv_sm0(texto):
    regs = _csv_linhas(texto, EMP_SM0_COLS, ["EMPRESA", "FILIAL"])
    out = []
    for r in regs:
        if not r.get("empresa") or not r.get("filial"):
            continue
        r["sizefil"] = int(float(r["sizefil"])) if (r.get("sizefil") or "").strip() else None
        r["qtd_tabelas"] = int(float(r["qtd_tabelas"])) if (r.get("qtd_tabelas") or "").strip() else None
        r["sx2"] = (r.get("sx2") or "S").upper() != "N"
        out.append(r)
    if not out:
        raise ValueError("CSV da SM0 sem nenhuma empresa/filial.")
    return out


def _niveis_leiaute(leiaute, tam):
    """Tamanho de cada nível no código da filial. 'EEUUFF' → (2, 2, 2).
    Sem leiaute de gestão corporativa, tudo é filial."""
    lei = (leiaute or "").strip().upper()
    if not lei:
        return 0, 0, int(tam or 0)
    return lei.count("E"), lei.count("U"), lei.count("F")


def _consistencia(it, tam_fil):
    """Confronta o conteúdo do campo filial com o compartilhamento da SX2.

    SX2 vem como 'EMP|UNID|FIL' (X2_MODOEMP|X2_MODOUN|X2_MODO). O tamanho
    esperado do conteúdo é o do último nível EXCLUSIVO: tabela exclusiva por
    filial guarda o código inteiro; compartilhada na filial mas exclusiva na
    empresa guarda só a parte da empresa; tudo compartilhado guarda branco."""
    sit = (it.get("situacao") or "").upper()
    if sit == "NAO EXISTE":
        return "TABELA NAO EXISTE"
    if sit == "VAZIA":
        return "VAZIA"
    if sit == "SEM CAMPO FILIAL":
        return "SEM CAMPO FILIAL"
    sx2 = (it.get("sx2") or "").upper()
    partes = sx2.split("|")
    if len(partes) != 3:
        return sx2 or "SEM SX2"          # SEM SX2 / SEM REGISTRO NA SX2
    m_emp, m_un, m_fil = [p.strip() for p in partes]
    ne, nu, nf = _niveis_leiaute(it.get("leiaute"), tam_fil)
    total = ne + nu + nf
    if m_fil == "E":
        esperado = total
    elif m_un == "E" and nu:
        esperado = ne + nu
    elif m_emp == "E" and ne:
        esperado = ne
    else:
        esperado = 0
    tipo = (it.get("filial_tipo") or "").upper()
    valor = "" if tipo == "BRANCO" else (it.get("filial") or "").strip()
    if esperado == 0:
        return "OK" if not valor else "PREENCHIDA EM TABELA COMPARTILHADA"
    if not valor:
        return "BRANCO EM TABELA EXCLUSIVA"
    if total and len(valor) != esperado:
        return f"TAMANHO {len(valor)} - ESPERADO {esperado}"
    if tipo in ("FILIAL DE OUTRA EMPRESA", "NAO CADASTRADA NA SM0"):
        return tipo
    return "OK"


def _csv_empresas(texto):
    regs = _csv_linhas(texto, EMP_COLS, ["EMPRESA", "TABELA", "SITUACAO"])
    if not regs:
        raise ValueError("CSV sem linhas.")
    dt = hora = None
    semana = None
    itens = []
    for r in regs:
        d, h = _dt_hora(r.pop("_dt", ""))
        dt, hora = dt or d, hora or h
        s = r.pop("_semana", "")
        semana = semana or (int(re.sub(r"^.*W", "", s) or 0) or None if s else None)
        r["qtde"] = float(str(r.get("qtde") or "0").replace(",", ".") or 0)
        for k in ("filial", "filial_tipo", "nome_filial"):
            r[k] = r.get(k) or None
        itens.append(r)
    if not dt:
        raise ValueError("CSV sem data válida em DT_LEITURA.")
    return {"data_iso": dt, "hora_medicao": hora, "semana": semana, "itens": itens}


def _gravar_sm0(customer, amb, linhas):
    """A leitura da SM0 SUBSTITUI a anterior daquela base: filial que saiu da
    SM0 não pode continuar aparecendo como válida no painel."""
    lido = None
    for r in linhas:
        lido = lido or " ".join(x for x in _dt_hora(r.get("_dt", "")) if x) or None
    execute("delete from cockpit.monitemp_sm0 where customer=%s and ambiente=%s", (customer, amb))
    vals = [(customer, amb, r["empresa"], r["filial"], r.get("nome_empresa"), r.get("nome_filial"),
             r.get("cnpj"), r.get("leiaute"), r.get("sizefil"), r.get("sx2"),
             r.get("qtd_tabelas"), r.get("tabelas"), lido, current_user() or "coleta")
            for r in linhas]
    with db() as c, c.cursor() as cur:
        execute_values(cur, """
            insert into cockpit.monitemp_sm0
              (customer, ambiente, empresa, filial, nome_empresa, nome_filial, cnpj, leiaute,
               sizefil, sx2, qtd_tabelas, tabelas, lido_em, updated_by)
            values %s""", vals, page_size=200)
    return {"empresas": len({r["empresa"] for r in linhas}), "filiais": len(linhas)}


def _gravar_empresas(customer, amb, body, origem="upload"):
    dt = _data_iso(body.get("data_iso"))
    tam = {r["empresa"]: r["sizefil"] for r in q(
        "select empresa, max(sizefil) sizefil from cockpit.monitemp_sm0 "
        "where customer=%s and ambiente=%s group by empresa", (customer, amb))}
    for a in q("""select id from cockpit.monitemp_medicoes
                   where customer=%s and ambiente=%s and data_medicao=%s""", (customer, amb, dt)):
        execute("delete from cockpit.monitemp_medicoes where id=%s", (a["id"],))   # cascade
    mid = q("""insert into cockpit.monitemp_medicoes
                 (customer, ambiente, data_medicao, hora_medicao, semana, origem)
               values (%s,%s,%s,%s,%s,%s) returning id""",
            (customer, amb, dt, body.get("hora_medicao"), body.get("semana"), origem),
            one=True)["id"]
    vals = []
    for it in body["itens"]:
        vals.append((mid, customer, amb, it.get("empresa"), it.get("nome_empresa"),
                     it.get("leiaute"), it.get("tipo"), (it.get("tabela") or "").upper(),
                     it.get("descricao"), it.get("tabela_fisica"), it.get("situacao"),
                     it.get("sx2"), it.get("filial"), it.get("filial_tipo"),
                     it.get("nome_filial"), it.get("qtde") or 0,
                     _consistencia(it, tam.get(it.get("empresa")))))
    # O P1 so mede o que o P0 achou: tabela inexistente nao entra no SQL, senao o
    # script passa de 300 KB e a base devolve HTTP 500. As linhas "NAO EXISTE"
    # sao recompostas aqui, cruzando o escopo da aba Estrutura com as empresas da
    # SM0 - elas sao justamente uma das respostas que a aba tem de dar, entao nao
    # podem sumir da matriz so porque nao ha o que medir.
    escopo = q("""select tabela, descricao, tipo from cockpit.estrut_tabelas
                    where customer=%s and ativo order by tabela""", (customer,))
    if escopo:
        emps = q("""select distinct on (empresa) empresa, nome_empresa, leiaute
                      from cockpit.monitemp_sm0 where customer=%s and ambiente=%s
                     order by empresa""", (customer, amb))
        vistos = {(v[3], v[7]) for v in vals}
        for inf in emps:
            e = inf["empresa"]
            for tb in escopo:
                alias = (tb["tabela"] or "").upper()
                if not alias or (e, alias) in vistos:
                    continue
                vals.append((mid, customer, amb, e, inf.get("nome_empresa"), inf.get("leiaute"),
                             (tb.get("tipo") or "cadastro").upper(), alias, tb.get("descricao"),
                             f"{alias}{e}0", "NAO EXISTE", None, None, None, None, 0,
                             _consistencia({"situacao": "NAO EXISTE"}, tam.get(e))))
    with db() as c, c.cursor() as cur:
        execute_values(cur, """
            insert into cockpit.monitemp_itens
              (medicao_id, customer, ambiente, empresa, nome_empresa, leiaute, tipo, tabela,
               descricao, tabela_fisica, situacao, sx2, filial, filial_tipo, nome_filial,
               qtde, consistencia)
            values %s""", vals, page_size=500)
    return {"medicoes": 1, "itens": len(vals), "ultima_medicao": dt}


def _texto_upload():
    if (f := request.files.get("arquivo")):
        bruto = f.read()
    else:
        bruto = request.get_data()
    if not bruto:
        return None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return bruto.decode(enc)
        except UnicodeDecodeError:
            continue
    return bruto.decode("utf-8", errors="replace")


@app.get("/api/monitemp/<customer>")
def api_monitemp(customer):
    """SM0 lida + última medição (itens) + lista de medições da base pedida."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "empresas")):
        return d
    amb = _ambiente()
    sm0 = q("""select empresa, filial, nome_empresa, nome_filial, cnpj, leiaute, sizefil,
                      sx2, qtd_tabelas, tabelas, lido_em, updated_by, updated_at
                 from cockpit.monitemp_sm0 where customer=%s and ambiente=%s
                order by empresa, filial""", (customer, amb))
    meds = q("""select id, data_medicao, hora_medicao, semana, origem
                  from cockpit.monitemp_medicoes where customer=%s and ambiente=%s
                 order by data_medicao""", (customer, amb))
    data = request.args.get("data")
    alvo = next((m for m in meds if str(m["data_medicao"]) == data), meds[-1] if meds else None)
    itens = []
    if alvo:
        itens = q("""select empresa, nome_empresa, leiaute, tipo, tabela, descricao, tabela_fisica,
                            situacao, sx2, filial, filial_tipo, nome_filial, qtde, consistencia
                       from cockpit.monitemp_itens where medicao_id=%s
                      order by empresa, tipo, tabela, filial nulls first""", (alvo["id"],))
    # O ESCOPO medido vem da aba Estrutura: o que a consultoria definiu para este
    # cliente e o que os scripts P0/P1 vao procurar. Vazio (cliente que ainda nao
    # definiu nada) cai na lista padrao do gerador - senao o painel nao mediria
    # nada e pareceria quebrado.
    escopo = q("""select modulo, tabela, descricao, tipo
                    from cockpit.estrut_tabelas where customer=%s and ativo
                   order by modulo, tabela""", (customer,))
    # A coluna "Tabelas do escopo" da SM0 e CONGELADA no P0: ela diz o que aquela
    # leitura procurou e achou. Se o escopo mudou depois, o numerador vira uma
    # mentira silenciosa ("25 de 131" quando so procurou 34), entao a tela precisa
    # saber que a leitura ficou para tras.
    lim = q("""select max(greatest(coalesce(definido_em, created_at), created_at)) em
                 from cockpit.estrut_tabelas where customer=%s and ativo""",
            (customer,), one=True)
    sm0_em = max([r["updated_at"] for r in sm0 if r.get("updated_at")], default=None)
    escopo_mudou = bool(lim and lim.get("em") and sm0_em and lim["em"] > sm0_em)
    return _json({"ok": True, "customer": customer, "ambiente": amb, "sm0": sm0,
                  "medicoes": meds, "medicao": alvo, "itens": itens,
                  "escopo": escopo, "escopo_mudou": escopo_mudou})


@app.post("/api/monitemp/<customer>/sm0")
def api_monitemp_sm0(customer):
    """Sobe o CSV do P0 (leitura da SM0) daquela base."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de subir a SM0.")
    texto = _texto_upload()
    if not texto:
        return _err(400, "Arquivo vazio.")
    try:
        linhas = _csv_sm0(texto)
    except Exception as e:
        return _err(400, f"CSV da SM0 inválido: {e}")
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")
    amb = _ambiente()
    return _json({"ok": True, "customer": customer, "ambiente": amb,
                  **_gravar_sm0(customer, amb, linhas)})


@app.post("/api/monitemp/<customer>/upload")
def api_monitemp_upload(customer):
    """Sobe o CSV da medição de Empresas (script P1). Regrava a data."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de subir medições.")
    texto = _texto_upload()
    if not texto:
        return _err(400, "Arquivo vazio.")
    try:
        body = _csv_empresas(texto)
    except Exception as e:
        return _err(400, f"CSV inválido: {e}")
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")
    amb = _ambiente()
    return _json({"ok": True, "customer": customer, "ambiente": amb,
                  **_gravar_empresas(customer, amb, body)})


@app.delete("/api/monitemp/<customer>/medicao")
def api_monitemp_medicao_remover(customer):
    """Apaga a medição de uma data — só admin, como em Cadastros."""
    if (r := require_admin()):
        return r
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de apagar medições.")
    if (d := deny_customer(customer)):
        return d
    data = _data_iso(request.args.get("data"))
    if not data:
        return _err(400, "Informe a data da medição (AAAA-MM-DD).")
    ids = q("""select id from cockpit.monitemp_medicoes
                where customer=%s and ambiente=%s and data_medicao=%s""",
            (customer, _ambiente(), data))
    for a in ids:
        execute("delete from cockpit.monitemp_medicoes where id=%s", (a["id"],))
    return _json({"ok": True, "removidas": len(ids)})


# ── MOVIMENTOS — cobertura de cenários, não volume por tabela ───────────────
# O total por tabela responde "tem movimento?". A pergunta que decide go-live é
# outra: "qual cenário ainda não rodou?" — venda para contribuinte de outro
# estado, NCM com ST, baixa com adiantamento, transferência entre contas. Numa
# contagem por tabela isso tudo vira uma linha só.
# Por isso a medição de movimentos carrega uma quebra dimensional
# (cockpit.monitmov_dimensoes, alimentada pelo SQL D5) e é cruzada com a lista
# de cenários combinados com o cliente (cockpit.monitmov_cenarios).
# Spec: docs/specs/2026-08-25-movimentos-cobertura-cenarios.md

# Colunas DIM1..DIM4 genéricas de propósito: incluir uma análise nova
# (SD1 de entrada, SD3 de estoque) não muda layout de arquivo nem banco nem tela.
MOV_CSV_COLS = {
    "ANALISE": "analise", "ANALISE1": "analise", "TIPOANALISE": "analise",
    "DESCRICAO": "descricao", "DESCR": "descricao", "OBSERVACAO": "descricao",
    "DIM1NOME": "dim1_nome", "DIM2NOME": "dim2_nome",
    "DIM3NOME": "dim3_nome", "DIM4NOME": "dim4_nome",
    "DIM1": "dim1", "DIM2": "dim2", "DIM3": "dim3", "DIM4": "dim4",
    "PERIODO": "periodo",
    "QTDE": "qtde", "QTD": "qtde", "QUANTIDADE": "qtde", "REGISTROS": "qtde",
    "CONTAGEM": "qtde", "LINHAS": "qtde", "REALIZADO": "qtde",
    "QTDDOC": "qtd_doc", "QTDDOCS": "qtd_doc", "DOCUMENTOS": "qtd_doc",
    "QTDDOCUMENTOS": "qtd_doc", "DOCS": "qtd_doc",
    "DTLEITURA": "_dt", "DATA": "_dt", "DATAMEDICAO": "_dt", "SEMANA": "_semana",
}

# Uma linha em monitmov_itens por análise, derivada da quebra — assim os KPIs e
# o gráfico por módulo que já existiam continuam de pé sem um segundo arquivo.
MOV_ANALISES = {
    "SD2_FISCAL":   ("SD2", "Fiscal",     "Itens de NF de saída (UF × contribuinte × NCM × CFOP)"),
    "SE5_BANCARIO": ("SE5", "Financeiro", "Movimento bancário (operação × sentido × tipo)"),
}


def _periodo_datas(txt):
    """'20260101-20991231' → (date, date). Aceita AAAAMMDD, AAAA-MM-DD e
    DD/MM/AAAA dos dois lados; qualquer coisa fora disso vira (None, None) e o
    período fica só como texto — não é motivo para recusar o arquivo."""
    partes = re.split(r"\s*(?:-{1,2}|até|ate|to|a)\s*", (txt or "").strip(),
                      maxsplit=1, flags=re.I)
    if len(partes) != 2:
        m = re.fullmatch(r"(\d{8})\D+(\d{8})", (txt or "").strip())
        partes = [m.group(1), m.group(2)] if m else []
    saida = []
    for p in partes[:2]:
        p = (p or "").strip()
        if re.fullmatch(r"\d{8}", p):
            saida.append(f"{p[0:4]}-{p[4:6]}-{p[6:8]}")
        else:
            saida.append(_dt_hora(p)[0])   # aceita DD/MM/AAAA além de ISO
    return (saida + [None, None])[:2]


def _num_br(v):
    """Número do export do Protheus: '1.240' é mil duzentos e quarenta, mas
    '1240.5' é um decimal. Só trata ponto como separador de milhar quando o
    formato é inequívoco — trocar sempre quebraria a segunda forma."""
    s = str(v if v is not None else "").strip()
    if not s:
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _csv_movimentos(texto):
    """CSV do D5 → mesmo formato do JSON de importação. Uma medição por data
    encontrada em DT_LEITURA."""
    try:
        dial = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dial = csv.excel
    linhas = list(csv.reader(io.StringIO(texto), dial))
    if not linhas:
        raise ValueError("CSV vazio.")
    cabec = [MOV_CSV_COLS.get(_norm_col(c)) for c in linhas[0]]
    if "analise" not in cabec:
        raise ValueError("CSV sem a coluna ANALISE. Esperado o layout do D5: "
                         "ANALISE, DESCRICAO, DIM1_NOME, DIM1 … QTDE, QTD_DOC.")
    # Mesma lição do import de cadastros: coluna de quantidade não reconhecida
    # entraria como medição inteira zerada e pareceria base vazia.
    if "qtde" not in cabec:
        raise ValueError("CSV sem coluna de quantidade — aceito: QTDE, QTD, "
                         "QUANTIDADE, REGISTROS, LINHAS. Cabeçalho recebido: "
                         + ", ".join(c.strip() for c in linhas[0] if c.strip()))

    por_data = {}
    for linha in linhas[1:]:
        if not any((c or "").strip() for c in linha):
            continue
        reg = {}
        for col, val in zip(cabec, linha):
            if col:
                reg[col] = (val or "").strip()
        if not reg.get("analise"):
            continue
        dt, hora = _dt_hora(reg.pop("_dt", ""))
        if not dt:
            raise ValueError("CSV sem data válida em DT_LEITURA (ex.: 25/08/2026 11:43).")
        semana = reg.pop("_semana", "")
        med = por_data.setdefault(dt, {
            "data_iso": dt, "hora_medicao": hora,
            "semana": int(re.sub(r"^.*W", "", semana) or 0) or None,
            "periodo": reg.get("periodo") or "",
            "dimensoes": [],
        })
        reg["qtde"] = _num_br(reg.get("qtde"))
        reg["qtd_doc"] = _num_br(reg.get("qtd_doc")) if reg.get("qtd_doc") else None
        med["dimensoes"].append(reg)
    return {"medicoes": list(por_data.values())}


def _norm_dim(v):
    return (v or "").strip().upper()


def _casa_cenario(cen, obs):
    """Coringa '*' (ou vazio) casa com qualquer valor. Comparação é por texto
    normalizado: o que vem do banco do cliente tem padding e caixa variável."""
    if _norm_dim(cen["analise"]) != _norm_dim(obs["analise"]):
        return False
    for k in ("dim1", "dim2", "dim3", "dim4"):
        alvo = _norm_dim(cen.get(k)) or "*"
        if alvo == "*":
            continue
        if _norm_dim(obs.get(k)) != alvo:
            return False
    return True


def _cobertura(dimensoes, cenarios):
    """Cruza esperado × observado. Devolve a lista de cenários com status e a
    lista do que foi observado sem cenário que case.

    NÃO PREVISTO não é erro: numa base de produção é o normal enquanto a lista
    de cenários não está fechada. Vira sinal depois que o cliente combinou tudo.
    """
    cobertura, casados = [], set()
    for c in cenarios:
        achou = [o for o in dimensoes if _casa_cenario(c, o)]
        qtde = sum(float(o.get("qtde") or 0) for o in achou)
        for o in achou:
            casados.add(id(o))
        cobertura.append({
            **c,
            "qtde": qtde,
            "qtd_doc": sum(float(o.get("qtd_doc") or 0) for o in achou),
            "ocorrencias": len(achou),
            "status": ("COBERTO" if qtde > 0
                       else ("FALTANTE" if c.get("esperado", True) else "OPCIONAL")),
        })
    # Só faz sentido apontar "não previsto" para a análise que já tem cenário
    # combinado — senão a primeira medição do cliente vira uma lista de acusações.
    com_cenario = {_norm_dim(c["analise"]) for c in cenarios}
    nao_previstos = [o for o in dimensoes
                     if id(o) not in casados
                     and _norm_dim(o.get("analise")) in com_cenario
                     and float(o.get("qtde") or 0) > 0]
    ordem = {"FALTANTE": 0, "COBERTO": 1, "OPCIONAL": 2}
    cobertura.sort(key=lambda c: (ordem.get(c["status"], 9), c["analise"],
                                  c.get("descricao") or ""))
    nao_previstos.sort(key=lambda o: -float(o.get("qtde") or 0))
    return cobertura, nao_previstos


SQL_MOV_EVO_DIM = """
  select medicao_id, analise, descricao,
         dim1_nome, dim1, dim2_nome, dim2, dim3_nome, dim3, dim4_nome, dim4,
         qtde, qtd_doc
    from cockpit.monitmov_dimensoes
   where medicao_id = any(%s)
"""


def _chave_dim(d):
    """Identidade de uma combinação observada, estável entre medições."""
    return "|".join(_norm_dim(d.get(k)) for k in ("analise", "dim1", "dim2", "dim3", "dim4"))


@app.get("/api/monitmov/<customer>/evolucao")
def api_monitmov_evolucao(customer):
    """Evolução dos MOVIMENTOS: uma linha por medição com movimentos, combinações
    e cobertura, mais a matriz cenário × medição.

    A cobertura é recalculada com `_cobertura()` — a MESMA função da "Situação
    atual". Reescrever a regra de casamento aqui (coringa, normalização, o que
    conta como não previsto) faria as duas visões divergirem no dia em que uma
    delas mudasse, e o usuário veria 7/11 numa aba e 8/11 na outra sem explicação.

    Uma consulta só traz as dimensões de TODAS as medições; agrupar em Python
    evita N+1 no Postgres — são poucas medições, mas centenas de combinações
    cada."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "cadastros")):
        return d
    amb = _ambiente()
    meds = q("""select id, data_medicao, semana, hora_medicao,
                       periodo_ini, periodo_fim, origem
                  from cockpit.monitmov_medicoes
                 where customer=%s and ambiente=%s order by data_medicao, id""",
             (customer, amb))
    cenarios = q("""select id, analise, dim1, dim2, dim3, dim4, descricao,
                           esperado, etapa, responsavel
                      from cockpit.monitmov_cenarios where customer=%s
                     order by analise, descricao""", (customer,))
    if not meds:
        return _json({"ok": True, "customer": customer, "ambiente": amb,
                      "vazio": True, "medicoes": [], "cenarios": [], "serie": []})

    dims = q(SQL_MOV_EVO_DIM, ([m["id"] for m in meds],))
    por_med = {}
    for d in dims:
        por_med.setdefault(d["medicao_id"], []).append(d)

    # Análises que TÊM cenário combinado. As que não têm (o SD2_FISCAL do Olim,
    # com 243 movimentos e 60 combinações) sumiam da evolução inteira: cobertura
    # não se aplica a elas, e sem um bloco próprio o lado fiscal ficava invisível.
    com_cenario = {_norm_dim(c["analise"]) for c in cenarios}

    serie, matriz, analises, fora = [], {}, {}, {}
    for m in meds:
        ds = por_med.get(m["id"], [])
        cob, nao_prev = _cobertura(ds, cenarios)
        cobertos = sum(1 for c in cob if c["status"] == "COBERTO")
        faltantes = sum(1 for c in cob if c["status"] == "FALTANTE")
        data = str(m["data_medicao"])

        # ── por análise: vale para TODAS, tenham cenário ou não ──────────────
        for d in ds:
            a = d["analise"]
            linha = analises.setdefault(_norm_dim(a), {
                "analise": a, "tem_cenario": _norm_dim(a) in com_cenario, "serie": {}})
            x = linha["serie"].setdefault(data, {"movimentos": 0.0, "documentos": 0.0,
                                                 "combinacoes": 0})
            x["movimentos"] += float(d["qtde"] or 0)
            x["documentos"] += float(d["qtd_doc"] or 0)
            x["combinacoes"] += 1
        for c in cob:
            linha = analises.get(_norm_dim(c["analise"]))
            if not linha:            # cenário de análise que não veio nesta medição
                continue
            x = linha["serie"].setdefault(data, {"movimentos": 0.0, "documentos": 0.0,
                                                 "combinacoes": 0})
            x["cenarios"] = x.get("cenarios", 0) + 1
            if c["status"] == "COBERTO":
                x["cobertos"] = x.get("cobertos", 0) + 1
            elif c["status"] == "FALTANTE":
                x["faltantes"] = x.get("faltantes", 0) + 1

        # ── fora do combinado: a lista, não só a contagem ────────────────────
        for o in nao_prev:
            k = _chave_dim(o)
            item = fora.setdefault(k, {
                "analise": o["analise"],
                "dims": [(o.get("dim1_nome"), o.get("dim1")), (o.get("dim2_nome"), o.get("dim2")),
                         (o.get("dim3_nome"), o.get("dim3")), (o.get("dim4_nome"), o.get("dim4"))],
                "serie": {}})
            item["serie"][data] = item["serie"].get(data, 0) + float(o.get("qtde") or 0)
        serie.append({
            "data_medicao": m["data_medicao"], "semana": m["semana"],
            "hora_medicao": m["hora_medicao"],
            "periodo_ini": m["periodo_ini"], "periodo_fim": m["periodo_fim"],
            "movimentos": sum(float(d["qtde"] or 0) for d in ds),
            "documentos": sum(float(d["qtd_doc"] or 0) for d in ds),
            "combinacoes": len(ds),
            "analises": len({_norm_dim(d["analise"]) for d in ds}),
            "cobertos": cobertos, "faltantes": faltantes,
            "cenarios": len(cob), "nao_previstos": len(nao_prev),
        })
        # A chave da matriz é o id do cenário: descrição repete entre análises.
        for c in cob:
            linha = matriz.setdefault(c["id"], {
                "id": c["id"], "analise": c["analise"], "descricao": c["descricao"],
                "esperado": c["esperado"], "etapa": c["etapa"],
                "responsavel": c["responsavel"], "serie": {}})
            linha["serie"][data] = {"qtde": c["qtde"], "status": c["status"]}

    # Faltante primeiro: é o que precisa de ação. Depois o que tem menos volume.
    ordem = {"FALTANTE": 0, "COBERTO": 1, "OPCIONAL": 2}
    ultima = str(meds[-1]["data_medicao"])
    linhas = sorted(matriz.values(),
                    key=lambda l: (ordem.get((l["serie"].get(ultima) or {}).get("status"), 9),
                                   l["analise"], l["descricao"] or ""))
    # Fora do combinado ordenado pelo que MAIS pesa na última medição: é a fila
    # de trabalho ("isto virou regra ou é ruído?"), não um relatório de acusação.
    lista_fora = sorted(fora.values(),
                        key=lambda f: -(f["serie"].get(ultima) or 0))
    return _json({"ok": True, "customer": customer, "ambiente": amb, "vazio": False,
                  "medicoes": [str(m["data_medicao"]) for m in meds],
                  "serie": serie, "cenarios": linhas,
                  "analises": sorted(analises.values(), key=lambda a: a["analise"]),
                  "fora_combinado": lista_fora})


@app.get("/api/monitmov/<customer>")
def api_monitmov(customer):
    """Última medição de MOVIMENTOS do cliente no ambiente pedido: totais por
    análise, quebra dimensional e cobertura de cenários."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "cadastros", "cobertura")):
        return d
    amb = _ambiente()
    meds = q("""select id, data_medicao, semana, periodo_ini, periodo_fim, origem
                  from cockpit.monitmov_medicoes
                 where customer=%s and ambiente=%s order by data_medicao""",
             (customer, amb))
    cenarios = q("""select id, analise, dim1, dim2, dim3, dim4, descricao,
                           esperado, etapa, responsavel
                      from cockpit.monitmov_cenarios where customer=%s
                     order by analise, descricao""", (customer,))
    if not meds:
        return _json({"ok": True, "customer": customer, "ambiente": amb, "vazio": True,
                      "itens": [], "modulos": [], "serie": [], "total_medicoes": 0,
                      "dimensoes": [], "analises": [], "cenarios": cenarios,
                      "cobertura": [], "nao_previstos": [], "a_classificar": []})
    ult = meds[-1]
    itens = q("""select tabela, descricao, modulo, filtro, quantidade, valor, periodo
                   from cockpit.monitmov_itens where medicao_id=%s
                  order by modulo nulls last, quantidade desc, tabela""", (ult["id"],))
    modulos = q("""select coalesce(modulo,'(sem módulo)') as modulo, count(*) as tabelas,
                          count(*) filter (where quantidade > 0) as com_movimento,
                          sum(quantidade) as registros
                     from cockpit.monitmov_itens where medicao_id=%s
                    group by 1 order by 4 desc, 1""", (ult["id"],))
    serie = q("""select m.data_medicao, m.semana, sum(i.quantidade) as realizado,
                        count(*) filter (where i.quantidade > 0) as tabelas_com_carga
                   from cockpit.monitmov_medicoes m
                   join cockpit.monitmov_itens i on i.medicao_id = m.id
                  where m.customer=%s and m.ambiente=%s
                  group by m.data_medicao, m.semana order by m.data_medicao""",
              (customer, amb))
    dimensoes = q("""select analise, descricao, dim1_nome, dim1, dim2_nome, dim2,
                            dim3_nome, dim3, dim4_nome, dim4, periodo, qtde, qtd_doc
                       from cockpit.monitmov_dimensoes where medicao_id=%s
                      order by analise, qtde desc, dim1, dim2, dim3, dim4""",
                  (ult["id"],))
    analises = q("""select analise, count(*) as combinacoes,
                           count(*) filter (where qtde > 0) as com_movimento,
                           count(distinct dim1) as valores_dim1,
                           max(dim1_nome) as dim1_nome, max(dim2_nome) as dim2_nome,
                           max(dim3_nome) as dim3_nome, max(dim4_nome) as dim4_nome,
                           sum(qtde) as qtde, sum(qtd_doc) as qtd_doc
                      from cockpit.monitmov_dimensoes where medicao_id=%s
                     group by analise order by analise""", (ult["id"],))
    cobertura, nao_previstos = _cobertura(dimensoes, cenarios)
    # Fila de trabalho do consultor: o que o de-para bancário ainda não sabe ler
    a_classificar = [d for d in dimensoes if _norm_dim(d.get("dim1")) == "(A CLASSIFICAR)"]
    return _json({"ok": True, "customer": customer, "ambiente": amb, "vazio": False,
                  "ultima_medicao": ult["data_medicao"], "total_medicoes": len(meds),
                  "periodo_ini": ult["periodo_ini"], "periodo_fim": ult["periodo_fim"],
                  "itens": itens, "modulos": modulos, "serie": serie,
                  "dimensoes": dimensoes, "analises": analises, "cenarios": cenarios,
                  "cobertura": cobertura, "nao_previstos": nao_previstos,
                  "a_classificar": a_classificar})


def _gravar_movimentos(customer, amb, body, origem="upload"):
    """Núcleo da carga de movimentos (cobertura de cenários). Mesmo motivo do
    _gravar_cadastros: upload manual e coleta automática entram pela mesma
    porta."""
    n_med = n_dim = 0
    ultima = None
    for m in body["medicoes"]:
        dt = _data_iso(m.get("data_iso") or m.get("data_medicao"))
        if not dt:
            continue
        dims = m.get("dimensoes") or []
        periodo = m.get("periodo") or (dims[0].get("periodo") if dims else "")
        p_ini, p_fim = _periodo_datas(periodo)
        for a in q("""select id from cockpit.monitmov_medicoes
                       where customer=%s and data_medicao=%s and ambiente=%s""",
                   (customer, dt, amb)):
            execute("delete from cockpit.monitmov_dimensoes where medicao_id=%s", (a["id"],))
            execute("delete from cockpit.monitmov_itens where medicao_id=%s", (a["id"],))
            execute("delete from cockpit.monitmov_medicoes where id=%s", (a["id"],))
        row = q("""insert into cockpit.monitmov_medicoes
                     (customer, ambiente, data_medicao, semana, hora_medicao,
                      periodo_ini, periodo_fim, origem, payload)
                   values (%s,%s,%s,%s,%s,%s,%s,%s,%s) returning id""",
                (customer, amb, dt, m.get("semana"), m.get("hora_medicao"),
                 p_ini, p_fim, origem,
                 json.dumps({k: v for k, v in m.items() if k != "dimensoes"})), one=True)
        mid = row["id"]
        n_med += 1
        ultima = max(ultima or dt, dt)

        linhas, porAnalise = [], {}
        for x in dims:
            analise = (x.get("analise") or "").strip().upper()
            if not analise:
                continue
            qt = float(x.get("qtde") or 0)
            qd = x.get("qtd_doc")
            linhas.append((mid, customer, amb, dt, analise, x.get("descricao"),
                           x.get("dim1_nome"), x.get("dim1"), x.get("dim2_nome"), x.get("dim2"),
                           x.get("dim3_nome"), x.get("dim3"), x.get("dim4_nome"), x.get("dim4"),
                           x.get("periodo") or periodo, qt,
                           float(qd) if qd not in (None, "") else None))
            ag = porAnalise.setdefault(analise, {"qtde": 0.0, "periodo": x.get("periodo") or periodo})
            ag["qtde"] += qt
        if linhas:
            with db() as c, c.cursor() as cur:
                execute_values(cur, """
                    insert into cockpit.monitmov_dimensoes
                      (medicao_id, customer, ambiente, data_medicao, analise, descricao,
                       dim1_nome, dim1, dim2_nome, dim2, dim3_nome, dim3,
                       dim4_nome, dim4, periodo, qtde, qtd_doc)
                    values %s""", linhas, page_size=200)
            n_dim += len(linhas)
        # itens = um resumo por análise, para os KPIs e o gráfico por módulo
        resumo = []
        for analise, ag in porAnalise.items():
            tab, mod, desc = MOV_ANALISES.get(
                analise, (analise.split("_")[0], "(sem módulo)", analise))
            resumo.append((mid, customer, amb, dt, tab, desc, mod, None,
                           ag["qtde"], None, ag["periodo"]))
        if resumo:
            with db() as c, c.cursor() as cur:
                execute_values(cur, """
                    insert into cockpit.monitmov_itens
                      (medicao_id, customer, ambiente, data_medicao, tabela, descricao,
                       modulo, filtro, quantidade, valor, periodo)
                    values %s""", resumo, page_size=100)
    return {"medicoes": n_med, "dimensoes": n_dim, "ultima_medicao": ultima}


@app.post("/api/monitmov/<customer>/upload")
def api_monitmov_upload(customer):
    """Importa o export do D5 (CSV) ou o mesmo conteúdo em JSON. Idempotente:
    regrava a medição inteira daquela data, naquele ambiente."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de subir movimentos.")

    if (f := request.files.get("arquivo")):
        bruto, nome = f.read(), (f.filename or "")
    else:
        bruto, nome = request.get_data(), ""
    if not bruto:
        return _err(400, "Arquivo vazio.")
    texto = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    texto = texto if texto is not None else bruto.decode("utf-8", errors="replace")

    ehcsv = nome.lower().endswith(".csv") or "csv" in (request.content_type or "") \
        or not texto.lstrip().startswith(("{", "["))
    try:
        body = _csv_movimentos(texto) if ehcsv else json.loads(texto)
    except Exception as e:
        return _err(400, f"{'CSV' if ehcsv else 'JSON'} inválido: {e}")
    if not isinstance(body, dict) or not body.get("medicoes"):
        return _err(400, "Arquivo sem medições.")
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")
    amb = _ambiente()
    r = _gravar_movimentos(customer, amb, body)
    return _json({"ok": True, "customer": customer, "ambiente": amb,
                  "formato": "csv" if ehcsv else "json", **r})


# ── Ambientes Protheus — coleta automática pela REST do cliente ─────────────
# A medição sempre nasceu de um caminho manual: abrir o modal "Script SQL",
# copiar, colar no console TCloud, exportar CSV e arrastar o arquivo aqui. O
# fonte TLPP TSCMONITREST (fontes/tlpp/ do cockpit) publica na base do cliente
# um POST /tscmonit/query que recebe ESSE MESMO script e devolve ESSE MESMO
# CSV — então a coleta automática entra pelo importador de sempre
# (_gravar_cadastros / _gravar_movimentos), sem uma segunda cópia das regras.
#
# Produção e teste são bases DIFERENTES: uma linha por (customer, ambiente),
# cada uma com URL, environment, empresa/filial, sufixo e banco próprios.
# É por isso que o teste de conexão compara environment e banco com o que o
# /ping devolve: coletar da base errada é o erro caro deste processo — o número
# chega bonito, plausível, e vem do lugar errado.

BANCOS = ("oracle", "mssql", "postgres")
# Com ssl_verificar=false o urllib3 imprime um InsecureRequestWarning por
# chamada. A escolha e do consultor, por ambiente, e ja aparece na tela — o
# aviso so polui o log da Vercel.
try:
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
except Exception:
    pass
TIMEOUT_TETO = 50          # a função da Vercel morre em 60s (vercel.json)


def _cred_key():
    b = base64.b64decode(os.environ.get("PROTHEUS_CRED_KEY", "") or "")
    if len(b) != 32:
        raise RuntimeError("PROTHEUS_CRED_KEY ausente ou inválida "
                           "(32 bytes em base64: openssl rand -base64 32).")
    return b


def _cifra(txt, aad):
    """base64(iv||tag||ct), AES-256-GCM — o mesmo formato do cockpit
    (src/lib/tasks-sc/crypto.ts). A AAD amarra o segredo ao par
    cliente/ambiente: um blob copiado para outra linha não decifra."""
    if not txt:
        return None
    iv = os.urandom(12)
    sel = AESGCM(_cred_key()).encrypt(iv, txt.encode(), aad.encode())  # ct||tag
    return base64.b64encode(iv + sel[-16:] + sel[:-16]).decode()


def _decifra(blob, aad):
    if not blob:
        return ""
    raw = base64.b64decode(blob)
    iv, tag, ct = raw[:12], raw[12:28], raw[28:]
    return AESGCM(_cred_key()).decrypt(iv, ct + tag, aad.encode()).decode()


def _aad(customer, ambiente):
    return f"{customer}|{ambiente}"


def _amb_path(ambiente):
    a = (ambiente or "").strip().lower()
    return a if a in AMBIENTES else None


def _cfg_ambiente(customer, ambiente):
    return q("select * from cockpit.protheus_ambientes where customer=%s and ambiente=%s",
             (customer, ambiente), one=True)


def _log_coleta(customer, ambiente, tipo, ini, ok, **kw):
    """Log de cada disparo. Sem isso, 'não atualizou' vira caça ao log da
    Vercel — e o cliente não tem como provar quando a leitura foi feita."""
    try:
        execute("""insert into cockpit.protheus_coletas
                     (customer, ambiente, tipo, iniciado_em, duracao_ms, ok, http_status,
                      linhas, medicoes, itens, bytes, data_medicao, erro, disparado_por)
                   values (%s,%s,%s, to_timestamp(%s), %s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                (customer, ambiente, tipo, ini, int((time.time() - ini) * 1000), ok,
                 kw.get("http_status"), kw.get("linhas"), kw.get("medicoes"),
                 kw.get("itens"), kw.get("bytes"), kw.get("data_medicao"),
                 (kw.get("erro") or "")[:2000] or None, current_user()))
    except Exception:
        pass          # log não pode derrubar a coleta


def _guarda_ambiente(customer, ambiente, escrita=False):
    """Devolve (erro_response, ambiente_normalizado). Guarda das rotas
    /api/protheus/* — credencial e coleta do ambiente do cliente são operação,
    não consulta: o usuário do cliente não passa por aqui."""
    if (r := require_interno()):
        return r, None
    if (d := deny_customer(customer)):
        return d, None
    amb = _amb_path(ambiente)
    if not amb:
        return _err(400, "Ambiente inválido — use 'producao' ou 'teste'."), None
    if escrita and effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de alterar ambientes."), None
    return None, amb


@app.get("/api/protheus/<customer>")
def api_protheus_listar(customer):
    """Parâmetros dos dois ambientes + as últimas coletas. Segredo NUNCA sai
    daqui: só os booleanos tem_senha/tem_token, que é o que a tela precisa."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    rows = q("""select customer, ambiente, url_rest, environment, empresa, filial,
                       usuario, banco, sufixo, timeout_s, limite_linhas, ativo,
                       coalesce(ssl_verificar, true) as ssl_verificar,
                       (senha_enc is not null) as tem_senha,
                       (token_enc is not null) as tem_token,
                       ultimo_teste_em, ultimo_teste_ok, ultimo_teste_msg,
                       observacao, updated_by, updated_at
                  from cockpit.protheus_ambientes
                 where customer=%s order by ambiente""", (customer,))
    coletas = q("""select ambiente, tipo, iniciado_em, duracao_ms, ok, linhas,
                          medicoes, data_medicao, erro, disparado_por
                     from cockpit.protheus_coletas
                    where customer=%s order by iniciado_em desc limit 12""", (customer,))
    return _json({"ok": True, "customer": customer, "ambientes": rows,
                  "coletas": coletas, "chave_ok": bool(os.environ.get("PROTHEUS_CRED_KEY"))})


@app.post("/api/protheus/<customer>/<ambiente>")
def api_protheus_salvar(customer, ambiente):
    """Salva a conexão. Senha e token só são regravados quando vêm preenchidos:
    reabrir a tela e salvar outro campo não pode apagar a credencial."""
    erro, amb = _guarda_ambiente(customer, ambiente, escrita=True)
    if erro:
        return erro
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")

    b = request.get_json(silent=True) or {}
    url = (b.get("url_rest") or "").strip().rstrip("/")
    if not re.match(r"^https?://", url, re.I):
        return _err(400, "URL REST inválida — informe http://host:porta/tscmonit.")
    banco = (b.get("banco") or "").strip().lower() or None
    if banco and banco not in BANCOS:
        return _err(400, f"Banco inválido — use um de: {', '.join(BANCOS)}.")

    senha_enc = token_enc = None
    try:
        if (s := (b.get("senha") or "").strip()):
            senha_enc = _cifra(s, _aad(customer, amb))
        if (t := (b.get("token") or "").strip()):
            token_enc = _cifra(t, _aad(customer, amb))
    except RuntimeError as e:
        return _err(500, str(e))

    execute("""insert into cockpit.protheus_ambientes
                 (customer, ambiente, url_rest, environment, empresa, filial, usuario,
                  senha_enc, token_enc, banco, sufixo, timeout_s, limite_linhas,
                  ativo, ssl_verificar, observacao, updated_by, updated_at)
               values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s, now())
               on conflict (customer, ambiente) do update set
                 url_rest=excluded.url_rest, environment=excluded.environment,
                 empresa=excluded.empresa, filial=excluded.filial,
                 usuario=excluded.usuario,
                 senha_enc=coalesce(excluded.senha_enc, cockpit.protheus_ambientes.senha_enc),
                 token_enc=coalesce(excluded.token_enc, cockpit.protheus_ambientes.token_enc),
                 banco=excluded.banco, sufixo=excluded.sufixo,
                 timeout_s=excluded.timeout_s, limite_linhas=excluded.limite_linhas,
                 ativo=excluded.ativo, ssl_verificar=excluded.ssl_verificar,
                 observacao=excluded.observacao,
                 updated_by=excluded.updated_by, updated_at=now()""",
            (customer, amb, url, (b.get("environment") or "").strip() or None,
             (b.get("empresa") or "01").strip(), (b.get("filial") or "01").strip(),
             (b.get("usuario") or "").strip() or None, senha_enc, token_enc, banco,
             (b.get("sufixo") or "").strip()[:10] or None,
             int(b.get("timeout_s") or 45), int(b.get("limite_linhas") or 50000),
             bool(b.get("ativo", True)), bool(b.get("ssl_verificar", True)),
             (b.get("observacao") or "").strip() or None,
             current_user()))
    return _json({"ok": True, "customer": customer, "ambiente": amb})


@app.delete("/api/protheus/<customer>/<ambiente>")
def api_protheus_remover(customer, ambiente):
    erro, amb = _guarda_ambiente(customer, ambiente, escrita=True)
    if erro:
        return erro
    execute("delete from cockpit.protheus_ambientes where customer=%s and ambiente=%s",
            (customer, amb))
    return _json({"ok": True, "customer": customer, "ambiente": amb, "removido": True})


def _json_da_resposta(r):
    """Le o corpo como JSON SEM exigir o content-type. O AppServer nem sempre
    carimba o cabecalho que o fonte pede, e recusar por causa disso transformava
    uma resposta boa em "HTTP 200 nao reconhecido"."""
    try:
        return r.json()
    except ValueError:
        return {}


def _trecho(r):
    """Primeiros caracteres do corpo — sem isso o erro nao diz NADA sobre o que
    a base respondeu, e o diagnostico vira adivinhacao."""
    txt = " ".join((r.text or "").split())
    return (txt[:160] + "...") if len(txt) > 160 else (txt or "(corpo vazio)")


def _verifica_tls(cfg):
    """Valida o certificado do servidor? Padrao SIM.

    Base TCloud publicada por IP tem certificado emitido para
    *.protheus.cloudtotvs.com.br: o nome nao bate com o IP e o `requests`
    recusa com SSLError antes de qualquer coisa do Protheus (o navegador
    tambem recusaria; ele so ja tem a excecao guardada). O certo e cadastrar
    a URL pelo HOSTNAME; quando o cliente so fornece o IP, esta chave desliga
    a validacao para aquele ambiente — e so para ele."""
    return cfg.get("ssl_verificar") is not False


def _motivo_rede(e):
    """Uma frase que diz o que fazer. 'SSLError' sozinho nao separa
    certificado invalido de porta errada, e foi isso que custou tempo na
    Acosul em 22/09/2026."""
    nome = type(e).__name__
    txt = " ".join(str(e).split())[:200]
    if isinstance(e, requests.exceptions.SSLError):
        dica = ("Certificado recusado. Se a URL usa IP, o certificado do TCloud e "
                "emitido para o hostname (*.protheus.cloudtotvs.com.br) e nunca vai "
                "casar: use o hostname, ou marque 'Nao validar certificado' aqui. "
                "Se o servico for HTTP puro, troque https:// por http://.")
    elif isinstance(e, requests.exceptions.ConnectTimeout):
        dica = "Conectou em nada dentro do tempo: confira porta, firewall e VPN."
    else:
        dica = "Confira URL, porta, VPN e o [HTTPURI] do appserver.ini."
    return f"Sem resposta: {nome}. {dica} Detalhe: {txt}"


def _sessao_protheus(cfg, customer, amb, tenant=True):
    """Monta (headers, auth, timeout) a partir do cadastro.

    `tenantId` e obrigatorio quando a URI REST do cliente esta com
    `PrepareIn=All` (caso do Olim): sem ele o AppServer nao sabe em que
    empresa/filial preparar o ambiente. Mandamos TAMBEM empresa/filial como
    query params porque instalacoes mais antigas so entendem daquele jeito —
    quem nao usa um dos dois simplesmente ignora.
    Basic auth vai sempre que houver usuario cadastrado: com `SECURITY=1` no
    [HTTPREST] a chamada sem credencial leva 401 mesmo com o token correto.

    `tenant=False` derruba o tenantId de proposito: e assim que descobrimos a
    empresa/filial REAIS da base. Sem o header o AppServer prepara o ambiente
    pelo PrepareIn e o /ping devolve o cEmpAnt/cFilAnt que valem ali. A
    mascara da filial muda de cliente para cliente (01, 0101, 010101) e
    digitar essa mascara a mao e o erro que essa segunda tentativa evita."""
    tok = _decifra(cfg.get("token_enc"), _aad(customer, amb))
    sen = _decifra(cfg.get("senha_enc"), _aad(customer, amb))
    heads = {"Accept": "application/json"}
    if tenant:
        heads["tenantId"] = f"{cfg.get('empresa') or '01'},{cfg.get('filial') or '01'}"
    if tok:
        heads["X-TSC-Token"] = tok
    auth = (cfg.get("usuario"), sen) if cfg.get("usuario") else None
    return heads, auth, tok, min(int(cfg.get("timeout_s") or 45), TIMEOUT_TETO)


def _ping_protheus(cfg, customer, amb, tenant=True):
    """GET /tscmonit/ping. Com `tenant=False` sai sem tenantId e sem os query
    params de empresa/filial — a chamada de DESCOBERTA."""
    heads, auth, _tok, tmo = _sessao_protheus(cfg, customer, amb, tenant)
    params = ({"empresa": cfg.get("empresa") or "01",
               "filial": cfg.get("filial") or "01"} if tenant else None)
    r = requests.get(f"{cfg['url_rest']}/ping", headers=heads, auth=auth,
                     params=params, timeout=min(tmo, 30), verify=_verifica_tls(cfg))
    return r, _json_da_resposta(r)


def _sugestao_tenant(cfg, dados):
    """Empresa/filial que a BASE informou, quando diferem do cadastro. Devolve
    sempre os dois campos: a tela aplica o par inteiro, nao metade dele."""
    emp, fil = str(dados.get("empresa") or "").strip(), str(dados.get("filial") or "").strip()
    if not emp and not fil:
        return None
    if emp == (cfg.get("empresa") or "").strip() and fil == (cfg.get("filial") or "").strip():
        return None
    return {"empresa": emp or (cfg.get("empresa") or ""),
            "filial": fil or (cfg.get("filial") or "")}


@app.post("/api/protheus/<customer>/<ambiente>/testar")
def api_protheus_testar(customer, ambiente):
    """GET /tscmonit/ping. Além de "respondeu", confere se o environment e o
    banco são os que foram cadastrados — divergência aqui é aviso, não erro:
    quem decide se a URL está apontando para a base certa é o consultor."""
    erro, amb = _guarda_ambiente(customer, ambiente)
    if erro:
        return erro
    cfg = _cfg_ambiente(customer, amb)
    if not cfg:
        return _err(404, f"Ambiente {amb} ainda não cadastrado para {customer}.")

    ini = time.time()
    try:
        r, dados = _ping_protheus(cfg, customer, amb)
    except requests.RequestException as e:
        return _fecha_teste(customer, amb, ini, False, None, _motivo_rede(e))

    if r.status_code >= 400 or not dados.get("ok"):
        # Segunda tentativa SEM tenantId: se a base responder assim, ela mesma
        # diz a empresa/filial que valem ali. E o caso do 403 "Usuario sem
        # acesso a empresa/filial", que quase sempre e mascara de filial
        # errada — 01 onde a base usa 010101.
        sug = None
        try:
            r2, d2 = _ping_protheus(cfg, customer, amb, tenant=False)
            if r2.status_code < 400 and d2.get("ok"):
                sug = _sugestao_tenant(cfg, d2)
        except requests.RequestException:
            pass
        msg = (dados.get("erro") or
               f"HTTP {r.status_code} na rota /ping, mas a resposta nao é o JSON do "
               f"TSCMONITREST. A base devolveu: {_trecho(r)}")
        if r.status_code == 404:
            # 404 = o AppServer respondeu, mas a rota nao existe nele. Quase
            # sempre e fonte nao compilado naquele RPO ou AppServer nao
            # reiniciado (as anotacoes so entram na subida do servico).
            msg = (f"HTTP 404: o servidor respondeu, mas a rota /tscmonit/ping nao existe "
                   f"nele. Confira se o TSCMONITREST.tlpp esta compilado NESTE RPO e se o "
                   f"AppServer do REST foi reiniciado depois (as rotas anotadas so entram na "
                   f"subida do servico); e se o prefixo do [HTTPURI] e o mesmo da URL.")
        if sug:
            msg = (f"HTTP {r.status_code} com empresa/filial "
                   f"{cfg.get('empresa')}/{cfg.get('filial')}. Sem o tenantId a base "
                   f"respondeu como {sug['empresa']}/{sug['filial']} — use esses valores.")
        return _fecha_teste(customer, amb, ini, False, r.status_code, msg, sugestao=sug)

    avisos = []
    if cfg.get("environment") and dados.get("environment") and \
            cfg["environment"].strip().lower() != str(dados["environment"]).strip().lower():
        avisos.append(f"environment cadastrado ({cfg['environment']}) ≠ do servidor "
                      f"({dados['environment']})")
    if cfg.get("banco") and dados.get("banco") and \
            cfg["banco"] != str(dados["banco"]).strip().lower():
        avisos.append(f"banco cadastrado ({cfg['banco']}) ≠ do servidor ({dados['banco']})")
    if (sug := _sugestao_tenant(cfg, dados)):
        avisos.append(f"empresa/filial cadastrada ({cfg.get('empresa')}/{cfg.get('filial')}) "
                      f"≠ do servidor ({sug['empresa']}/{sug['filial']})")
    if not dados.get("token_ok"):
        avisos.append("token recusado pela base (confira MV_TSCTOKN) — a coleta vai falhar")

    msg = "Conexão OK." if not avisos else "Conectou, mas: " + "; ".join(avisos)
    return _fecha_teste(customer, amb, ini, not avisos, r.status_code, msg, dados, sug)


def _fecha_teste(customer, amb, ini, ok, status, msg, dados=None, sugestao=None):
    execute("""update cockpit.protheus_ambientes
                  set ultimo_teste_em=now(), ultimo_teste_ok=%s, ultimo_teste_msg=%s
                where customer=%s and ambiente=%s""", (ok, msg[:500], customer, amb))
    _log_coleta(customer, amb, "ping", ini, ok, http_status=status,
                erro=None if ok else msg)
    return _json({"ok": ok, "customer": customer, "ambiente": amb, "mensagem": msg,
                  "http_status": status, "servidor": dados or {},
                  "sugestao": sugestao}, 200)


_RE_COMENT_INI = re.compile(r"^\s*(?:--[^\n]*(?:\n|$)|/\*.*?\*/)", re.S)


def _sql_execucao(texto):
    """Deixa o script como o banco aceita receber: sem cabeçalho de comentário e
    sem o ';' do fim.

    Comentário no MEIO da consulta o banco aceita. No INÍCIO, não: o texto chega
    ao TCGenQry sem começar por SELECT/WITH, o driver não abre cursor e o erro
    volta como "TC_GetError - NO CONNECTION" — mensagem sem nenhuma relação com a
    causa. Foi ela que custou a madrugada de 06→07/09/2026, porque o mesmo script
    sem o banner respondia 200 no Postman. O ';' final é normal num console SQL e
    o Oracle recusa via TCGenQry (ORA-00911).

    A rota TSCMONITREST v1.3 também corta os dois. Aqui é de propósito: a coleta
    não pode depender de qual versão do fonte está compilada na base do cliente,
    e o script salvo continua guardando o banner que documenta a consulta.
    """
    sql = texto or ""
    while True:
        novo = _RE_COMENT_INI.sub("", sql, count=1)
        if novo == sql:
            break
        sql = novo
    return sql.strip().rstrip(";").strip()


@app.post("/api/protheus/<customer>/<ambiente>/coletar")
def api_protheus_coletar(customer, ambiente):
    """Executa na base do cliente o script salvo em cockpit.monitcad_scripts e
    grava a medição — o mesmo destino do arquivo subido à mão.

        POST /api/protheus/TFEHXQ00/teste/coletar?tipo=cadastros
    """
    erro, amb = _guarda_ambiente(customer, ambiente, escrita=True)
    if erro:
        return erro
    tipo = _tipo_script()
    cfg = _cfg_ambiente(customer, amb)
    if not cfg:
        return _err(404, f"Ambiente {amb} ainda não cadastrado para {customer}.")
    if not cfg.get("ativo"):
        return _err(409, f"Ambiente {amb} está inativo — reative antes de coletar.")
    script = q("""select sql, dialeto, sufixo from cockpit.monitcad_scripts
                   where customer=%s and tipo=%s""", (customer, tipo), one=True)
    if not script or not (script.get("sql") or "").strip():
        return _err(409, f"Nenhum script de {tipo} salvo para {customer}. "
                         f"Salve no modal 'Script SQL' antes de coletar.")

    ini = time.time()
    try:
        heads, auth, tok, tmo = _sessao_protheus(cfg, customer, amb)
    except RuntimeError as e:
        return _err(500, str(e))
    if not tok:
        return _err(409, "Sem token cadastrado para este ambiente — a REST da base recusa "
                         "a chamada sem o X-TSC-Token.")

    sql_exec = _sql_execucao(script["sql"])
    if not sql_exec:
        return _err(409, f"O script de {tipo} salvo para {customer} só tem comentários.")

    corpo = {"token": tok, "tipo": tipo, "sql": sql_exec,
             "limite": int(cfg.get("limite_linhas") or 50000)}
    try:
        r = requests.post(f"{cfg['url_rest']}/query", json=corpo, headers=heads, auth=auth,
                          params={"empresa": cfg.get("empresa") or "01",
                                  "filial": cfg.get("filial") or "01"}, timeout=tmo,
                          verify=_verifica_tls(cfg))
        dados = _json_da_resposta(r)
    except requests.RequestException as e:
        msg = (f"Falha ao chamar a base: {_motivo_rede(e)} "
               f"Timeout do ambiente = {tmo}s (a função da Vercel morre em 60s).")
        _log_coleta(customer, amb, tipo, ini, False, erro=msg)
        return _err(502, msg)

    if r.status_code >= 400 or not dados.get("ok"):
        msg = dados.get("erro") or (f"HTTP {r.status_code} na rota /query, mas a resposta "
                                    f"nao é o JSON do TSCMONITREST. A base devolveu: {_trecho(r)}")
        _log_coleta(customer, amb, tipo, ini, False, http_status=r.status_code, erro=msg)
        return _err(502, f"A base recusou a consulta: {msg}")

    # Truncado NÃO grava. Uma medição cortada entra sem erro nenhum: os números
    # chegam bonitos e errados e ninguém percebe olhando o painel. Melhor recusar
    # e pedir mais limite do que sujar o histórico com um número que parece bom.
    if dados.get("truncado"):
        msg = (f"A consulta bateu no limite de {corpo['limite']} linha(s) e voltou "
               f"cortada — nada foi gravado. Aumente o limite de linhas do ambiente "
               f"ou reduza o alcance do script antes de coletar de novo.")
        _log_coleta(customer, amb, tipo, ini, False, http_status=r.status_code,
                    linhas=dados.get("linhas"), erro=msg)
        return _err(409, msg)

    csv_txt = dados.get("csv") or ""
    if not csv_txt.strip():
        msg = "A consulta rodou e voltou vazia — nenhuma linha para gravar."
        _log_coleta(customer, amb, tipo, ini, False, http_status=r.status_code,
                    linhas=0, erro=msg)
        return _err(409, msg)

    try:
        if tipo == "sm0":
            body = _csv_sm0(csv_txt)
        elif tipo == "empresas":
            body = _csv_empresas(csv_txt)
        elif tipo == "movimentos":
            body = _csv_movimentos(csv_txt)
        else:
            body = _csv_para_body(csv_txt)
    except Exception as e:
        msg = f"CSV devolvido pela base é inválido: {e}"
        _log_coleta(customer, amb, tipo, ini, False, http_status=r.status_code,
                    linhas=dados.get("linhas"), bytes=len(csv_txt), erro=msg)
        return _err(422, msg)

    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")

    if tipo == "sm0":
        res = _gravar_sm0(customer, amb, body)
        res["ultima_medicao"] = None
        itens = res.get("filiais")
    elif tipo == "empresas":
        res = _gravar_empresas(customer, amb, body, origem="coleta")
        itens = res.get("itens")
    elif tipo == "movimentos":
        res = _gravar_movimentos(customer, amb, body, origem="coleta")
        itens = res.get("dimensoes")
    else:
        res = _gravar_cadastros(customer, amb, body, origem="coleta")
        itens = res.get("tabelas")

    _log_coleta(customer, amb, tipo, ini, True, http_status=r.status_code,
                linhas=dados.get("linhas"), medicoes=res.get("medicoes"), itens=itens,
                bytes=len(csv_txt), data_medicao=res.get("ultima_medicao"))
    return _json({"ok": True, "customer": customer, "ambiente": amb, "tipo": tipo,
                  "environment": dados.get("environment"), "banco": dados.get("banco"),
                  "linhas": dados.get("linhas"), "truncado": dados.get("truncado"),
                  "tempo_base_ms": dados.get("tempo_ms"), **res})




CEN_CAMPOS = ("analise", "dim1", "dim2", "dim3", "dim4", "descricao",
              "esperado", "etapa", "responsavel")


def _cenarios_do_txt(texto):
    """Lê o CENARIOS_MONITOR.TXT — o mesmo arquivo que vai para a \\system\\ do
    Protheus. Uma fonte, dois consumidores: o job TLPP e este painel."""
    fora = []
    for n, linha in enumerate(texto.splitlines(), 1):
        linha = linha.strip()
        if not linha or linha.startswith("#"):
            continue
        p = [c.strip() for c in linha.split(";")]
        if len(p) < 5:
            raise ValueError(f"Linha {n}: esperado ANALISE;DIM1;DIM2;DIM3;DIM4;"
                             f"DESCRICAO;ESPERADO;ETAPA;RESPONSAVEL — veio {len(p)} campo(s).")
        p += [""] * (9 - len(p))
        fora.append({
            "analise": p[0].upper(),
            "dim1": p[1] or "*", "dim2": p[2] or "*",
            "dim3": p[3] or "*", "dim4": p[4] or "*",
            "descricao": p[5] or None,
            "esperado": (p[6] or "S").strip().upper() not in ("N", "NAO", "NÃO", "0", "FALSE"),
            "etapa": p[7] or None, "responsavel": p[8] or None,
        })
    if not fora:
        raise ValueError("Nenhum cenário no arquivo (só comentários?).")
    return fora


@app.get("/api/monitmov/<customer>/cenarios")
def api_monitmov_cenarios(customer):
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    return _json({"ok": True, "customer": customer,
                  "cenarios": q("""select id, analise, dim1, dim2, dim3, dim4, descricao,
                                          esperado, etapa, responsavel, updated_by, updated_at
                                     from cockpit.monitmov_cenarios where customer=%s
                                    order by analise, descricao""", (customer,))})


@app.post("/api/monitmov/<customer>/cenarios")
def api_monitmov_cenarios_salvar(customer):
    """Substitui a lista inteira. O arquivo é a verdade — mesma regra do
    TABELAS_MONITOR.TXT: editar em dois lugares é como as listas divergem."""
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de salvar cenários.")
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")
    b = request.get_json(silent=True) or {}
    if isinstance(b.get("texto"), str):
        try:
            cens = _cenarios_do_txt(b["texto"])
        except Exception as e:
            return _err(400, str(e))
    elif isinstance(b.get("cenarios"), list):
        cens = [{k: c.get(k) for k in CEN_CAMPOS} for c in b["cenarios"]]
    else:
        return _err(400, "Envie o conteúdo do CENARIOS_MONITOR.TXT em 'texto' "
                         "ou a lista em 'cenarios'.")
    execute("delete from cockpit.monitmov_cenarios where customer=%s", (customer,))
    gravados = 0
    for c in cens:
        if not (c.get("analise") or "").strip():
            continue
        execute("""insert into cockpit.monitmov_cenarios
                     (customer, analise, dim1, dim2, dim3, dim4, descricao,
                      esperado, etapa, responsavel, updated_by, updated_at)
                   values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s, now())
                   on conflict (customer, analise, dim1, dim2, dim3, dim4)
                   do update set descricao=excluded.descricao,
                                 esperado=excluded.esperado, etapa=excluded.etapa,
                                 responsavel=excluded.responsavel,
                                 updated_by=excluded.updated_by, updated_at=now()""",
                (customer, (c["analise"] or "").strip().upper(),
                 (c.get("dim1") or "*").strip(), (c.get("dim2") or "*").strip(),
                 (c.get("dim3") or "*").strip(), (c.get("dim4") or "*").strip(),
                 c.get("descricao"), bool(c.get("esperado", True)),
                 c.get("etapa"), c.get("responsavel"), current_user()))
        gravados += 1
    return _json({"ok": True, "customer": customer, "cenarios": gravados})


@app.delete("/api/monitmov/<customer>/cenarios")
def api_monitmov_cenarios_remover(customer):
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de remover cenários.")
    execute("delete from cockpit.monitmov_cenarios where customer=%s", (customer,))
    return _json({"ok": True, "customer": customer, "removidos": True})


# ── GAPs (aba "GAPs") ───────────────────────────────────────────────────────
# GAP = ticket com a tag EXATA "GAP" em cockpit.ticket_tags. Existe também o
# campo tipo_atividade='GAP', que quase coincide (630 nos dois, 15 só em cada) —
# a tag é a fonte da verdade aqui, por decisão de 11/08/2026.
DECISOES = ("aprovar", "segunda_fase", "contorno", "entendimento_projeto",
            "recusar", "pendente")

SQL_GAPS = """
  select t.uuid_ticket, t.raw->>'id' as id, t.titulo, t.status_tasks,
         t.status_temporario, t.etapa_gap, t.classificacao_gap, t.produto,
         t.competencia, t.projeto, t.prioridade, t.time_estimate, t.aging_dias,
         t.due_date, t.atrasado, t.bloqueado, t.user_assigned,
         u.nome as responsavel_nome, t.assigned_customer, t.observador,
         t.ult_ocorr_texto, t.ult_ocorr_data, t.ult_ocorr_autor, t.updated_at,
         d.decisao, d.estimativa as decisao_estimativa, d.observacao as decisao_obs,
         d.decided_by, d.decided_at,
         (a.uuid_ticket is not null) as tem_alinhamento,
         coalesce((select array_agg(g2.raw_tag order by g2.raw_tag)
                     from cockpit.ticket_tags g2
                    where g2.uuid_ticket = t.uuid_ticket
                      and g2.raw_tag <> 'GAP'), '{}') as tags
    from cockpit.tickets t
    join cockpit.ticket_tags g
      on g.uuid_ticket = t.uuid_ticket and g.raw_tag = 'GAP'
    left join cockpit.usuarios u on u.codigo = t.user_assigned
    left join cockpit.decisoes d on d.uuid_ticket = t.uuid_ticket
    left join cockpit.gap_alinhamentos a on a.uuid_ticket = t.uuid_ticket
   where t.customer = %s
   order by t.time_estimate desc nulls last, t.titulo
"""


def _customer_do_ticket(uuid):
    r = q("select customer from cockpit.tickets where uuid_ticket=%s", (uuid,), one=True)
    return r["customer"] if r else None


# ── TAREFAS — a mesma leitura dos GAPs, sem a tag GAP obrigatória ───────────
# A aba GAPs existe para decidir escopo, então parte de uma tag só. Esta aqui é
# de garimpo: o recorte é o que o usuário marcar. Duas diferenças que importam:
#   1. sem o `join ... raw_tag = 'GAP'`, entra TODO ticket do cliente;
#   2. o array de tags NÃO exclui 'GAP' — aqui ela é uma tag como outra qualquer
#      e precisa aparecer para poder ser marcada no filtro.
# `eh_gap` vem junto para a lista poder marcar quais já são GAP sem uma segunda
# consulta.
SQL_TAREFAS = """
  select t.uuid_ticket, t.raw->>'id' as id, t.titulo, t.status_tasks,
         t.status_temporario, t.etapa_gap, t.classificacao_gap, t.produto,
         t.competencia, t.projeto, t.prioridade, t.time_estimate, t.aging_dias,
         t.due_date, t.atrasado, t.bloqueado, t.user_assigned,
         u.nome as responsavel_nome, t.assigned_customer, t.observador,
         t.ult_ocorr_texto, t.ult_ocorr_data, t.ult_ocorr_autor, t.updated_at,
         d.decisao, d.estimativa as decisao_estimativa, d.observacao as decisao_obs,
         d.decided_by, d.decided_at,
         (a.uuid_ticket is not null) as tem_alinhamento,
         exists (select 1 from cockpit.ticket_tags g3
                  where g3.uuid_ticket = t.uuid_ticket and g3.raw_tag = 'GAP') as eh_gap,
         coalesce((select array_agg(g2.raw_tag order by g2.raw_tag)
                     from cockpit.ticket_tags g2
                    where g2.uuid_ticket = t.uuid_ticket), '{}') as tags
    from cockpit.tickets t
    left join cockpit.usuarios u on u.codigo = t.user_assigned
    left join cockpit.decisoes d on d.uuid_ticket = t.uuid_ticket
    left join cockpit.gap_alinhamentos a on a.uuid_ticket = t.uuid_ticket
   where t.customer = %s
   order by t.time_estimate desc nulls last, t.titulo
"""


@app.get("/api/tarefas/<customer>")
def api_tarefas(customer):
    """Todos os tickets do cliente. O detalhe continua em /api/gaps/ticket/<uuid>:
    aquela rota autoriza pelo cliente DONO do ticket e nunca exigiu a tag GAP,
    então serve para qualquer tarefa sem rota nova."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "tarefas")):
        return d
    itens = q(SQL_TAREFAS, (customer,))
    horas = sum(float(t["time_estimate"] or 0) for t in itens)
    return _json({"ok": True, "customer": customer, "total": len(itens),
                  "horas_estimadas": horas,
                  "gaps": sum(1 for t in itens if t["eh_gap"]),
                  "tarefas": itens})


def _guarda_ticket(uuid):
    """Autoriza pelo cliente DONO do ticket — nunca pelo que o front mandou."""
    cust = _customer_do_ticket(uuid)
    if not cust:
        return _err(404, "GAP não encontrado."), None
    return deny_aba(cust, "gaps", "tarefas"), cust


@app.get("/api/gaps/<customer>")
def api_gaps(customer):
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "gaps")):
        return d
    gaps = q(SQL_GAPS, (customer,))
    horas = sum(float(g["time_estimate"] or 0) for g in gaps)
    return _json({"ok": True, "customer": customer, "total": len(gaps),
                  "horas_estimadas": horas, "gaps": gaps})


@app.get("/api/gaps/ticket/<uuid>")
def api_gap_detalhe(uuid):
    if (r := require_auth()):
        return r
    negado, _ = _guarda_ticket(uuid)
    if negado:
        return negado
    t = q("""select t.uuid_ticket, t.raw->>'id' as id, t.titulo, t.descricao, t.cliente,
                    t.status_tasks, t.etapa_gap, t.classificacao_gap, t.produto,
                    t.competencia, t.time_estimate, t.due_date, t.aging_dias,
                    t.user_assigned, u.nome as responsavel_nome, t.assigned_customer,
                    t.observador, t.updated_at, t.synced_at
               from cockpit.tickets t
               left join cockpit.usuarios u on u.codigo = t.user_assigned
              where t.uuid_ticket = %s""", (uuid,), one=True)
    tags = q("""select raw_tag, dimensao, valor from cockpit.ticket_tags
                 where uuid_ticket=%s order by dimensao_idx nulls last, raw_tag""", (uuid,))
    ocorr = q("""select uuid_history, tipo, details, autor, occurred_at
                   from cockpit.ocorrencias where uuid_ticket=%s
                  order by occurred_at desc nulls last limit 30""", (uuid,))
    dec = q("select * from cockpit.decisoes where uuid_ticket=%s", (uuid,), one=True)
    ali = q("select * from cockpit.gap_alinhamentos where uuid_ticket=%s", (uuid,), one=True)
    return _json({"ok": True, "ticket": t, "tags": tags, "ocorrencias": ocorr,
                  "decisao": dec, "alinhamento": ali})


@app.post("/api/gaps/ticket/<uuid>/decisao")
def api_gap_decisao(uuid):
    """Grava a decisão LOCAL (cockpit.decisoes). Não toca no Tasks SC."""
    if (r := require_interno()):
        return r
    negado, _ = _guarda_ticket(uuid)
    if negado:
        return negado
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de decidir.")
    b = request.get_json(silent=True) or {}
    decisao = (b.get("decisao") or "pendente").strip()
    if decisao not in DECISOES:
        return _err(400, f"Decisão inválida: {decisao}")
    est = b.get("estimativa")
    est = float(est) if est not in (None, "") else None
    execute("""insert into cockpit.decisoes
                 (uuid_ticket, decisao, estimativa, observacao, classe,
                  decided_by, decided_at, updated_at)
               values (%s,%s,%s,%s,%s,%s, now(), now())
               on conflict (uuid_ticket) do update set
                 decisao=excluded.decisao, estimativa=excluded.estimativa,
                 observacao=excluded.observacao, classe=excluded.classe,
                 decided_by=excluded.decided_by, decided_at=now(), updated_at=now()""",
            (uuid, decisao, est, b.get("observacao"), b.get("classe"), current_user()))
    return _json({"ok": True, "uuid_ticket": uuid, "decisao": decisao})


@app.post("/api/gaps/ticket/<uuid>/alinhamento")
def api_gap_alinhamento(uuid):
    """Alinhamento comercial do GAP. ATENÇÃO: argumentacao_interna é INTERNA —
    não expor ao cliente em e-mail nem em painel compartilhado."""
    if (r := require_interno()):
        return r
    negado, cust = _guarda_ticket(uuid)
    if negado:
        return negado
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de gravar.")
    b = request.get_json(silent=True) or {}
    tid = q("select raw->>'id' as id from cockpit.tickets where uuid_ticket=%s",
            (uuid,), one=True)
    execute("""insert into cockpit.gap_alinhamentos
                 (uuid_ticket, task_id, customer, questionamento_cliente,
                  argumentacao_interna, alinhamento_reuniao, retorno_cliente,
                  created_by, updated_by, updated_at)
               values (%s,%s,%s,%s,%s,%s,%s,%s,%s, now())
               on conflict (uuid_ticket) do update set
                 questionamento_cliente=excluded.questionamento_cliente,
                 argumentacao_interna=excluded.argumentacao_interna,
                 alinhamento_reuniao=excluded.alinhamento_reuniao,
                 retorno_cliente=excluded.retorno_cliente,
                 updated_by=excluded.updated_by, updated_at=now()""",
            (uuid, (tid or {}).get("id"), cust, b.get("questionamento_cliente"),
             b.get("argumentacao_interna"), b.get("alinhamento_reuniao"),
             b.get("retorno_cliente"), current_user(), current_user()))
    return _json({"ok": True, "uuid_ticket": uuid})


# ── Gmail: rascunho por IMAP APPEND ─────────────────────────────────────────
# A App Password fica cifrada (Fernet) em cockpit.gmail_credenciais, com a chave
# derivada do SESSION_SECRET da Vercel — trocar o SESSION_SECRET invalida todas
# as credenciais salvas, e cada usuário só precisa salvar de novo pela própria
# tela. A senha NUNCA volta para o front.
IMAP_HOST = "imap.gmail.com"


def _fernet():
    from cryptography.fernet import Fernet
    if SESSION_SECRET == "dev-insecure-secret":
        raise RuntimeError("SESSION_SECRET não configurado — não dá para cifrar a senha.")
    chave = base64.urlsafe_b64encode(hashlib.sha256(SESSION_SECRET.encode()).digest())
    return Fernet(chave)


def _imap_login(gmail, senha):
    m = imaplib.IMAP4_SSL(IMAP_HOST, timeout=25)
    m.login(gmail, senha)
    return m


def _pasta_rascunhos(m):
    """O nome da pasta muda com o idioma da conta ([Gmail]/Drafts, /Rascunhos…).
    Acha pela flag \\Drafts, que é estável."""
    ok, linhas = m.list()
    if ok == "OK":
        for l in linhas:
            txt = l.decode("utf-8", "replace") if isinstance(l, bytes) else str(l)
            if "\\Drafts" in txt:
                return '"' + txt.split(' "/" ')[-1].strip().strip('"') + '"'
    return '"[Gmail]/Drafts"'


def _cred_gmail(usuario):
    row = q("select gmail_email, senha_cif from cockpit.gmail_credenciais where usuario=%s",
            (usuario,), one=True)
    if not row:
        return None, None
    try:
        return row["gmail_email"], _fernet().decrypt(row["senha_cif"].encode()).decode()
    except Exception:
        return row["gmail_email"], None      # cifrada com outro SESSION_SECRET


@app.get("/api/gmail/cred")
def api_gmail_cred_get():
    if (r := require_auth()):
        return r
    gmail, senha = _cred_gmail(current_user())
    return _json({"ok": True, "configurado": bool(gmail and senha), "gmail_email": gmail,
                  "precisa_resalvar": bool(gmail and not senha)})


@app.post("/api/gmail/cred")
def api_gmail_cred_set():
    """Valida a App Password fazendo login IMAP de verdade antes de gravar."""
    if (r := require_auth()):
        return r
    b = request.get_json(silent=True) or {}
    gmail = (b.get("gmail_email") or "").strip().lower()
    senha = (b.get("app_password") or "").replace(" ", "")
    if not gmail or not senha:
        return _err(400, "Informe o e-mail do Gmail e a App Password.")
    try:
        m = _imap_login(gmail, senha)
        m.logout()
    except Exception as e:
        return _err(400, f"O Gmail recusou a credencial: {e}")
    execute("""insert into cockpit.gmail_credenciais
                 (usuario, gmail_email, senha_cif, validado_em, updated_at)
               values (%s,%s,%s, now(), now())
               on conflict (usuario) do update set
                 gmail_email=excluded.gmail_email, senha_cif=excluded.senha_cif,
                 validado_em=now(), updated_at=now()""",
            (current_user(), gmail, _fernet().encrypt(senha.encode()).decode()))
    return _json({"ok": True, "gmail_email": gmail})


@app.delete("/api/gmail/cred")
def api_gmail_cred_del():
    if (r := require_auth()):
        return r
    execute("delete from cockpit.gmail_credenciais where usuario=%s", (current_user(),))
    return _json({"ok": True})


@app.post("/api/gmail/rascunho")
def api_gmail_rascunho():
    """Cria um rascunho HTML na conta do usuário logado (IMAP APPEND)."""
    if (r := require_interno()):
        return r
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de criar rascunhos.")
    b = request.get_json(silent=True) or {}
    assunto = (b.get("assunto") or "").strip()
    html = b.get("html") or ""
    para = (b.get("para") or "").strip()
    if not assunto or not html:
        return _err(400, "Informe assunto e corpo do e-mail.")

    gmail, senha = _cred_gmail(current_user())
    if not gmail:
        return _err(428, "Sem credencial do Gmail. Cadastre uma App Password.")
    if not senha:
        return _err(428, "A credencial do Gmail não pôde ser lida (SESSION_SECRET mudou). "
                         "Cadastre a App Password de novo.")

    msg = email.message.EmailMessage()
    msg["From"] = gmail
    if para:
        msg["To"] = para
    if cc := (b.get("cc") or "").strip():
        msg["Cc"] = cc
    msg["Subject"] = assunto
    msg.set_content("Este e-mail tem formatação HTML — abra num cliente compatível.")
    msg.add_alternative(html, subtype="html")

    try:
        m = _imap_login(gmail, senha)
        try:
            m.append(_pasta_rascunhos(m), "\\Draft", imaplib.Time2Internaldate(time.time()),
                     msg.as_bytes())
        finally:
            m.logout()
    except Exception as e:
        return _err(502, f"Falha ao criar o rascunho no Gmail: {e}")
    return _json({"ok": True, "gmail_email": gmail, "assunto": assunto, "para": para})


# ── Painéis do cliente (aba "Transição") ────────────────────────────────────
# O HTML fica em cockpit.paineis_cliente, NÃO no git: este repo é público e o
# painel carrega o plano de atividades, consultores e GAPs do cliente.
@app.get("/api/paineis")
def api_paineis():
    """Painéis disponíveis para os clientes que o usuário enxerga."""
    if (r := require_auth()):
        return r
    allowed = allowed_customers()
    rows = q("""select p.customer, p.slug, p.titulo, p.descricao, p.updated_at,
                       c.nome as cliente_nome
                  from cockpit.paineis_cliente p
                  join cockpit.clientes c on c.customer = p.customer
                 where p.ativo order by c.nome, p.slug""")
    out = [r for r in rows if allowed is None or r["customer"] in allowed]
    return _json({"ok": True, "paineis": out, "pode_subir": is_admin(current_user())})


@app.get("/painel/<customer>/<slug>")
def page_painel(customer, slug):
    """Serve o painel dentro do iframe da aba Transição — exige sessão e acesso
    ao cliente. É por isso que o HTML não pode morar em web/ (estático livre)."""
    if not current_user():
        return redirect("/login", 302)
    if (d := deny_aba(customer, "transicao")):
        return d
    row = q("select html from cockpit.paineis_cliente "
            "where customer=%s and slug=%s and ativo", (customer, slug), one=True)
    if not row:
        return _err(404, "Painel não encontrado.")
    return Response(row["html"], mimetype="text/html; charset=utf-8")


@app.post("/api/paineis/<customer>/<slug>")
def api_painel_upload(customer, slug):
    """Publica/atualiza o HTML do painel. Só admin; o corpo é o arquivo inteiro."""
    if (r := require_admin()):
        return r
    if not q("select 1 from cockpit.clientes where customer=%s", (customer,), one=True):
        return _err(409, f"Cliente {customer} não cadastrado em cockpit.clientes.")
    html = request.get_data(as_text=True) or ""
    if (f := request.files.get("arquivo")):
        html = f.read().decode("utf-8", errors="replace")
    if "<" not in html or len(html) < 200:
        return _err(400, "Envie o HTML completo do painel.")
    titulo = request.args.get("titulo") or slug
    execute("""insert into cockpit.paineis_cliente
                 (customer, slug, titulo, html, updated_by, updated_at)
               values (%s,%s,%s,%s,%s, now())
               on conflict (customer, slug) do update set
                 titulo=excluded.titulo, html=excluded.html,
                 updated_by=excluded.updated_by, updated_at=now(), ativo=true""",
            (customer, slug, titulo, html, current_user()))
    return _json({"ok": True, "customer": customer, "slug": slug, "bytes": len(html)})


# detalhe AO VIVO (cache curto por instância quente)
_cache: dict = {}
TTL = 90


def _cached(key, loader):
    e = _cache.get(key)
    now = time.time()
    if e and now - e["t"] < TTL:
        return e["v"]
    try:
        v = loader()
        _cache[key] = {"t": now, "v": v}
        return v
    except PCIUnavailable:
        if e:  # serve o cache vencido em vez de tela de erro
            v = e["v"]
            if isinstance(v, dict):
                return {**v, "_stale": True, "_stale_idade_min": int((now - e["t"]) / 60)}
            return v
        raise


@app.get("/api/projeto/<cli>/<cod>/<ver>/<loja>/<kind>")
def api_projeto_detalhe(cli, cod, ver, loja, kind):
    if (r := require_auth()):
        return r
    if (d := deny_customer(cli)):
        return d
    if kind not in ("mapa", "cronograma"):
        return _err(400, "Esperado mapa|cronograma.")
    params = {"CLIENTE": cli, "CODIGO": cod, "VERSAO": ver, "LOJA": loja or "undefined"}
    if kind == "mapa":
        params["t"] = int(time.time() * 1000)
    url = MAPA_URL if kind == "mapa" else CRONO_URL
    return _json(_cached(f"{kind}:{cli}:{cod}:{ver}:{loja}", lambda: pci_get(url, params)))


# ============================================================================
#  PROTÓTIPO — roteiro MIT045, ciclos, papéis e aproveitamento
# ============================================================================
# A PERGUNTA QUE ESTA ÁREA RESPONDE não é "qual o status do item", é
# **"o aproveitamento ANDOU do ciclo anterior para este?"**. Por isso o
# resultado NÃO mora no item: mora no par (item × ciclo). O mesmo roteiro é
# executado de novo a cada ciclo — interno, isolado, integrado — e comparar os
# ciclos é o produto final. Guardar o status no item apagaria a história.
#
# DOIS EIXOS DE RESPOSTA, e confundi-los fecha o protótipo sozinho:
#   status    → o que o CLIENTE respondeu (executado, dúvida, erro, n/a)
#   validacao → o que o CONSULTOR disse   (pendente, ok, reprovado)
# "Executei" não é "está certo". O APROVEITAMENTO conta somente `validacao=ok`:
# sem isso o cliente marca 40 itens como executados e o protótipo fecha em 100%
# sem ninguém da TOTVS ter olhado.
#
# TRÊS INDICADORES, cada um responde uma pergunta diferente:
#   EXECUÇÃO       = % respondido como executado (o cliente andou?)
#   APROVEITAMENTO = % validado com OK           (o trabalho presta?)
#   MATURAÇÃO      = média das notas 0–10        (o usuário se sente seguro?)
# A distância entre execução e aproveitamento é a FILA DE VALIDAÇÃO — é ela que
# denuncia consultor parado, não o número de itens executados.

PROTO_STATUS = ("nao_iniciado", "executado", "duvida", "erro", "nao_aplicavel")
PROTO_STATUS_LABEL = {
    "nao_iniciado": "Não iniciado",
    "executado": "Executado",
    "duvida": "Com dúvida",
    "erro": "Com erro",
    "nao_aplicavel": "Não aplicável",
}
PROTO_VALID = ("pendente", "ok", "reprovado")
PROTO_VALID_LABEL = {"pendente": "Aguardando validação", "ok": "OK do consultor",
                     "reprovado": "Reprovado"}
PROTO_TIPOS = ("interno", "isolado", "integrado")
PROTO_TIPO_LABEL = {"interno": "Interno (consultoria)", "isolado": "Isolado",
                    "integrado": "Integrado"}

# Papéis NO CLIENTE (não no sistema). O perfil de usuarios_login diz o que a
# pessoa é no sistema; o papel aqui diz o que ela é NAQUELE PROJETO — a mesma
# consultora é cp_totvs na Dígitro e consultor no Olim.
PROTO_PAPEIS = ("cp_totvs", "consultor", "cp_cliente", "usuario_chave")
PROTO_PAPEL_LABEL = {
    "cp_totvs": "CP TOTVS — coordena o projeto",
    "consultor": "Consultor TOTVS — valida os módulos dele",
    "cp_cliente": "CP do Cliente — distribui e acompanha",
    "usuario_chave": "Usuário-chave — executa o que lhe foi atribuído",
}
PROTO_PAPEIS_TOTVS = ("cp_totvs", "consultor")

# De-para do texto da planilha para o domínio. Chave normalizada por _norm_col.
PROTO_STATUS_DE = {
    "NAOINICIADO": "nao_iniciado", "": "nao_iniciado",
    "EXECUTADOCOMEXITO": "executado", "EXITO": "executado", "OK": "executado",
    "EXECUTADO": "executado",
    "EXECUTADOCOMRESSALVA": "duvida", "COMRESSALVA": "duvida", "RESSALVA": "duvida",
    "COMDUVIDA": "duvida", "DUVIDA": "duvida",
    "ERRONECESSARIOAJUSTE": "erro", "ERRO": "erro", "NECESSARIOAJUSTE": "erro",
    "COMERRO": "erro",
    "NAOAPLICAVEL": "nao_aplicavel", "NA": "nao_aplicavel",
}

# Cabeçalho do MIT045. Aceita com e sem acento porque a planilha viaja entre
# xlsx, Sheets e CSV exportado, e cada um mexe no acento de um jeito.
PROTO_COLS = {
    "ID": "ordem", "ITEM": "ordem", "SEQ": "ordem",
    "MODULO": "modulo", "PROCESSO": "processo", "SUBPROCESSO": "subprocesso",
    "DESCRICAO": "descricao", "ATIVIDADE": "descricao",
    "CONSULTOR": "consultor", "USUARIO": "usuario", "USUARIOCHAVE": "usuario",
    "DATAPLANEJADA": "data_planejada", "DATAPREVISTA": "data_planejada",
    "STATUS": "status", "SITUACAO": "status",
    "DATACONCLUSAO": "data_conclusao", "OCORRENCIA": "ocorrencia",
    "OBSERVACAO": "ocorrencia", "OBSERVACOES": "ocorrencia",
}


def _proto_data(v):
    """Datas do roteiro chegam em três dialetos: '1/5/2026' (Sheets), o date do
    openpyxl e o ISO. Nada de dateutil — a função tem que ser previsível."""
    if v is None:
        return None
    if hasattr(v, "isoformat"):                     # date/datetime do openpyxl
        try:
            return v.date().isoformat() if hasattr(v, "date") else v.isoformat()
        except Exception:
            return None
    s = str(v).strip()
    if not s:
        return None
    # ARMADILHA: a célula de data já chegou aqui como TEXTO ("2026-05-01 00:00:00"),
    # porque quem lê a planilha normaliza tudo para string antes. Sem cortar a
    # hora, _data_iso não casa e a Data do Protótipo entra NULA — o roteiro
    # importa 40 itens certinhos e sem data, e ninguém percebe.
    if (m := re.match(r"(\d{4}-\d{2}-\d{2})[T ]", s)):
        return m.group(1)
    if (iso := _data_iso(s)):
        return iso
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
    if m:
        d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        # dd/mm/aaaa é o padrão daqui; só inverte quando o primeiro campo não
        # pode ser dia. Adivinhar mais do que isso troca 05/01 por 01/05.
        if d > 12 and mth <= 12:
            pass
        elif mth > 12 and d <= 12:
            d, mth = mth, d
        try:
            return f"{y:04d}-{mth:02d}-{d:02d}"
        except Exception:
            return None
    return None


def _proto_linhas_xlsx(dados):
    """Lê a primeira aba do .xlsx. openpyxl é dependência declarada; se faltar,
    a mensagem tem que dizer o que fazer, não estourar ImportError na cara."""
    try:
        import openpyxl
    except ImportError:
        raise ValueError("Leitura de .xlsx indisponível no servidor (openpyxl). "
                         "Exporte o roteiro como CSV e suba de novo.")
    wb = openpyxl.load_workbook(io.BytesIO(dados), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _proto_linhas_csv(texto):
    amostra = texto[:4096]
    try:
        dial = csv.Sniffer().sniff(amostra, delimiters=",;\t|")
    except csv.Error:
        dial = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(texto), dial)]


def _proto_cel(v):
    return "" if v is None else str(v).strip()


def _proto_parse(linhas):
    """Devolve {meta:{projeto,data_prototipo}, itens:[...]}.

    ARMADILHA: o MIT045 tem uma COLUNA VAZIA à esquerda e três linhas de
    cabeçalho antes da tabela. Procurar a linha de cabeçalho pelo CONTEÚDO
    (tem 'descricao' e ('modulo' ou 'processo')) é o que faz o parser aguentar
    a planilha ganhar uma linha de logo amanhã.
    """
    idx_cabec, mapa = None, {}
    for i, linha in enumerate(linhas[:40]):
        m = {}
        for j, c in enumerate(linha):
            alvo = PROTO_COLS.get(_norm_col(_proto_cel(c)))
            if alvo and alvo not in m:
                m[alvo] = j
        if "descricao" in m and ("modulo" in m or "processo" in m):
            idx_cabec, mapa = i, m
            break
    if idx_cabec is None:
        raise ValueError("Não achei o cabeçalho do roteiro. Esperado uma linha com "
                         "DESCRIÇÃO e MÓDULO/PROCESSO (padrão MIT045).")

    # Metadados ficam ACIMA do cabeçalho, no formato rótulo | valor.
    meta = {"projeto": None, "data_prototipo": None}
    for linha in linhas[:idx_cabec]:
        celulas = [_proto_cel(c) for c in linha]
        for j, c in enumerate(celulas):
            rot = _norm_col(c)
            val = next((x for x in celulas[j + 1:] if x), "")
            if rot == "PROJETO" and val and not meta["projeto"]:
                meta["projeto"] = val
            if rot in ("DATAPROTOTIPO", "DATADOPROTOTIPO") and val:
                meta["data_prototipo"] = _proto_data(val)

    def pega(linha, campo):
        j = mapa.get(campo)
        return _proto_cel(linha[j]) if (j is not None and j < len(linha)) else ""

    itens, ordem_auto = [], 0
    for linha in linhas[idx_cabec + 1:]:
        if not any(_proto_cel(c) for c in linha):
            continue
        desc = pega(linha, "descricao")
        if not desc:
            continue                       # linha de nota/rodapé, não é item
        bruto = pega(linha, "ordem")
        try:
            ordem = int(float(bruto))
        except (TypeError, ValueError):
            ordem_auto += 1
            ordem = ordem_auto
        ordem_auto = max(ordem_auto, ordem)
        itens.append({
            "ordem": ordem,
            "modulo": pega(linha, "modulo") or None,
            "processo": pega(linha, "processo") or None,
            "subprocesso": pega(linha, "subprocesso") or None,
            "descricao": desc,
            "consultor": pega(linha, "consultor") or None,
            "usuario": pega(linha, "usuario") or None,
            "data_planejada": _proto_data(pega(linha, "data_planejada")),
            # o status que JÁ estava na planilha: não vira item, mas serve para
            # semear o primeiro ciclo (senão o CP redigita 40 linhas).
            "status_planilha": PROTO_STATUS_DE.get(_norm_col(pega(linha, "status"))),
            "ocorrencia_planilha": pega(linha, "ocorrencia") or None,
            "data_conclusao_planilha": _proto_data(pega(linha, "data_conclusao")),
        })
    if not itens:
        raise ValueError("Cabeçalho encontrado, mas nenhuma linha de item com descrição.")
    # ordem repetida quebraria o unique (roteiro_id, ordem) no meio do insert
    vistos, saida = set(), []
    for it in itens:
        while it["ordem"] in vistos:
            it["ordem"] += 1
        vistos.add(it["ordem"])
        saida.append(it)
    return {"meta": meta, "itens": saida}
# ── PAPÉIS NO CLIENTE ──────────────────────────────────────────────────────
# DOIS EIXOS, e confundi-los abre acesso:
#   usuarios_login.perfil   → o que a pessoa é no SISTEMA (admin/comum/cliente/leitor)
#   usuario_clientes.papel  → o que ela é NAQUELE PROJETO (cp_totvs/consultor/…)
# O perfil continua mandando em "pode gravar no Tasks SC"; o papel manda em
# "pode coordenar ESTE protótipo".
def papel_no_cliente(customer, email=None):
    email = (email or effective_user() or "").lower()
    if not email:
        return None
    r = q("select papel from cockpit.usuario_clientes "
          "where customer=%s and lower(email)=%s", (customer, email), one=True)
    return (r or {}).get("papel")


def eh_cp_totvs(customer, email=None):
    """Admin sempre é — senão, marcar o primeiro CP seria impossível
    (bootstrap). Fora isso, exige o papel NESTE cliente e que a pessoa seja
    interna: papel de coordenação não se concede a login de cliente."""
    email = email or effective_user()
    if perfil_do(email) == "admin":
        return True
    return (papel_no_cliente(customer, email) == "cp_totvs"
            and perfil_do(email) in PERFIS_INTERNOS)


def require_cp(customer):
    """Portão da COORDENAÇÃO do protótipo: importar/criar roteiro, criar ciclo,
    montar a equipe, atribuir responsável. Consultor comum não passa — foi
    exatamente isso que você pediu ao separar o papel de CP."""
    if (r := require_auth()):
        return r
    if not eh_cp_totvs(customer):
        return _err(403, "Ação restrita ao Coordenador de Projetos da TOTVS neste cliente.")
    return None


# ── CATÁLOGO DE MÓDULOS DO CLIENTE ─────────────────────────────────────────
# FONTE ÚNICA, de propósito: a tela de Acessos e o modal "Equipe do protótipo"
# liberam módulo para a MESMA pessoa — duas listas diferentes é como se libera
# num lugar um módulo que o outro nem enxerga.
#
# ARMADILHA (11/09/2026): a lista vinha SÓ dos itens do roteiro do protótipo.
# Cliente sem MIT045 importado abria o modal com "* todos" e mais nada — e o
# admin concluía que o recurso estava quebrado. Módulo também mora em Cadastros
# (monitcad_tabelas) e em Movimentos (monitmov_itens), que existem bem antes do
# roteiro: é a união das três que dá uma lista útil no primeiro dia do projeto.
_SQL_MODULOS = """
    select r.customer as customer, i.modulo as modulo
      from cockpit.proto_itens i
      join cockpit.proto_roteiros r on r.id = i.roteiro_id
     where r.ativo
    union
    select t.customer, t.modulo from cockpit.monitcad_tabelas t
    union
    select m.customer, m.modulo from cockpit.monitmov_itens m
"""


# '(sem módulo)' é RÓTULO de tela que já entrou no dado (agrupamento de
# medição sem módulo preenchido). Liberar alguém nele não significa nada.
_MODULO_VALIDO = ("modulo is not null and btrim(modulo) <> '' "
                  "and lower(btrim(modulo)) not in "
                  "('(sem módulo)','(sem modulo)','sem módulo','sem modulo')")


def modulos_do_cliente(customer):
    """Módulos que EXISTEM neste cliente, ordenados. Lista vazia é resposta
    legítima: o projeto ainda não tem roteiro nem medição."""
    rows = q(f"""select distinct modulo from ({_SQL_MODULOS}) x
                  where customer=%s and {_MODULO_VALIDO}
                  order by 1""", (customer,))
    return [r["modulo"] for r in rows]


def modulos_por_cliente():
    """{customer: [módulos]} numa consulta só — a tela de Acessos mexe em
    vários clientes de uma vez e não pode fazer uma ida ao banco por cliente."""
    rows = q(f"""select customer, modulo from ({_SQL_MODULOS}) x
                  where {_MODULO_VALIDO}
                  group by customer, modulo order by customer, modulo""")
    mapa = {}
    for r in rows:
        mapa.setdefault(r["customer"], []).append(r["modulo"])
    return mapa


def proto_modulos_do_usuario(customer, email=None, papel="executa"):
    """Módulos em que a pessoa EXECUTA (ou VALIDA, com papel='valida').
    None = todos. Conjunto vazio = nenhum."""
    email = (email or effective_user() or "").lower()
    if not email:
        return set()
    # CP TOTVS e CP do Cliente enxergam e agem no cliente inteiro.
    if eh_cp_totvs(customer, email):
        return None
    if papel == "executa" and papel_no_cliente(customer, email) == "cp_cliente":
        return None
    rows = q("""select modulo from cockpit.proto_usuario_modulos
                 where customer=%s and lower(email)=%s and papel=%s""",
             (customer, email, papel))
    mods = {r["modulo"] for r in rows}
    return None if "*" in mods else mods


# ── ATRIBUIÇÃO DE RESPONSÁVEL ──────────────────────────────────────────────
# PRECEDÊNCIA, do mais específico para o mais geral. É isto que permite
# "define uma vez no roteiro e sobrescreve só onde precisa":
#   1) exceção do CICLO   : item > processo > módulo
#   2) padrão do ROTEIRO  : item > processo > módulo
#   3) a coluna Usuário que veio da planilha
# Inverter essa ordem faria a atribuição do módulo apagar a da linha.
_PESO_ESCOPO = {"item": 3, "processo": 2, "modulo": 1}


def proto_responsaveis(roteiro_id, ciclo_id, itens):
    """{item_id: email}. Uma consulta só; a precedência é resolvida aqui."""
    regras = q("""select ciclo_id, escopo, alvo, email
                    from cockpit.proto_atribuicoes
                   where roteiro_id=%s and (ciclo_id is null or ciclo_id=%s)""",
               (roteiro_id, ciclo_id))
    saida = {}
    for it in itens:
        melhor, peso = None, 0
        for r in regras:
            if r["escopo"] == "item" and str(r["alvo"]) != str(it["id"]):
                continue
            if r["escopo"] == "processo" and (r["alvo"] or "") != (it.get("processo") or ""):
                continue
            if r["escopo"] == "modulo" and (r["alvo"] or "") != (it.get("modulo") or ""):
                continue
            # exceção do ciclo pesa mais que qualquer padrão do roteiro
            p = _PESO_ESCOPO[r["escopo"]] + (10 if r["ciclo_id"] else 0)
            if p > peso:
                melhor, peso = r["email"], p
        saida[it["id"]] = melhor or (it.get("usuario") or None)
    return saida


# ── QUEM PODE RESPONDER E QUEM PODE VALIDAR ────────────────────────────────
def _proto_ciclo_aberto(ciclo):
    if not ciclo:
        return "Ciclo não encontrado."
    if not ciclo["aberto"]:
        return "Este ciclo está fechado — a foto dele já foi congelada."
    if not eh_interno() and not ciclo["visivel_cliente"]:
        return "Ciclo não disponível para você."
    return None


def proto_pode_responder(customer, ciclo, item, responsavel=None):
    """Mensagem do erro ou None. Responder é do CLIENTE (e da TOTVS no ciclo
    interno). Passa quem: é interno; é CP do Cliente; tem o módulo com papel
    'executa'; ou é o RESPONSÁVEL atribuído àquela linha — este último é o que
    faz a atribuição por linha valer alguma coisa."""
    if (m := _proto_ciclo_aberto(ciclo)):
        return m
    email = (effective_user() or "").lower()
    if eh_interno(email):
        return None
    if papel_no_cliente(customer, email) == "cp_cliente":
        return None
    if responsavel and responsavel.strip().lower() == email:
        return None
    mods = proto_modulos_do_usuario(customer, email, "executa")
    if mods is None:
        return None
    if not mods:
        return ("Você não é o responsável por este item e não tem módulo liberado "
                "neste cliente. Fale com o coordenador do projeto.")
    if (item.get("modulo") or "") not in mods:
        return f"O módulo {item.get('modulo') or '(sem módulo)'} não está liberado para você."
    return None


def proto_pode_validar(customer, ciclo, item):
    """Validar é da TOTVS. Passa o CP TOTVS e o consultor com o módulo no papel
    'valida'. O cliente NUNCA valida o próprio trabalho — é essa separação que
    dá sentido ao OK do consultor."""
    if (m := _proto_ciclo_aberto(ciclo)):
        return m
    email = (effective_user() or "").lower()
    if not eh_interno(email):
        return "Validar o item é ação do consultor da TOTVS."
    if eh_cp_totvs(customer, email):
        return None
    mods = proto_modulos_do_usuario(customer, email, "valida")
    if mods is None:
        return None
    if (item.get("modulo") or "") not in mods:
        return (f"Você não é o consultor responsável por validar o módulo "
                f"{item.get('modulo') or '(sem módulo)'} neste cliente.")
    return None


# ── Leitura ────────────────────────────────────────────────────────────────
def _proto_ctx(customer):
    """O que a tela precisa saber sobre QUEM está olhando, num lugar só."""
    email = (effective_user() or "").lower()
    papel = papel_no_cliente(customer, email)
    exec_ = proto_modulos_do_usuario(customer, email, "executa")
    vali = proto_modulos_do_usuario(customer, email, "valida")
    return {
        "email": email, "papel": papel, "interno": eh_interno(email),
        "cp_totvs": eh_cp_totvs(customer, email),
        "cp_cliente": papel == "cp_cliente",
        "modulos_executa": None if exec_ is None else sorted(exec_),
        "modulos_valida": None if vali is None else sorted(vali),
    }


@app.get("/api/proto/<customer>")
def api_proto_lista(customer):
    """Roteiros do cliente + ciclos. O cliente NUNCA vê o ciclo interno: ele é
    a consultoria ensaiando, não resultado."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    ctx = _proto_ctx(customer)
    roteiros = q("""select r.*, (select count(*) from cockpit.proto_itens i
                                  where i.roteiro_id=r.id and i.ativo) as n_itens
                      from cockpit.proto_roteiros r
                     where r.customer=%s and r.ativo
                     order by r.escopo_nome, r.created_at desc""", (customer,))
    ids = [r["id"] for r in roteiros]
    ciclos = q("""select c.*,
                    (select count(*) from cockpit.proto_resultados x
                      where x.ciclo_id=c.id and x.status <> 'nao_iniciado') as respondidos
                   from cockpit.proto_ciclos c
                  where c.roteiro_id = any(%s::uuid[]) order by c.tipo, c.numero""",
               (ids,)) if ids else []
    if not ctx["interno"]:
        ciclos = [c for c in ciclos if c["visivel_cliente"]]
    return _json({"ok": True, "customer": customer, "roteiros": roteiros,
                  "ciclos": ciclos, "ctx": ctx,
                  "status": [{"id": k, "label": PROTO_STATUS_LABEL[k]} for k in PROTO_STATUS],
                  "validacoes": [{"id": k, "label": PROTO_VALID_LABEL[k]} for k in PROTO_VALID],
                  "tipos": [{"id": t, "label": PROTO_TIPO_LABEL[t]} for t in PROTO_TIPOS],
                  "papeis": [{"id": p, "label": PROTO_PAPEL_LABEL[p]} for p in PROTO_PAPEIS]})


@app.get("/api/proto/<customer>/roteiro/<rid>")
def api_proto_roteiro(customer, rid):
    """Itens + ciclos + resultados + RESPONSÁVEL resolvido por item."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    rot = q("select * from cockpit.proto_roteiros where id=%s and customer=%s",
            (rid, customer), one=True)
    if not rot:
        return _err(404, "Roteiro não encontrado.")
    ctx = _proto_ctx(customer)
    itens = q("select * from cockpit.proto_itens where roteiro_id=%s and ativo order by ordem",
              (rid,))
    ciclos = q("select * from cockpit.proto_ciclos where roteiro_id=%s order by tipo, numero",
               (rid,))
    if not ctx["interno"]:
        ciclos = [c for c in ciclos if c["visivel_cliente"]]
    cids = [c["id"] for c in ciclos]
    res = q("select * from cockpit.proto_resultados where ciclo_id = any(%s::uuid[])",
            (cids,)) if cids else []
    # Responsável é por CICLO (a exceção do ciclo sobrepõe o padrão do roteiro),
    # então vem um mapa por ciclo — a tela troca de ciclo sem nova requisição.
    resp = {c["id"]: proto_responsaveis(rid, c["id"], itens) for c in ciclos}
    atrib = q("select * from cockpit.proto_atribuicoes where roteiro_id=%s", (rid,))
    return _json({"ok": True, "roteiro": rot, "itens": itens, "ciclos": ciclos,
                  "resultados": res, "responsaveis": resp, "atribuicoes": atrib,
                  "ctx": ctx,
                  "status": [{"id": k, "label": PROTO_STATUS_LABEL[k]} for k in PROTO_STATUS],
                  "validacoes": [{"id": k, "label": PROTO_VALID_LABEL[k]} for k in PROTO_VALID]})


# ── Importação e criação do roteiro ────────────────────────────────────────
def _proto_grava_itens(roteiro_id, itens):
    with db() as conn, conn.cursor() as cur:
        execute_values(cur,
            """insert into cockpit.proto_itens
                 (roteiro_id, ordem, modulo, processo, subprocesso, descricao,
                  consultor, usuario, data_planejada) values %s""",
            [(roteiro_id, i["ordem"], i.get("modulo"), i.get("processo"),
              i.get("subprocesso"), i["descricao"], i.get("consultor"),
              i.get("usuario"), i.get("data_planejada")) for i in itens])


def _proto_novo_roteiro(customer, a, **kw):
    escopo = (a.get("escopo") or "modulo").strip().lower()
    if escopo not in ("modulo", "processo"):
        raise ValueError("escopo deve ser 'modulo' ou 'processo'.")
    nome = (a.get("escopo_nome") or "").strip()
    if not nome:
        raise ValueError("Informe o nome do módulo ou do processo do roteiro.")
    return q("""insert into cockpit.proto_roteiros
                 (customer, codigo_projeto, escopo, escopo_nome, titulo, fonte_url,
                  fonte_csv_url, origem, arquivo_nome, data_prototipo, created_by, updated_by)
               values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) returning *""",
             (customer, (a.get("codigo_projeto") or "").strip() or None, escopo, nome,
              (a.get("titulo") or nome).strip(), (a.get("fonte_url") or "").strip() or None,
              kw.get("csv_url"), kw.get("origem", "upload"), kw.get("arquivo"),
              kw.get("data_prototipo"), current_user(), current_user()), one=True)


@app.post("/api/proto/<customer>/roteiro")
def api_proto_importar(customer):
    """Importa um MIT045. Três caminhos, um parser só: arquivo no corpo,
    ?csv_url= (a URL de 'Publicar na web > CSV') ou o CSV colado.

    O LINK DO DRIVE NÃO É FONTE DE DADOS: o backend na Vercel não tem
    credencial Google. Ele é gravado em fonte_url para rastreabilidade.
    """
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    a = request.args
    csv_url = (a.get("csv_url") or "").strip() or None
    nome = (a.get("arquivo") or "").strip()
    origem, linhas = "upload", None
    if csv_url:
        origem = "csv"
        try:
            rr = requests.get(csv_url, timeout=25)
            rr.raise_for_status()
            rr.encoding = rr.encoding or "utf-8"
            linhas = _proto_linhas_csv(rr.text)
        except Exception as e:
            return _err(502, f"Não consegui ler o CSV publicado: {e}")
    else:
        arq = request.files.get("arquivo")
        dados = arq.read() if arq else request.get_data()
        nome = nome or (arq.filename if arq else "")
        if not dados:
            return _err(400, "Envie o arquivo do roteiro (.xlsx ou .csv) ou uma csv_url.")
        try:
            if nome.lower().endswith(".xlsx") or dados[:2] == b"PK":
                linhas = _proto_linhas_xlsx(dados)
            else:
                linhas = _proto_linhas_csv(dados.decode("utf-8-sig", "replace"))
        except ValueError as e:
            return _err(422, str(e))
    try:
        parsed = _proto_parse(linhas)
        rot = _proto_novo_roteiro(customer, a, csv_url=csv_url, origem=origem,
                                  arquivo=nome or None,
                                  data_prototipo=parsed["meta"]["data_prototipo"])
    except ValueError as e:
        return _err(422, str(e))

    itens = parsed["itens"]
    _proto_grava_itens(rot["id"], itens)
    ciclo = None
    if a.get("semear") == "1":
        ciclo = _proto_cria_ciclo(rot["id"], "interno", 1, visivel_cliente=False)
        por_ordem = {r["ordem"]: r["id"] for r in
                     q("select id, ordem from cockpit.proto_itens where roteiro_id=%s",
                       (rot["id"],))}
        with db() as conn, conn.cursor() as cur:
            execute_values(cur,
                """insert into cockpit.proto_resultados
                     (ciclo_id, item_id, status, ocorrencia, data_conclusao, respondido_por)
                   values %s on conflict (ciclo_id, item_id) do nothing""",
                [(ciclo["id"], por_ordem[i["ordem"]], i["status_planilha"] or "nao_iniciado",
                  i["ocorrencia_planilha"], i["data_conclusao_planilha"], current_user())
                 for i in itens if i["ordem"] in por_ordem])
    return _json({"ok": True, "roteiro": rot, "itens": len(itens),
                  "modulos": sorted({i["modulo"] for i in itens if i["modulo"]}),
                  "ciclo_semeado": ciclo, "meta": parsed["meta"]})


@app.post("/api/proto/<customer>/roteiro/manual")
def api_proto_roteiro_manual(customer):
    """Cria um roteiro SEM planilha: do zero (itens no corpo, podendo vir
    vazio) ou copiando outro roteiro já existente (?copiar_de=<rid>).

    Copiar é o que faz a biblioteca crescer sozinha: o roteiro de 'Gestão de
    Contratos' de um cliente vira o ponto de partida do próximo. Copia a
    ESTRUTURA (itens), nunca os resultados — resultado é do projeto onde foi
    executado.
    """
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    b = request.get_json(silent=True) or {}
    copiar = (request.args.get("copiar_de") or b.get("copiar_de") or "").strip()
    itens = []
    if copiar:
        origem = q("select * from cockpit.proto_roteiros where id=%s", (copiar,), one=True)
        if not origem:
            return _err(404, "Roteiro de origem não encontrado.")
        if (d2 := deny_customer(origem["customer"])):
            return d2          # não copia de cliente que você não enxerga
        itens = [dict(x) for x in q(
            """select ordem, modulo, processo, subprocesso, descricao, consultor,
                      null::date as data_planejada, null as usuario
                 from cockpit.proto_itens where roteiro_id=%s and ativo order by ordem""",
            (copiar,))]
        b.setdefault("titulo", origem["titulo"])
        b.setdefault("escopo", origem["escopo"])
        b.setdefault("escopo_nome", origem["escopo_nome"])
    else:
        for n, it in enumerate(b.get("itens") or [], start=1):
            desc = (it.get("descricao") or "").strip()
            if not desc:
                continue
            itens.append({"ordem": int(it.get("ordem") or n),
                          "modulo": (it.get("modulo") or "").strip() or None,
                          "processo": (it.get("processo") or "").strip() or None,
                          "subprocesso": (it.get("subprocesso") or "").strip() or None,
                          "descricao": desc,
                          "consultor": (it.get("consultor") or "").strip() or None,
                          "usuario": (it.get("usuario") or "").strip() or None,
                          "data_planejada": _proto_data(it.get("data_planejada"))})
    try:
        rot = _proto_novo_roteiro(customer, b, origem="cowork" if copiar else "upload")
    except ValueError as e:
        return _err(422, str(e))
    if itens:
        _proto_grava_itens(rot["id"], itens)
    return _json({"ok": True, "roteiro": rot, "itens": len(itens),
                  "copiado_de": copiar or None})


@app.get("/api/proto/modelos")
def api_proto_modelos():
    """Roteiros que dá para copiar — só dos clientes que o usuário enxerga."""
    if (r := require_auth()):
        return r
    if not eh_interno():
        return _err(403, "Ação restrita à equipe TOTVS.")
    permitidos = allowed_customers()
    rows = q("""select r.id, r.customer, r.titulo, r.escopo, r.escopo_nome,
                       c.nome as cliente_nome,
                       (select count(*) from cockpit.proto_itens i
                         where i.roteiro_id=r.id and i.ativo) as n_itens
                  from cockpit.proto_roteiros r
                  left join cockpit.clientes c on c.customer=r.customer
                 where r.ativo order by c.nome, r.titulo""")
    if permitidos is not None:
        rows = [x for x in rows if x["customer"] in permitidos]
    return _json({"ok": True, "roteiros": rows})


@app.route("/api/proto/<customer>/roteiro/<rid>/itens", methods=["POST", "DELETE"])
def api_proto_itens(customer, rid):
    """Acrescenta itens ao roteiro (linha a linha ou em bloco) ou inativa um."""
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    if not q("select 1 from cockpit.proto_roteiros where id=%s and customer=%s",
             (rid, customer), one=True):
        return _err(404, "Roteiro não encontrado.")
    if request.method == "DELETE":
        iid = (request.args.get("item") or "").strip()
        # inativa, não apaga: o item pode ter resultado de ciclo anterior
        execute("update cockpit.proto_itens set ativo=false where id=%s and roteiro_id=%s",
                (iid, rid))
        return _json({"ok": True})
    b = request.get_json(silent=True) or {}
    prox = (q("select coalesce(max(ordem),0)+1 n from cockpit.proto_itens where roteiro_id=%s",
              (rid,), one=True) or {}).get("n", 1)
    novos = []
    for it in (b.get("itens") or [b]):
        desc = (it.get("descricao") or "").strip()
        if not desc:
            continue
        novos.append({"ordem": prox + len(novos),
                      "modulo": (it.get("modulo") or "").strip() or None,
                      "processo": (it.get("processo") or "").strip() or None,
                      "subprocesso": (it.get("subprocesso") or "").strip() or None,
                      "descricao": desc,
                      "consultor": (it.get("consultor") or "").strip() or None,
                      "usuario": (it.get("usuario") or "").strip() or None,
                      "data_planejada": _proto_data(it.get("data_planejada"))})
    if not novos:
        return _err(400, "Nenhum item com descrição.")
    _proto_grava_itens(rid, novos)
    return _json({"ok": True, "itens": len(novos)})


@app.delete("/api/proto/<customer>/roteiro/<rid>")
def api_proto_remover(customer, rid):
    """Inativa o roteiro. NÃO apaga: os resultados dos ciclos são histórico."""
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    execute("update cockpit.proto_roteiros set ativo=false, updated_at=now(), updated_by=%s "
            "where id=%s and customer=%s", (current_user(), rid, customer))
    return _json({"ok": True})


# ── Ciclos ─────────────────────────────────────────────────────────────────
def _proto_cria_ciclo(roteiro_id, tipo, numero=None, visivel_cliente=None, **kw):
    if numero is None:
        numero = (q("select coalesce(max(numero),0)+1 n from cockpit.proto_ciclos "
                    "where roteiro_id=%s and tipo=%s", (roteiro_id, tipo), one=True) or {})["n"]
    if visivel_cliente is None:
        # Interno é ensaio da consultoria: nasce fechado para o cliente.
        visivel_cliente = tipo != "interno"
    return q("""insert into cockpit.proto_ciclos
                  (roteiro_id, tipo, numero, visivel_cliente, data_inicio, observacao, created_by)
                values (%s,%s,%s,%s,%s,%s,%s)
                on conflict (roteiro_id, tipo, numero) do update
                  set visivel_cliente = excluded.visivel_cliente
                returning *""",
             (roteiro_id, tipo, numero, visivel_cliente, kw.get("data_inicio"),
              kw.get("observacao"), current_user()), one=True)


@app.post("/api/proto/<customer>/roteiro/<rid>/ciclo")
def api_proto_ciclo_novo(customer, rid):
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    if not q("select 1 from cockpit.proto_roteiros where id=%s and customer=%s",
             (rid, customer), one=True):
        return _err(404, "Roteiro não encontrado.")
    b = request.get_json(silent=True) or {}
    tipo = (b.get("tipo") or "").strip().lower()
    if tipo not in PROTO_TIPOS:
        return _err(400, f"tipo deve ser um de: {', '.join(PROTO_TIPOS)}.")
    vis = b.get("visivel_cliente")
    c = _proto_cria_ciclo(rid, tipo, b.get("numero"), None if vis is None else bool(vis),
                          data_inicio=_proto_data(b.get("data_inicio")),
                          observacao=(b.get("observacao") or "").strip() or None)
    return _json({"ok": True, "ciclo": c})


@app.route("/api/proto/<customer>/ciclo/<cid>", methods=["PATCH", "POST"])
def api_proto_ciclo_editar(customer, cid):
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    if not q("""select 1 from cockpit.proto_ciclos c join cockpit.proto_roteiros r
                  on r.id=c.roteiro_id where c.id=%s and r.customer=%s""",
             (cid, customer), one=True):
        return _err(404, "Ciclo não encontrado.")
    b = request.get_json(silent=True) or {}
    campos, vals = [], []
    for k, conv in (("aberto", bool), ("visivel_cliente", bool),
                    ("data_inicio", _proto_data), ("data_fim", _proto_data),
                    ("observacao", lambda v: (v or "").strip() or None)):
        if k in b:
            campos.append(f"{k}=%s"); vals.append(conv(b[k]))
    if not campos:
        return _err(400, "Nada para alterar.")
    vals.append(cid)
    execute(f"update cockpit.proto_ciclos set {', '.join(campos)} where id=%s", tuple(vals))
    return _json({"ok": True, "ciclo": q("select * from cockpit.proto_ciclos where id=%s",
                                         (cid,), one=True)})


@app.delete("/api/proto/<customer>/ciclo/<cid>")
def api_proto_ciclo_remover(customer, cid):
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    if effective_user() != current_user():
        return _err(409, "Saia da simulação ('ver como') antes de apagar ciclos.")
    n = (q("""select count(*) n from cockpit.proto_resultados x
                join cockpit.proto_ciclos c on c.id=x.ciclo_id
                join cockpit.proto_roteiros r on r.id=c.roteiro_id
               where c.id=%s and r.customer=%s and x.status <> 'nao_iniciado'""",
           (cid, customer), one=True) or {}).get("n", 0)
    if n and request.args.get("confirmar") != "1":
        return _err(409, f"Este ciclo já tem {n} item(ns) respondido(s). Reenvie com "
                         "confirmar=1 se é isso mesmo — não há desfazer.")
    execute("""delete from cockpit.proto_ciclos c using cockpit.proto_roteiros r
                where c.roteiro_id=r.id and c.id=%s and r.customer=%s""", (cid, customer))
    return _json({"ok": True})


# ── Resposta do cliente e validação do consultor ───────────────────────────
def _proto_ciclo_do(customer, cid):
    return q("""select c.* from cockpit.proto_ciclos c
                  join cockpit.proto_roteiros r on r.id=c.roteiro_id
                 where c.id=%s and r.customer=%s""", (cid, customer), one=True)


@app.post("/api/proto/<customer>/ciclo/<cid>/resultado")
def api_proto_resultado(customer, cid):
    """A RESPOSTA do cliente. Body: {itens:[{item_id,status,nota,ocorrencia,
    data_conclusao}]} — ou os campos soltos para um item só.

    Responder um item que já tinha OK do consultor devolve a validação para
    'pendente': mudou a resposta, o OK anterior era sobre outra coisa.
    """
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    ciclo = _proto_ciclo_do(customer, cid)
    if not ciclo:
        return _err(404, "Ciclo não encontrado.")
    b = request.get_json(silent=True) or {}
    lote = b.get("itens") if isinstance(b.get("itens"), list) else [b]
    ids = [str(x.get("item_id") or "") for x in lote if x.get("item_id")]
    if not lote or len(ids) != len(lote):
        return _err(400, "Todo item precisa de item_id.")
    itens = q("select * from cockpit.proto_itens where id = any(%s::uuid[]) and roteiro_id=%s",
              (ids, ciclo["roteiro_id"]))
    donos = {str(i["id"]): i for i in itens}
    if len(donos) != len(set(ids)):
        return _err(404, "Item que não pertence a este roteiro.")
    resp = proto_responsaveis(ciclo["roteiro_id"], cid, itens)

    linhas = []
    for x in lote:
        item = donos[str(x["item_id"])]
        if (msg := proto_pode_responder(customer, ciclo, item, resp.get(item["id"]))):
            return _err(403, msg)
        st = (x.get("status") or "nao_iniciado").strip()
        if st not in PROTO_STATUS:
            return _err(400, f"Status inválido: {st}")
        nota = x.get("nota")
        if nota in ("", None):
            nota = None
        else:
            try:
                nota = int(nota)
            except (TypeError, ValueError):
                return _err(400, "Nota tem que ser um número de 0 a 10.")
            if not 0 <= nota <= 10:
                return _err(400, "Nota fora da faixa 0–10.")
        linhas.append((cid, item["id"], st, nota,
                       (x.get("ocorrencia") or "").strip() or None,
                       _proto_data(x.get("data_conclusao")), effective_user()))
    with db() as conn, conn.cursor() as cur:
        execute_values(cur,
            """insert into cockpit.proto_resultados
                 (ciclo_id, item_id, status, nota, ocorrencia, data_conclusao, respondido_por)
               values %s
               on conflict (ciclo_id, item_id) do update set
                 status=excluded.status, nota=excluded.nota,
                 ocorrencia=excluded.ocorrencia, data_conclusao=excluded.data_conclusao,
                 respondido_por=excluded.respondido_por, respondido_em=now(),
                 -- mudou a resposta: o OK anterior era sobre outra coisa
                 validacao='pendente', validado_por=null, validado_em=null""", linhas)
    return _json({"ok": True, "gravados": len(linhas)})


@app.post("/api/proto/<customer>/ciclo/<cid>/validar")
def api_proto_validar(customer, cid):
    """O OK DO CONSULTOR. Body: {itens:[{item_id, validacao, parecer}]}.

    Só a TOTVS valida — o cliente nunca dá OK no próprio trabalho. É essa
    separação que faz o aproveitamento significar alguma coisa.
    """
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    ciclo = _proto_ciclo_do(customer, cid)
    if not ciclo:
        return _err(404, "Ciclo não encontrado.")
    b = request.get_json(silent=True) or {}
    lote = b.get("itens") if isinstance(b.get("itens"), list) else [b]
    ids = [str(x.get("item_id") or "") for x in lote if x.get("item_id")]
    if not lote or len(ids) != len(lote):
        return _err(400, "Todo item precisa de item_id.")
    donos = {str(i["id"]): i for i in q(
        "select * from cockpit.proto_itens where id = any(%s::uuid[]) and roteiro_id=%s",
        (ids, ciclo["roteiro_id"]))}
    if len(donos) != len(set(ids)):
        return _err(404, "Item que não pertence a este roteiro.")
    linhas = []
    for x in lote:
        item = donos[str(x["item_id"])]
        if (msg := proto_pode_validar(customer, ciclo, item)):
            return _err(403, msg)
        v = (x.get("validacao") or "pendente").strip()
        if v not in PROTO_VALID:
            return _err(400, f"Validação inválida: {v}")
        linhas.append((cid, item["id"], v, (x.get("parecer") or "").strip() or None,
                       effective_user()))
    with db() as conn, conn.cursor() as cur:
        execute_values(cur,
            """insert into cockpit.proto_resultados
                 (ciclo_id, item_id, validacao, parecer, validado_por, validado_em, status)
               values %s
               on conflict (ciclo_id, item_id) do update set
                 validacao=excluded.validacao, parecer=excluded.parecer,
                 validado_por=excluded.validado_por, validado_em=now()""",
            [(c, i, v, p, u, "nao_iniciado") for c, i, v, p, u in linhas],
            # 6 %s + now() = as 7 colunas do insert. Contar errado aqui estoura
            # "not all arguments converted" só quando alguém validar de verdade.
            template="(%s,%s,%s,%s,%s,now(),%s)")
    return _json({"ok": True, "validados": len(linhas)})


# ── Equipe: papéis + módulos de quem executa e de quem valida ──────────────
@app.route("/api/proto/<customer>/equipe", methods=["GET", "POST", "DELETE"])
def api_proto_equipe(customer):
    """O CP TOTVS monta a equipe aqui: define o papel de cada um NESTE cliente
    e os módulos em que cada consultor VALIDA / cada usuário-chave EXECUTA."""
    if (r := require_cp(customer)):
        return r
    if (d := deny_customer(customer)):
        return d
    if request.method == "GET":
        pessoas = q("""select uc.email, uc.papel, u.nome, u.perfil, u.ativo
                         from cockpit.usuario_clientes uc
                         join cockpit.usuarios_login u on u.email = uc.email
                        where uc.customer=%s order by coalesce(u.nome, uc.email)""",
                    (customer,))
        # Mesma lista da tela de Acessos (roteiro + Cadastros + Movimentos):
        # antes vinha só do roteiro e o modal abria vazio em cliente sem MIT045.
        mods = modulos_do_cliente(customer)
        procs = q("""select distinct i.processo from cockpit.proto_itens i
                       join cockpit.proto_roteiros r on r.id=i.roteiro_id
                      where r.customer=%s and r.ativo and i.processo is not null
                      order by 1""", (customer,))
        return _json({"ok": True, "pessoas": pessoas,
                      "liberacoes": q("""select * from cockpit.proto_usuario_modulos
                                          where customer=%s order by email, papel, modulo""",
                                      (customer,)),
                      "modulos": mods,
                      "processos": [p["processo"] for p in procs],
                      "papeis": [{"id": p, "label": PROTO_PAPEL_LABEL[p]} for p in PROTO_PAPEIS]})

    b = request.get_json(silent=True) or {}
    emails = emails_do_texto(b.get("emails"))
    if not emails:
        return _err(400, "Informe pelo menos um e-mail.")
    # Só mexe em quem JÁ está liberado no cliente. Papel não cria acesso — o
    # acesso continua sendo decisão da tela de Acessos.
    conhecidos = {r["email"].lower(): r["email"] for r in q(
        "select email from cockpit.usuario_clientes where customer=%s and lower(email)=any(%s)",
        (customer, emails))}
    if (fora := [e for e in emails if e not in conhecidos]):
        return _err(409, "Estes e-mails ainda não têm acesso a este cliente — libere em "
                         "Acessos primeiro: " + ", ".join(fora))
    alvos = [conhecidos[e] for e in emails]

    if request.method == "DELETE":
        modulos = [str(m).strip() for m in (b.get("modulos") or []) if str(m).strip()]
        papel = (b.get("papel_modulo") or "executa").strip()
        if modulos:
            execute("""delete from cockpit.proto_usuario_modulos
                        where customer=%s and lower(email)=any(%s) and modulo=any(%s)
                          and papel=%s""", (customer, emails, modulos, papel))
        else:
            execute("""update cockpit.usuario_clientes set papel=null
                        where customer=%s and lower(email)=any(%s)""", (customer, emails))
        return _json({"ok": True})

    papel = (b.get("papel") or "").strip() or None
    if papel and papel not in PROTO_PAPEIS:
        return _err(400, f"Papel desconhecido: {papel}")
    if papel in PROTO_PAPEIS_TOTVS:
        # papel de coordenação/validação não se concede a login de cliente
        externos = [e for e in alvos if perfil_do(e) not in PERFIS_INTERNOS]
        if externos:
            return _err(409, "Papel da TOTVS não pode ser dado a login de cliente: "
                             + ", ".join(externos))
    if papel:
        execute("""update cockpit.usuario_clientes set papel=%s
                    where customer=%s and lower(email)=any(%s)""",
                (papel, customer, emails))

    modulos = [str(m).strip() for m in (b.get("modulos") or []) if str(m).strip()]
    papel_mod = (b.get("papel_modulo") or "").strip()
    if modulos:
        if papel_mod not in ("executa", "valida"):
            return _err(400, "papel_modulo deve ser 'executa' ou 'valida'.")
        proc = (b.get("processo") or "").strip() or None
        with db() as conn, conn.cursor() as cur:
            execute_values(cur,
                """insert into cockpit.proto_usuario_modulos
                     (customer, email, modulo, papel, processo, created_by) values %s
                   on conflict do nothing""",
                [(customer, e, m, papel_mod, proc, current_user()) for e in alvos for m in modulos])
    return _json({"ok": True, "emails": alvos, "papel": papel,
                  "modulos": len(modulos)})


# ── Atribuição: quem executa o quê ─────────────────────────────────────────
@app.route("/api/proto/<customer>/roteiro/<rid>/atribuicoes", methods=["POST", "DELETE"])
def api_proto_atribuicoes(customer, rid):
    """Body: {ciclo_id: null|uuid, regras:[{escopo,alvo,email}]}.

    ciclo_id NULL = padrão do roteiro (vale para todos os ciclos).
    ciclo_id preenchido = exceção daquele ciclo. É isso que evita redistribuir
    40 linhas a cada ciclo e ainda permite trocar uma pessoa só onde precisa.

    O CP do CLIENTE também atribui — distribuir trabalho entre quem já tem
    acesso é o trabalho dele. O que ele não faz é liberar acesso novo.
    """
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    if not (eh_cp_totvs(customer) or papel_no_cliente(customer) == "cp_cliente"):
        return _err(403, "Atribuir responsável é do Coordenador de Projetos "
                         "(TOTVS ou do cliente).")
    if not q("select 1 from cockpit.proto_roteiros where id=%s and customer=%s",
             (rid, customer), one=True):
        return _err(404, "Roteiro não encontrado.")
    b = request.get_json(silent=True) or {}
    ciclo_id = (b.get("ciclo_id") or None)
    regras = b.get("regras") if isinstance(b.get("regras"), list) else [b]
    limpas = []
    for g in regras:
        escopo = (g.get("escopo") or "").strip()
        alvo = str(g.get("alvo") or "").strip()
        if escopo not in ("modulo", "processo", "item") or not alvo:
            return _err(400, "Cada regra precisa de escopo (modulo|processo|item) e alvo.")
        limpas.append((escopo, alvo, (g.get("email") or "").strip().lower()))
    if not limpas:
        return _err(400, "Nada para atribuir.")

    if request.method == "DELETE":
        for escopo, alvo, _ in limpas:
            execute("""delete from cockpit.proto_atribuicoes
                        where roteiro_id=%s and escopo=%s and alvo=%s
                          and ciclo_id is not distinct from %s""",
                    (rid, escopo, alvo, ciclo_id))
        return _json({"ok": True})

    # Atribuir a quem não tem acesso ao cliente cria responsável fantasma: a
    # linha fica com dono e ninguém consegue responder.
    emails = {e for _, _, e in limpas if e}
    if emails:
        tem = {r["email"].lower() for r in q(
            "select email from cockpit.usuario_clientes where customer=%s and lower(email)=any(%s)",
            (customer, sorted(emails)))}
        if (fora := sorted(emails - tem)):
            return _err(409, "Sem acesso a este cliente (libere em Acessos antes): "
                             + ", ".join(fora))
    with db() as conn, conn.cursor() as cur:
        execute_values(cur,
            """insert into cockpit.proto_atribuicoes
                 (roteiro_id, ciclo_id, escopo, alvo, email, created_by) values %s
               on conflict (roteiro_id, ciclo_id, escopo, alvo) do update
                 set email = excluded.email, created_by = excluded.created_by""",
            [(rid, ciclo_id, escopo, alvo, email, current_user())
             for escopo, alvo, email in limpas])
    return _json({"ok": True, "regras": len(limpas), "ciclo_id": ciclo_id})


# ── Indicadores ────────────────────────────────────────────────────────────
def _proto_agrega(chave, res_por_item, itens, resp=None):
    """Agrega por uma chave qualquer (módulo, consultor, usuário, responsável).

    'nao_aplicavel' SAI do denominador: item que o cliente não usa reprovaria o
    módulo inteiro sem nada de errado ter havido.

    EXECUÇÃO e APROVEITAMENTO são numeradores DIFERENTES de propósito:
      execução       = o cliente respondeu 'executado'
      aproveitamento = o consultor deu OK
    A diferença entre os dois é a fila de validação.
    """
    saida = {}
    for it in itens:
        k = (resp.get(it["id"]) if (chave == "responsavel" and resp) else it.get(chave)) \
            or "(sem informação)"
        b = saida.setdefault(k, {"chave": k, "total": 0, "aplicaveis": 0, "executado": 0,
                                 "duvida": 0, "erro": 0, "nao_iniciado": 0,
                                 "nao_aplicavel": 0, "ok": 0, "reprovado": 0,
                                 "aguardando": 0, "notas": []})
        r = res_por_item.get(it["id"]) or {}
        st = r.get("status") or "nao_iniciado"
        val = r.get("validacao") or "pendente"
        b["total"] += 1
        b[st] = b.get(st, 0) + 1
        if st != "nao_aplicavel":
            b["aplicaveis"] += 1
        if val == "ok":
            b["ok"] += 1
        elif val == "reprovado":
            b["reprovado"] += 1
        # aguardando = trabalho FEITO que ninguém conferiu ainda. É este número
        # que denuncia consultor parado — não a contagem de executados.
        if st in ("executado", "duvida") and val == "pendente":
            b["aguardando"] += 1
        if r.get("nota") is not None:
            b["notas"].append(r["nota"])
    for b in saida.values():
        ap = b["aplicaveis"]
        b["execucao"] = round(100.0 * b["executado"] / ap, 1) if ap else None
        b["aproveitamento"] = round(100.0 * b["ok"] / ap, 1) if ap else None
        b["maturacao"] = round(sum(b["notas"]) / len(b["notas"]), 1) if b["notas"] else None
        b["respondidos"] = b["total"] - b["nao_iniciado"]
        b.pop("notas", None)
    return sorted(saida.values(), key=lambda x: str(x["chave"]))


@app.get("/api/proto/<customer>/indicadores/<rid>")
def api_proto_indicadores(customer, rid):
    """Um retrato por ciclo + o comparativo entre ciclos. O comparativo É o
    produto: número solto de um ciclo não diz se o protótipo andou."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    rot = q("select * from cockpit.proto_roteiros where id=%s and customer=%s",
            (rid, customer), one=True)
    if not rot:
        return _err(404, "Roteiro não encontrado.")
    itens = q("select * from cockpit.proto_itens where roteiro_id=%s and ativo", (rid,))
    ciclos = q("select * from cockpit.proto_ciclos where roteiro_id=%s", (rid,))
    if not eh_interno():
        ciclos = [c for c in ciclos if c["visivel_cliente"]]
    cids = [c["id"] for c in ciclos]
    res = q("select * from cockpit.proto_resultados where ciclo_id = any(%s::uuid[])",
            (cids,)) if cids else []
    por_ciclo = {}
    for r_ in res:
        por_ciclo.setdefault(r_["ciclo_id"], {})[r_["item_id"]] = r_

    ORDEM = {"interno": 0, "isolado": 1, "integrado": 2}
    saida = []
    for c in sorted(ciclos, key=lambda c: (ORDEM.get(c["tipo"], 9), c["numero"])):
        m = por_ciclo.get(c["id"], {})
        resp = proto_responsaveis(rid, c["id"], itens)
        geral = _proto_agrega("__todos__", m, [{**i, "__todos__": "Geral"} for i in itens])
        saida.append({
            "ciclo": c, "rotulo": f"{PROTO_TIPO_LABEL[c['tipo']]} · {c['numero']}",
            "geral": geral[0] if geral else None,
            "por_modulo": _proto_agrega("modulo", m, itens),
            "por_consultor": _proto_agrega("consultor", m, itens),
            "por_responsavel": _proto_agrega("responsavel", m, itens, resp),
        })
    # A evolução é contra o ciclo ANTERIOR na ordem metodológica, não na de criação.
    for i, s in enumerate(saida):
        ant = saida[i - 1]["geral"] if i and saida[i - 1]["geral"] else None
        g = s["geral"]
        for campo in ("aproveitamento", "execucao"):
            s[f"delta_{campo}"] = (
                None if not (g and ant and g[campo] is not None and ant[campo] is not None)
                else round(g[campo] - ant[campo], 1))
    return _json({"ok": True, "roteiro": rot, "ciclos": saida, "itens": len(itens)})


# ── Exportar de volta no formato MIT045 ────────────────────────────────────
@app.get("/api/proto/<customer>/roteiro/<rid>/export")
def api_proto_export(customer, rid):
    """Gera o .xlsx no layout MIT045 com o preenchimento de UM ciclo, para
    anexar no Drive do projeto. Mantém a coluna vazia à esquerda e as três
    linhas de cabeçalho — é assim que o arquivo original é, e é assim que o
    nosso próprio parser reconhece na volta."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "prototipo")):
        return d
    rot = q("select * from cockpit.proto_roteiros where id=%s and customer=%s",
            (rid, customer), one=True)
    if not rot:
        return _err(404, "Roteiro não encontrado.")
    ciclo = None
    if (cid := (request.args.get("ciclo") or "").strip()):
        ciclo = _proto_ciclo_do(customer, cid)
        if not ciclo:
            return _err(404, "Ciclo não encontrado.")
        if not eh_interno() and not ciclo["visivel_cliente"]:
            return _err(403, "Ciclo não disponível para você.")
    itens = q("select * from cockpit.proto_itens where roteiro_id=%s and ativo order by ordem",
              (rid,))
    res = {r["item_id"]: r for r in q(
        "select * from cockpit.proto_resultados where ciclo_id=%s", (cid,))} if ciclo else {}
    resp = proto_responsaveis(rid, ciclo["id"] if ciclo else None, itens)
    cli = (q("select nome from cockpit.clientes where customer=%s", (customer,), one=True)
           or {}).get("nome") or customer
    try:
        import openpyxl
    except ImportError:
        return _err(503, "Exportação .xlsx indisponível no servidor (openpyxl).")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roteiro"
    ws.append([None, "Roteiro Protótipo ", "Roteiro de Protótipo - MIT045"])
    ws.append([None, "Projeto", cli, None, None, None, None,
               "Data Protótipo", rot["data_prototipo"]])
    ws.append([None, "Ciclo", (f"{PROTO_TIPO_LABEL[ciclo['tipo']]} · {ciclo['numero']}"
                               if ciclo else "sem ciclo")])
    ws.append([None, "ID", "Módulo", "Processo", "Subprocesso", "Descrição", "Consultor",
               "Usuário", "Data planejada", "Status", "Data conclusão", "Ocorrência",
               "Validação", "Parecer do consultor", "Nota"])
    for i in itens:
        r = res.get(i["id"]) or {}
        ws.append([None, i["ordem"], i["modulo"], i["processo"], i["subprocesso"],
                   i["descricao"], i["consultor"], resp.get(i["id"]) or i["usuario"],
                   i["data_planejada"],
                   PROTO_STATUS_LABEL.get(r.get("status") or "nao_iniciado"),
                   r.get("data_conclusao"), r.get("ocorrencia"),
                   PROTO_VALID_LABEL.get(r.get("validacao") or "pendente"),
                   r.get("parecer"), r.get("nota")])
    buf = io.BytesIO()
    wb.save(buf)
    nome = _slug(f"MIT045 {rot['titulo']} {cli}") + ".xlsx"
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nome}"'})


# ═══════════════════════════════════════════════════════════════════════════
#  ESTRUTURA DE EMPRESAS — a DEFINIÇÃO (o que foi decidido na implantação)
# ═══════════════════════════════════════════════════════════════════════════
# NÃO confundir com a aba Empresas e Compartilhamento (monitemp_*): lá se MEDE
# a base do cliente (SM0 + SX2 lidas do banco dele); aqui se DESENHA, antes de
# existir base. Uma responde "como está", a outra "como combinamos que seria".
# A TDN é explícita: "a estrutura das tabelas deve ser definida na implantação
# do sistema (...) não aconselhamos a alteração no compartilhamento das tabelas
# após a implantação" — por isso esta tela existe, e existe ANTES da carga.
#
# COMPARTILHAMENTO SEMPRE EM 3 POSIÇÕES: Empresa · Unidade de Negócio · Filial,
# na MESMA ordem de X2_MODOEMP|X2_MODOUN|X2_MODO que a aba de medição já mostra.
# 'E' = exclusiva, 'C' = compartilhada. Trocar essa ordem entre as duas telas
# faria o consultor comparar coisas diferentes achando que são a mesma.

ESTRUT_LEIAUTES = {
    "FF":     {"label": "FF · só filial", "tam": (0, 0, 2),
               "ajuda": "Uma empresa só. O código de filial é o código inteiro."},
    "EEFF":   {"label": "EEFF · empresa + filial", "tam": (2, 0, 2),
               "ajuda": "O código da filial começa pelo código da empresa (ex.: 01 + 01 = 0101)."},
    "EEUUFF": {"label": "EEUUFF · empresa + unidade + filial", "tam": (2, 2, 2),
               "ajuda": "Empresa + unidade de negócio + filial (ex.: 01 + 01 + 02 = 010102)."},
}
# Índice de cada nível dentro da string de 3 posições do compartilhamento.
ESTRUT_NIVEL = {"empresa": 0, "unidade": 1, "filial": 2}
ESTRUT_NIVEL_LABEL = ("Empresa", "Unidade de Negócio", "Filial")


def _estrut_niveis(leiaute):
    """Quais posições do compartilhamento VALEM neste leiaute.
    Sem unidade de negócio no código, a posição do meio é decorativa: cobrar
    coerência dela geraria alerta que ninguém consegue resolver."""
    t = ESTRUT_LEIAUTES.get(leiaute, ESTRUT_LEIAUTES["EEFF"])["tam"]
    return [i for i in range(3) if t[i] > 0]


def _estrut_compart_ok(v):
    return isinstance(v, str) and len(v) == 3 and all(c in "EC" for c in v)


def _estrut_mais_compart(a, b, niveis):
    """'a' é MAIS compartilhada que 'b'? (C é mais amplo que E).
    Só é verdade quando 'a' é >= em TODOS os níveis que valem e > em algum —
    dois modos cruzados (um mais amplo aqui, menos ali) não são comparáveis e
    não viram alerta: inventar ordem onde não há vira alarme falso."""
    if not (_estrut_compart_ok(a) and _estrut_compart_ok(b)):
        return False
    maior = False
    for i in niveis:
        pa, pb = a[i] == "C", b[i] == "C"
        if pb and not pa:
            return False
        if pa and not pb:
            maior = True
    return maior


def _estrut_cfg(customer):
    r = q("select * from cockpit.estrut_config where customer=%s", (customer,), one=True)
    if r:
        return dict(r)
    return {"customer": customer, "leiaute": "EEFF", "tam_empresa": 2,
            "tam_unidade": 0, "tam_filial": 2, "observacao": None,
            "definido_por": None, "definido_em": None, "novo": True}


def _estrut_tam_esperado(cfg):
    t = ESTRUT_LEIAUTES.get(cfg["leiaute"], ESTRUT_LEIAUTES["EEFF"])["tam"]
    return (cfg.get("tam_empresa") or t[0]) + (cfg.get("tam_unidade") or t[1]) \
         + (cfg.get("tam_filial") or t[2])


def _estrut_alertas(cfg, empresas, filiais, tabelas):
    """As três famílias de problema que só aparecem quando alguém confere:

      1. CÓDIGO que não fecha com o leiaute (tamanho, ou filial que não começa
         pelo código da empresa dona). É o erro que só aparece na carga.
      2. Tabela DEFINIDA fora do padrão sugerido — não é erro, é decisão; entra
         como aviso para a ata da reunião, nunca como bloqueio.
      3. REGRAS da TDN (cockpit.estrut_regras): movimento mais compartilhado que
         o cadastro (1:n) e pares que têm de ser idênticos. Essa é a que quebra
         o go-live em silêncio.
    """
    niveis = _estrut_niveis(cfg["leiaute"])
    tam = _estrut_tam_esperado(cfg)
    tam_emp = cfg.get("tam_empresa") or 0
    al = []
    emp_por_id = {e["id"]: e for e in empresas}

    for e in empresas:
        if tam_emp and len(e["codigo"] or "") != tam_emp:
            al.append({"nivel": "erro", "onde": "empresa", "chave": e["codigo"],
                       "texto": f"Código da empresa com {len(e['codigo'] or '')} caractere(s); "
                                f"o leiaute {cfg['leiaute']} espera {tam_emp}."})
    for f in filiais:
        cod = f["codigo"] or ""
        if len(cod) != tam:
            al.append({"nivel": "erro", "onde": "filial", "chave": cod,
                       "texto": f"Código da filial com {len(cod)} caractere(s); "
                                f"o leiaute {cfg['leiaute']} espera {tam}."})
        dona = emp_por_id.get(f["empresa_id"])
        if tam_emp and dona and not cod.startswith(dona["codigo"] or ""):
            al.append({"nivel": "erro", "onde": "filial", "chave": cod,
                       "texto": f"A filial é da empresa {dona['codigo']} mas o código não "
                                f"começa por {dona['codigo']}."})
    if not any(f.get("matriz") for f in filiais) and filiais:
        al.append({"nivel": "aviso", "onde": "filial", "chave": "—",
                   "texto": "Nenhuma filial marcada como matriz."})

    por_tab = {t["tabela"]: t for t in tabelas if t.get("compart")}
    for t in tabelas:
        if t.get("compart") and t.get("sugestao") and t["compart"] != t["sugestao"]:
            al.append({"nivel": "info", "onde": "tabela", "chave": t["tabela"],
                       "texto": f"Definida como {t['compart']} — a sugestão era {t['sugestao']}. "
                                f"Registre o porquê na observação."})
        if not t.get("compart"):
            al.append({"nivel": "aviso", "onde": "tabela", "chave": t["tabela"],
                       "texto": "Sem compartilhamento definido."})

    for r in q("select * from cockpit.estrut_regras"):
        a, b = por_tab.get(r["tab_a"]), por_tab.get(r["tab_b"])
        if not (a and b):
            continue                      # regra sobre tabela que este projeto não usa
        if r["tipo"] == "nao_mais" and _estrut_mais_compart(b["compart"], a["compart"], niveis):
            al.append({"nivel": "erro", "onde": "regra", "chave": f"{r['tab_b']} x {r['tab_a']}",
                       "texto": f"{r['tab_b']} ({b['compart']}) está MAIS compartilhada que "
                                f"{r['tab_a']} ({a['compart']}). {r['motivo']}",
                       "fonte": r["fonte"], "tabelas": [r["tab_a"], r["tab_b"]]})
        if r["tipo"] == "igual" and a["compart"] != b["compart"]:
            al.append({"nivel": "erro", "onde": "regra", "chave": f"{r['tab_a']} x {r['tab_b']}",
                       "texto": f"{r['tab_a']} ({a['compart']}) e {r['tab_b']} ({b['compart']}) "
                                f"precisam ser iguais. {r['motivo']}",
                       "fonte": r["fonte"], "tabelas": [r["tab_a"], r["tab_b"]]})
    ordem = {"erro": 0, "aviso": 1, "info": 2}
    return sorted(al, key=lambda x: (ordem.get(x["nivel"], 9), x["onde"], str(x["chave"])))


# ── Trilha de alterações e propostas (01/10/2026) ──────────────────────────
# Spec: docs/specs/2026-10-01-estrutura-alteracoes-e-propostas.md
# Dois consultores mexem na mesma definição. A trilha diz QUEM trocou o
# compartilhamento (o ✓ da linha); a proposta deixa o segundo consultor
# registrar a opinião dele SEM sobrescrever a do primeiro, e a conversa
# acontece em cima do registro, não de memória.

def _estrut_hist(tabela_id, customer, tabela, antes, depois, origem="edicao",
                 proposta_id=None):
    """Grava UMA troca real. Sem mudança, não grava: reabrir o modal e salvar
    igual não pode virar 'alterado por fulano' e acusar quem só olhou."""
    if (antes or None) == (depois or None):
        return
    execute("""insert into cockpit.estrut_tabelas_hist
                 (tabela_id, customer, tabela, compart_antes, compart_depois, origem,
                  proposta_id, alterado_por)
               values (%s,%s,%s,%s,%s,%s,%s,%s)""",
            (tabela_id, customer, tabela, antes or None, depois or None, origem,
             proposta_id, current_user()))


def _estrut_anexa_trilha(customer, tabelas):
    """Pendura em cada tabela a trilha (mais nova primeiro) e as propostas.
    Duas queries para o cliente inteiro — por linha seriam ~130 idas ao banco."""
    hist, props = {}, {}
    for h in q("""select id, tabela_id, compart_antes, compart_depois, origem,
                         alterado_por, alterado_em
                    from cockpit.estrut_tabelas_hist where customer=%s
                   order by alterado_em desc""", (customer,)) or []:
        hist.setdefault(h["tabela_id"], []).append(h)
    for p in q("""select id, tabela_id, compart, compart_na_hora, motivo, autor, criado_em,
                         status, resolvido_por, resolvido_em, resolucao
                    from cockpit.estrut_propostas where customer=%s
                   order by criado_em desc""", (customer,)) or []:
        props.setdefault(p["tabela_id"], []).append(p)
    for t in tabelas:
        t["historico"] = hist.get(t["id"], [])
        t["propostas"] = props.get(t["id"], [])


@app.post("/api/estrutura/<customer>/proposta")
def api_estrutura_proposta(customer):
    """Abre, aplica ou descarta uma proposta de compartilhamento.
      {tabela_id, compart, motivo}            -> abre
      {id, acao: 'aplicar'|'descartar', resolucao} -> resolve
    Aplicar troca a definição E grava a trilha com origem 'proposta', para o ✓
    contar que a mudança saiu de uma discussão, não de um clique solto."""
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    if b.get("id"):
        acao = b.get("acao")
        if acao not in ("aplicar", "descartar"):
            return _err(422, "Ação inválida: use aplicar ou descartar.")
        p = q("""select * from cockpit.estrut_propostas
                  where id=%s and customer=%s""", (b["id"], customer), one=True)
        if not p:
            return _err(404, "Proposta não encontrada.")
        if p["status"] != "aberta":
            return _err(409, f"Esta proposta já foi {p['status']}.")
        resolucao = (b.get("resolucao") or "").strip() or None
        if acao == "descartar" and not resolucao:
            return _err(422, "Diga por que a proposta foi descartada — é o que fica para a ata.")
        if acao == "aplicar":
            t = q("select * from cockpit.estrut_tabelas where id=%s and customer=%s",
                  (p["tabela_id"], customer), one=True)
            if not t:
                return _err(404, "A tabela da proposta não existe mais na definição.")
            execute("""update cockpit.estrut_tabelas set compart=%s, definido_por=%s,
                         definido_em=now() where id=%s""",
                    (p["compart"], current_user(), t["id"]))
            _estrut_hist(t["id"], customer, t["tabela"], t["compart"], p["compart"],
                         "proposta", p["id"])
        execute("""update cockpit.estrut_propostas set status=%s, resolvido_por=%s,
                     resolvido_em=now(), resolucao=%s where id=%s""",
                ("aplicada" if acao == "aplicar" else "descartada", current_user(),
                 resolucao, p["id"]))
        return _json({"ok": True})
    compart = (b.get("compart") or "").strip().upper()
    motivo = (b.get("motivo") or "").strip()
    if not _estrut_compart_ok(compart):
        return _err(422, "Proposta precisa de 3 posições E/C (Empresa·Unidade·Filial).")
    if not motivo:
        return _err(422, "Explique o motivo da proposta — sem ele não há o que discutir.")
    t = q("select * from cockpit.estrut_tabelas where id=%s and customer=%s and ativo",
          (b.get("tabela_id"), customer), one=True)
    if not t:
        return _err(404, "Tabela não encontrada na definição deste cliente.")
    if compart == (t["compart"] or ""):
        return _err(422, f"{t['tabela']} já está definida como {compart}.")
    execute("""insert into cockpit.estrut_propostas
                 (tabela_id, customer, tabela, compart, compart_na_hora, motivo, autor)
               values (%s,%s,%s,%s,%s,%s,%s)""",
            (t["id"], customer, t["tabela"], compart, t["compart"], motivo, current_user()))
    return _json({"ok": True})


@app.get("/api/estrutura/<customer>")
def api_estrutura(customer):
    """Tudo da aba numa requisição: sem isso a tela faria cinco chamadas e o
    consultor veria a árvore montar em pedaços na frente do cliente."""
    if (r := require_auth()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    cfg = _estrut_cfg(customer)
    grupos = q("select * from cockpit.estrut_grupos where customer=%s and ativo "
               "order by ordem, codigo", (customer,))
    empresas = q("select * from cockpit.estrut_empresas where customer=%s and ativo "
                 "order by ordem, codigo", (customer,))
    filiais = q("select * from cockpit.estrut_filiais where customer=%s and ativo "
                "order by ordem, codigo", (customer,))
    tabelas = q("select * from cockpit.estrut_tabelas where customer=%s and ativo "
                "order by modulo, tabela", (customer,))
    catalogo = q("select modulo, count(*) n from cockpit.estrut_catalogo "
                 "group by modulo order by min(ordem), modulo")
    _estrut_anexa_trilha(customer, tabelas)
    # Referencia, so leitura: o que a aba Empresas ja leu da base do cliente
    # (monitemp_sm0). NAO alimenta a arvore acima - serve para o consultor ver
    # lado a lado o que foi combinado e o que a base tem. Producao na frente;
    # sem ela, a de testes, porque em projeto novo so existe a de testes.
    sm0, sm0_amb = [], None
    for amb in ("producao", "teste"):
        linhas = q("""select empresa, filial, nome_empresa, nome_filial, cnpj, leiaute,
                             sizefil, lido_em, updated_by
                        from cockpit.monitemp_sm0 where customer=%s and ambiente=%s
                       order by empresa, filial""", (customer, amb))
        if linhas:
            sm0, sm0_amb = linhas, amb
            break
    return _json({"ok": True, "customer": customer, "config": cfg,
                  "sm0": sm0, "sm0_ambiente": sm0_amb,
                  "leiautes": [{"id": k, **{x: v[x] for x in ("label", "ajuda")}}
                               for k, v in ESTRUT_LEIAUTES.items()],
                  "niveis": _estrut_niveis(cfg["leiaute"]),
                  "niveis_label": list(ESTRUT_NIVEL_LABEL),
                  "grupos": grupos, "empresas": empresas, "filiais": filiais,
                  "tabelas": tabelas, "catalogo": catalogo,
                  "alertas": _estrut_alertas(cfg, empresas, filiais, tabelas),
                  "interno": eh_interno()})


@app.get("/api/estrutura/catalogo/<modulo>")
def api_estrutura_catalogo(modulo):
    if (r := require_interno()):
        return r
    return _json({"ok": True, "modulo": modulo,
                  "tabelas": q("select * from cockpit.estrut_catalogo where modulo=%s "
                               "order by tipo desc, tabela", (modulo,))})


@app.post("/api/estrutura/<customer>/config")
def api_estrutura_config(customer):
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    leiaute = (b.get("leiaute") or "EEFF").upper()
    if leiaute not in ESTRUT_LEIAUTES:
        return _err(422, "Leiaute inválido.")
    t = ESTRUT_LEIAUTES[leiaute]["tam"]
    execute("""insert into cockpit.estrut_config
                 (customer, leiaute, tam_empresa, tam_unidade, tam_filial,
                  observacao, definido_por, definido_em)
               values (%s,%s,%s,%s,%s,%s,%s, now())
               on conflict (customer) do update
                  set leiaute=excluded.leiaute, tam_empresa=excluded.tam_empresa,
                      tam_unidade=excluded.tam_unidade, tam_filial=excluded.tam_filial,
                      observacao=excluded.observacao, definido_por=excluded.definido_por,
                      definido_em=now()""",
            (customer, leiaute,
             int(b.get("tam_empresa") or t[0]), int(b.get("tam_unidade") or t[1]),
             int(b.get("tam_filial") or t[2]),
             (b.get("observacao") or "").strip() or None, current_user()))
    return _json({"ok": True, "config": _estrut_cfg(customer)})


def _estrut_txt(b, campo, obrig=False):
    v = (b.get(campo) or "").strip()
    if obrig and not v:
        raise ValueError(f"Informe {campo}.")
    return v or None


@app.post("/api/estrutura/<customer>/grupo")
def api_estrutura_grupo(customer):
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    if b.get("excluir") and b.get("id"):
        # Inativar, não apagar: o grupo pode já ter virado decisão em ata.
        execute("update cockpit.estrut_grupos set ativo=false where id=%s and customer=%s",
                (b["id"], customer))
        return _json({"ok": True})
    try:
        codigo, nome = _estrut_txt(b, "codigo", True), _estrut_txt(b, "nome", True)
    except ValueError as e:
        return _err(422, str(e))
    if b.get("id"):
        execute("""update cockpit.estrut_grupos set codigo=%s, nome=%s, observacao=%s,
                     ordem=%s where id=%s and customer=%s""",
                (codigo, nome, _estrut_txt(b, "observacao"), int(b.get("ordem") or 0),
                 b["id"], customer))
    else:
        execute("""insert into cockpit.estrut_grupos
                     (customer, codigo, nome, observacao, ordem, created_by)
                   values (%s,%s,%s,%s,%s,%s)
                   on conflict (customer, codigo) do update
                      set nome=excluded.nome, observacao=excluded.observacao, ativo=true""",
                (customer, codigo, nome, _estrut_txt(b, "observacao"),
                 int(b.get("ordem") or 0), current_user()))
    return _json({"ok": True})


@app.post("/api/estrutura/<customer>/empresa")
def api_estrutura_empresa(customer):
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    if b.get("excluir") and b.get("id"):
        execute("update cockpit.estrut_empresas set ativo=false where id=%s and customer=%s",
                (b["id"], customer))
        return _json({"ok": True})
    try:
        codigo, nome = _estrut_txt(b, "codigo", True), _estrut_txt(b, "nome", True)
    except ValueError as e:
        return _err(422, str(e))
    campos = (b.get("grupo_id") or None, codigo, nome, _estrut_txt(b, "cnpj"),
              (_estrut_txt(b, "uf") or "").upper() or None, _estrut_txt(b, "observacao"),
              int(b.get("ordem") or 0))
    if b.get("id"):
        execute("""update cockpit.estrut_empresas set grupo_id=%s, codigo=%s, nome=%s,
                     cnpj=%s, uf=%s, observacao=%s, ordem=%s
                   where id=%s and customer=%s""", campos + (b["id"], customer))
    else:
        execute("""insert into cockpit.estrut_empresas
                     (customer, grupo_id, codigo, nome, cnpj, uf, observacao, ordem, created_by)
                   values (%s,%s,%s,%s,%s,%s,%s,%s,%s)
                   on conflict (customer, codigo) do update
                      set grupo_id=excluded.grupo_id, nome=excluded.nome, cnpj=excluded.cnpj,
                          uf=excluded.uf, observacao=excluded.observacao, ativo=true""",
                (customer,) + campos + (current_user(),))
    return _json({"ok": True})


@app.post("/api/estrutura/<customer>/filial")
def api_estrutura_filial(customer):
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    if b.get("excluir") and b.get("id"):
        execute("update cockpit.estrut_filiais set ativo=false where id=%s and customer=%s",
                (b["id"], customer))
        return _json({"ok": True})
    if not b.get("empresa_id"):
        return _err(422, "A filial precisa pertencer a uma empresa.")
    try:
        codigo, nome = _estrut_txt(b, "codigo", True), _estrut_txt(b, "nome", True)
    except ValueError as e:
        return _err(422, str(e))
    # Matriz é UMA por empresa: duas matrizes viram divergência silenciosa no
    # cadastro da SM0 e ninguém repara até a emissão do primeiro documento.
    if b.get("matriz"):
        execute("""update cockpit.estrut_filiais set matriz=false
                    where customer=%s and empresa_id=%s""", (customer, b["empresa_id"]))
    campos = (b["empresa_id"], codigo, _estrut_txt(b, "unidade"), nome,
              _estrut_txt(b, "cnpj"), (_estrut_txt(b, "uf") or "").upper() or None,
              _estrut_txt(b, "municipio"), bool(b.get("matriz")),
              _estrut_txt(b, "observacao"), int(b.get("ordem") or 0))
    if b.get("id"):
        execute("""update cockpit.estrut_filiais set empresa_id=%s, codigo=%s, unidade=%s,
                     nome=%s, cnpj=%s, uf=%s, municipio=%s, matriz=%s, observacao=%s, ordem=%s
                   where id=%s and customer=%s""", campos + (b["id"], customer))
    else:
        execute("""insert into cockpit.estrut_filiais
                     (customer, empresa_id, codigo, unidade, nome, cnpj, uf, municipio,
                      matriz, observacao, ordem, created_by)
                   values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                   on conflict (customer, codigo) do update
                      set empresa_id=excluded.empresa_id, unidade=excluded.unidade,
                          nome=excluded.nome, cnpj=excluded.cnpj, uf=excluded.uf,
                          municipio=excluded.municipio, matriz=excluded.matriz,
                          observacao=excluded.observacao, ativo=true""",
                (customer,) + campos + (current_user(),))
    return _json({"ok": True})


@app.post("/api/estrutura/<customer>/semear")
def api_estrutura_semear(customer):
    """Traz um módulo inteiro do catálogo para a definição do cliente.
    'do nothing' no conflito de propósito: semear de novo NÃO pode apagar o que
    a consultoria já decidiu — semeadura é ponto de partida, não reset."""
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    mods = [m for m in (b.get("modulos") or []) if isinstance(m, str)]
    if not mods:
        return _err(422, "Escolha ao menos um módulo.")
    linhas = q("select * from cockpit.estrut_catalogo where modulo = any(%s::text[]) "
               "order by modulo, tabela", (mods,))
    if not linhas:
        return _json({"ok": True, "incluidas": 0})
    usar_sugestao = bool(b.get("usar_sugestao", True))
    with db() as conn, conn.cursor() as cur:
        execute_values(cur,
            """insert into cockpit.estrut_tabelas
                 (customer, modulo, tabela, descricao, tipo, compart, sugestao, origem,
                  definido_por, definido_em)
               values %s on conflict (customer, modulo, tabela) do nothing""",
            [(customer, l["modulo"], l["tabela"], l["descricao"], l["tipo"],
              l["sugestao"] if usar_sugestao else None, l["sugestao"], "catalogo",
              current_user())
             for l in linhas],
            template="(%s,%s,%s,%s,%s,%s,%s,%s,%s, now())")
    return _json({"ok": True, "incluidas": len(linhas), "modulos": mods})


@app.post("/api/estrutura/<customer>/tabela")
def api_estrutura_tabela(customer):
    """Inclusão manual e edição da definição. A inclusão manual é obrigatória de
    verdade: tabela customizada (Z*) e alias de módulo que não está no catálogo
    aparecem em todo projeto, e sem isso a planilha volta a ser do Excel."""
    if (r := require_interno()):
        return r
    if (d := deny_aba(customer, "estrutura")):
        return d
    b = request.get_json(silent=True) or {}
    if b.get("excluir") and b.get("id"):
        execute("update cockpit.estrut_tabelas set ativo=false where id=%s and customer=%s",
                (b["id"], customer))
        return _json({"ok": True})
    compart = (b.get("compart") or "").strip().upper() or None
    if compart and not _estrut_compart_ok(compart):
        return _err(422, "Compartilhamento deve ter 3 posições E/C (Empresa·Unidade·Filial), "
                         "por exemplo ECC ou CCC.")
    if b.get("id"):
        antes = q("select tabela, compart from cockpit.estrut_tabelas where id=%s and customer=%s",
                  (b["id"], customer), one=True)
        if not antes:
            return _err(404, "Tabela não encontrada.")
        execute("""update cockpit.estrut_tabelas set descricao=%s, tipo=%s, compart=%s,
                     observacao=%s, definido_por=%s, definido_em=now()
                   where id=%s and customer=%s""",
                (_estrut_txt(b, "descricao"), b.get("tipo") or "cadastro", compart,
                 _estrut_txt(b, "observacao"), current_user(), b["id"], customer))
        _estrut_hist(b["id"], customer, antes["tabela"], antes["compart"], compart)
        return _json({"ok": True})
    tabela = (b.get("tabela") or "").strip().upper()
    modulo = (b.get("modulo") or "").strip()
    if not tabela or not modulo:
        return _err(422, "Informe o módulo e o nome da tabela.")
    antes = q("""select compart, ativo from cockpit.estrut_tabelas
                  where customer=%s and modulo=%s and tabela=%s""",
              (customer, modulo, tabela), one=True)
    execute("""insert into cockpit.estrut_tabelas
                 (customer, modulo, tabela, descricao, tipo, compart, origem,
                  observacao, definido_por, definido_em)
               values (%s,%s,%s,%s,%s,%s,'manual',%s,%s, now())
               on conflict (customer, modulo, tabela) do update
                  set descricao=excluded.descricao, tipo=excluded.tipo,
                      compart=excluded.compart, observacao=excluded.observacao,
                      ativo=true, definido_por=excluded.definido_por, definido_em=now()""",
            (customer, modulo, tabela, _estrut_txt(b, "descricao"),
             b.get("tipo") or "cadastro", compart, _estrut_txt(b, "observacao"),
             current_user()))
    # Inclusão nova não é "alteração" (ninguém tinha decidido antes); regravar
    # por cima de uma existente com outro compartilhamento, é.
    if antes and antes["compart"] and antes["compart"] != compart:
        novo = q("""select id from cockpit.estrut_tabelas
                     where customer=%s and modulo=%s and tabela=%s""",
                 (customer, modulo, tabela), one=True)
        _estrut_hist(novo["id"], customer, tabela, antes["compart"], compart, "manual")
    return _json({"ok": True})


# ── Compara: dicionário (SX2/SX3/SX6) entre as empresas da MESMA base ───────
# A pergunta desta tela é "as empresas desta base foram montadas iguais?". No
# Protheus o dicionário é por GRUPO DE EMPRESAS: SX2010, SX2020 e SX2030 são
# dicionários DIFERENTES, e é aí que mora a divergência que ninguém vê até o
# cliente reclamar que o campo existe numa empresa e não na outra.
#
# Transporte: o /tscmonit/query que já existe (só leitura). A agregação do X3 e
# do X6 roda NO BANCO de propósito — trazer o SX3 inteiro de três empresas são
# ~90 mil linhas, acima do limite e do tempo da função da Vercel. Volta só o que
# DIVERGE. O X2 vem inteiro porque "a tabela existe só na empresa 1 e 2" é uma
# pergunta sobre AUSÊNCIA: sem a lista completa não dá para saber quem falta.
CMP_RE_SUFIXO = re.compile(r"^SX([236])(\d{3})$")
CMP_SQL_SUFIXOS = ("SELECT TABLE_NAME AS TABELA FROM ALL_TABLES "
                   "WHERE TABLE_NAME LIKE 'SX2%' OR TABLE_NAME LIKE 'SX3%' "
                   "OR TABLE_NAME LIKE 'SX6%' ORDER BY TABLE_NAME")


def _cmp_union(sufixos, corpo):
    return "\n  UNION ALL\n".join(corpo(s) for s in sufixos)


def _cmp_sql_x2(sufixos):
    corpo = lambda s: (
        f"  SELECT '{s}' AS EMP, RTRIM(X2_CHAVE) AS CHAVE, RTRIM(X2_NOME) AS NOME, "
        f"RTRIM(X2_ARQUIVO) AS ARQ, RTRIM(X2_MODOEMP) AS MODOEMP, "
        f"RTRIM(X2_MODOUN) AS MODOUN, RTRIM(X2_MODO) AS MODO "
        f"FROM SX2{s} WHERE D_E_L_E_T_ = ' '")
    return ("WITH T AS (\n" + _cmp_union(sufixos, corpo) + "\n)\n"
            "SELECT CHAVE, EMP, NOME, ARQ, MODOEMP, MODOUN, MODO FROM T "
            "ORDER BY CHAVE, EMP")


def _cmp_sql_x3(sufixos):
    """Campos por tabela. Só volta (arquivo, campo) que falta em alguma empresa
    OU que muda de tipo/tamanho/decimal entre elas — o resto é ruído."""
    corpo = lambda s: (
        f"  SELECT '{s}' AS EMP, RTRIM(X3_ARQUIVO) AS ARQ, RTRIM(X3_CAMPO) AS CAMPO, "
        f"RTRIM(X3_TIPO) AS TIPO, TO_CHAR(X3_TAMANHO) AS TAM, "
        f"TO_CHAR(X3_DECIMAL) AS DECIMAIS, RTRIM(SUBSTR(X3_TITULO, 1, 40)) AS TITULO "
        f"FROM SX3{s} WHERE D_E_L_E_T_ = ' '")
    return ("WITH T AS (\n" + _cmp_union(sufixos, corpo) + "\n),\n"
            "A AS (SELECT ARQ, CAMPO, COUNT(DISTINCT EMP) AS EMPS, "
            "COUNT(DISTINCT TIPO || '/' || TAM || '/' || DECIMAIS) AS VARI "
            "FROM T GROUP BY ARQ, CAMPO)\n"
            "SELECT T.ARQ, T.CAMPO, T.EMP, T.TIPO, T.TAM, T.DECIMAIS, T.TITULO "
            "FROM T JOIN A ON A.ARQ = T.ARQ AND A.CAMPO = T.CAMPO "
            f"WHERE A.EMPS < {len(sufixos)} OR A.VARI > 1 "
            "ORDER BY T.ARQ, T.CAMPO, T.EMP")


def _cmp_sql_x6(sufixos):
    """Parâmetros. Divergência vem em três sabores, nesta ordem: o parâmetro não
    existe na empresa; existe com compartilhamento diferente (X6_FIL em branco =
    vale para todas as filiais, preenchido = só naquela); existe com conteúdo
    diferente. O ';' do conteúdo vira ',' porque o transporte é CSV com ';'."""
    cont = "RTRIM(REPLACE(REPLACE(SUBSTR(X6_CONTEUD, 1, 180), ';', ','), CHR(10), ' '))"
    corpo = lambda s: (
        f"  SELECT '{s}' AS EMP, RTRIM(X6_VAR) AS PARAM, RTRIM(X6_FIL) AS FIL, "
        f"RTRIM(X6_TIPO) AS TIPO, {cont} AS CONTEUDO, "
        f"RTRIM(SUBSTR(X6_DESCRIC, 1, 60)) AS DESCR "
        f"FROM SX6{s} WHERE D_E_L_E_T_ = ' '")
    return ("WITH T AS (\n" + _cmp_union(sufixos, corpo) + "\n),\n"
            "A AS (SELECT PARAM, COUNT(DISTINCT EMP) AS EMPS, COUNT(DISTINCT FIL) AS FILS, "
            "COUNT(DISTINCT CONTEUDO) AS CONTS FROM T GROUP BY PARAM)\n"
            "SELECT T.PARAM, T.EMP, T.FIL, T.TIPO, T.CONTEUDO, T.DESCR "
            "FROM T JOIN A ON A.PARAM = T.PARAM "
            f"WHERE A.EMPS < {len(sufixos)} OR A.FILS > 1 OR A.CONTS > 1 "
            "ORDER BY T.PARAM, T.EMP, T.FIL")


def _cmp_cel(linha, i):
    return (linha[i] if i < len(linha) and linha[i] is not None else "").strip()


def _cmp_agrega_x2(linhas, sufixos):
    itens = {}
    for ln in linhas[1:]:
        chave = _cmp_cel(ln, 0)
        if not chave:
            continue
        it = itens.setdefault(chave, {"chave": chave, "nome": _cmp_cel(ln, 2),
                                      "arquivo": _cmp_cel(ln, 3), "empresas": {}})
        it["empresas"][_cmp_cel(ln, 1)] = (_cmp_cel(ln, 4) + _cmp_cel(ln, 5)
                                           + _cmp_cel(ln, 6))
        it["nome"] = it["nome"] or _cmp_cel(ln, 2)
    saida = []
    for it in itens.values():
        presentes = [s for s in sufixos if s in it["empresas"]]
        modos = {it["empresas"][s] for s in presentes}
        it["presentes"] = presentes
        it["falta_em"] = [s for s in sufixos if s not in it["empresas"]]
        it["modo_diverge"] = len(modos) > 1
        saida.append(it)
    saida.sort(key=lambda x: x["chave"])
    return saida


def _cmp_agrega_x3(linhas, sufixos):
    tab = {}
    for ln in linhas[1:]:
        arq, campo, emp = _cmp_cel(ln, 0), _cmp_cel(ln, 1), _cmp_cel(ln, 2)
        if not arq or not campo:
            continue
        t = tab.setdefault(arq, {"arquivo": arq, "campos": {}})
        c = t["campos"].setdefault(campo, {"campo": campo, "titulo": _cmp_cel(ln, 6),
                                           "empresas": {}})
        c["empresas"][emp] = {"tipo": _cmp_cel(ln, 3), "tamanho": _cmp_cel(ln, 4),
                              "decimal": _cmp_cel(ln, 5)}
        c["titulo"] = c["titulo"] or _cmp_cel(ln, 6)
    saida = []
    for t in tab.values():
        campos = []
        for c in t["campos"].values():
            c["falta_em"] = [s for s in sufixos if s not in c["empresas"]]
            fmts = {f"{v['tipo']}/{v['tamanho']}/{v['decimal']}" for v in c["empresas"].values()}
            c["formato_diverge"] = len(fmts) > 1
            campos.append(c)
        campos.sort(key=lambda x: x["campo"])
        t["campos"] = campos
        t["ausentes"] = sum(1 for c in campos if c["falta_em"])
        t["formatos"] = sum(1 for c in campos if c["formato_diverge"])
        saida.append(t)
    saida.sort(key=lambda x: x["arquivo"])
    return saida


def _cmp_agrega_x6(linhas, sufixos):
    itens = {}
    for ln in linhas[1:]:
        param = _cmp_cel(ln, 0)
        if not param:
            continue
        it = itens.setdefault(param, {"param": param, "descricao": _cmp_cel(ln, 5),
                                      "empresas": {}})
        it["empresas"].setdefault(_cmp_cel(ln, 1), []).append(
            {"filial": _cmp_cel(ln, 2), "tipo": _cmp_cel(ln, 3),
             "conteudo": _cmp_cel(ln, 4)})
        it["descricao"] = it["descricao"] or _cmp_cel(ln, 5)
    saida = []
    for it in itens.values():
        it["falta_em"] = [s for s in sufixos if s not in it["empresas"]]
        # Compartilhamento = o CONJUNTO de filiais em que o parâmetro existe.
        # Empresa com {''} tem o parâmetro global; {'0101'} tem só numa filial.
        compart = {s: sorted({r["filial"] for r in rs})
                   for s, rs in it["empresas"].items()}
        it["compartilhamento"] = compart
        it["compart_diverge"] = len({tuple(v) for v in compart.values()}) > 1
        conteudos = {r["conteudo"] for rs in it["empresas"].values() for r in rs}
        it["conteudo_diverge"] = len(conteudos) > 1
        saida.append(it)
    saida.sort(key=lambda x: x["param"])
    return saida


@app.post("/api/estrutura/<customer>/comparar")
def api_estrutura_comparar(customer):
    """Lê SX2/SX3/SX6 de todas as empresas da base e devolve as divergências.

        POST /api/estrutura/TFEHXQ00/comparar?ambiente=producao
    """
    if (r := require_interno()):
        return r
    if (d := deny_customer(customer)):
        return d
    if (d := deny_aba(customer, "compara")):
        return d
    amb = _amb_path(request.args.get("ambiente") or "producao")
    if not amb:
        return _err(400, "Ambiente inválido — use 'producao' ou 'teste'.")
    cfg = _cfg_ambiente(customer, amb)
    if not cfg:
        return _err(404, f"Ambiente {amb} ainda não cadastrado para {customer} — "
                         "preencha a REST em ⚙ Ambiente antes de comparar.")
    if not cfg.get("ativo"):
        return _err(409, f"Ambiente {amb} está inativo.")
    banco = (cfg.get("banco") or "").lower()
    if banco and banco != "oracle":
        return _err(409, "A comparação de dicionário só tem SQL de Oracle por enquanto; "
                         f"este ambiente está cadastrado como {banco}.")

    ini = time.time()
    try:
        heads, auth, tok, tmo = _sessao_protheus(cfg, customer, amb)
    except RuntimeError as e:
        return _err(500, str(e))
    if not tok:
        return _err(409, "Sem token cadastrado para este ambiente — a REST da base "
                         "recusa a chamada sem o X-TSC-Token.")
    prazo = ini + min(int(cfg.get("timeout_s") or 45), TIMEOUT_TETO)

    def roda(sql, rotulo):
        resta = prazo - time.time()
        if resta < 5:
            raise RuntimeError(
                f"Tempo esgotado antes de ler {rotulo}. A função da Vercel morre em 60s "
                "— rode numa base com menos empresas ou fora do horário de pico.")
        r = requests.post(f"{cfg['url_rest']}/query",
                          json={"token": tok, "tipo": "cadastros", "sql": sql,
                                "limite": int(cfg.get("limite_linhas") or 50000)},
                          headers=heads, auth=auth,
                          params={"empresa": cfg.get("empresa") or "01",
                                  "filial": cfg.get("filial") or "01"},
                          timeout=max(5, int(resta)))
        dados = _json_da_resposta(r)
        if r.status_code >= 400 or not dados.get("ok"):
            raise RuntimeError(dados.get("erro")
                               or f"HTTP {r.status_code} ao ler {rotulo}: {_trecho(r)}")
        if dados.get("truncado"):
            raise RuntimeError(f"{rotulo} bateu o limite de linhas — comparação truncada é "
                               "comparação errada. Suba o limite no ⚙ Ambiente.")
        return _proto_linhas_csv(dados.get("csv") or "")

    try:
        achados = roda(CMP_SQL_SUFIXOS, "a lista de dicionários")
        por_dic = {"2": set(), "3": set(), "6": set()}
        for ln in achados[1:]:
            m = CMP_RE_SUFIXO.match(_cmp_cel(ln, 0).upper())
            if m:
                por_dic[m.group(1)].add(m.group(2))
        # Só compara o sufixo que tem os TRÊS dicionários: um SX3 sem o SX2 do
        # mesmo grupo é sobra de migração, não empresa.
        sufixos = sorted(por_dic["2"] & por_dic["3"] & por_dic["6"])
        parciais = sorted((por_dic["2"] | por_dic["3"] | por_dic["6"]) - set(sufixos))
        if len(sufixos) < 2:
            return _err(409, "Encontrei menos de dois dicionários completos nesta base "
                             f"(SX2/SX3/SX6 por grupo de empresas): {sufixos or 'nenhum'}. "
                             "Sem dois não há o que comparar.")
        x2 = _cmp_agrega_x2(roda(_cmp_sql_x2(sufixos), "o SX2"), sufixos)
        x3 = _cmp_agrega_x3(roda(_cmp_sql_x3(sufixos), "o SX3"), sufixos)
        x6 = _cmp_agrega_x6(roda(_cmp_sql_x6(sufixos), "o SX6"), sufixos)
    except requests.RequestException as e:
        return _err(502, f"Falha ao chamar a base: {type(e).__name__}. "
                         f"Timeout do ambiente = {cfg.get('timeout_s')}s.")
    except RuntimeError as e:
        return _err(502, str(e))

    return _json({
        "ok": True, "customer": customer, "ambiente": amb,
        "url_rest": cfg.get("url_rest"), "environment": cfg.get("environment"),
        # Epoch, não string: o servidor roda em UTC na Vercel e quem precisa da
        # hora de Florianópolis é a tela — ela formata no fuso do navegador.
        "lido_em_ts": int(time.time()),
        "duracao_ms": int((time.time() - ini) * 1000),
        # sufixo 010 = grupo de empresas 01; o nome vem do SM0 que a aba já tem.
        "empresas": [{"sufixo": s, "codigo": s[:-1]} for s in sufixos],
        "sufixos_parciais": parciais,
        "x2": x2, "x3": x3, "x6": x6,
    })


# ── estáticos do web/ (assets) ──────────────────────────────────────────────
@app.get("/<path:asset>")
def static_assets(asset):
    if asset.startswith("api/"):
        return _err(404, "Rota de API desconhecida.")
    safe = (WEB_DIR / asset).resolve()
    if WEB_DIR in safe.parents and safe.is_file():
        # Nenhum HTML (fora o login) sai sem sessão — inclui transicao.html
        if safe.suffix.lower() == ".html" and safe.name != "login.html" and not current_user():
            return redirect("/login", 302)
        ext = safe.suffix.lower()
        ctype = {".html": "text/html; charset=utf-8", ".js": "application/javascript",
                 ".css": "text/css", ".json": "application/json", ".png": "image/png",
                 ".jpg": "image/jpeg", ".svg": "image/svg+xml", ".ico": "image/x-icon",
                 ".woff2": "font/woff2"}.get(ext, "application/octet-stream")
        return Response(safe.read_bytes(), mimetype=ctype)
    return _err(404, "Não encontrado.")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5055)), debug=True)
